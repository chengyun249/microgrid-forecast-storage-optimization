#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""第四问4-2：独立训练、融合、价格预测、购电与储能控制、费用核对及Excel导出。

普通Python单文件，不导入旧项目模块；仅需公开依赖和显式指定的附件。
archive读取四分量基础预测，其余融合、价格预测、1月预运行和全年调度全部重算。
retrain从原附件重新训练基础四分量。研究配置曾按全年费用选择，详见交付说明。
"""
from __future__ import annotations
ATTACHMENT2_PATH = r""
ATTACHMENT3_PATH = r""
ATTACHMENT4_PATH = r""
RESULT_TEMPLATE_PATH = r""
OUTPUT_DIRECTORY = ""
ARCHIVED_PREDICTIONS_PATH = r""
RUN_MODE = "retrain"
TRAINING_WORKERS = 2
LIGHTGBM_THREADS = 1
QUESTION = 2
import argparse
from concurrent.futures import ProcessPoolExecutor
import copy
from dataclasses import dataclass
from datetime import date, datetime
from functools import lru_cache
import hashlib
import importlib.metadata as metadata
import json
import math
import os
from pathlib import Path
import sys
import time
from typing import Any

import lightgbm as lgb
import numpy as np
import pandas as pd
from numba import njit
from openpyxl import load_workbook
from scipy.ndimage import gaussian_filter1d
from scipy.optimize import linprog, minimize
from scipy.sparse import coo_matrix
from sklearn.linear_model import Ridge
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from threadpoolctl import threadpool_limits

PERIODS_PER_DAY = 144
FEATURE_START_INDEX = 14
START = 15
FORMAL = 31
DT = 1 / 6
RANDOM_STATE = 2026
ETA = 0.9
S_MIN, S_MAX, P_MAX = 1200., 10800., 5000. / 6
BLOCK = np.arange(144) // 24
COMPONENT_NAMES = ['legacy_daily_lgb', 'legacy_daily_ridge', 'periodic_lgb_l1', 'calendar_curve']

DATES=pd.date_range("2025-01-01","2025-12-31")
_RUN_INPUTS={}

# 一、数据保存和基础供需模型：逐日训练，历史月度选参。
def write_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def atomic_npz(path, **data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + '.tmp')
    with temporary.open('wb') as stream:
        np.savez_compressed(stream, **data)
    os.replace(temporary, path)


def read_npz(path):
    with np.load(path, allow_pickle=False) as z:
        return {name: z[name].copy() for name in z.files}


def required_input(value,label):
    if not str(value).strip():
        raise ValueError(label+'路径尚未填写。请编辑脚本顶部相关路径，或使用命令行参数。')
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(label+'文件不存在，请检查配置或命令行参数。')
    return path


def array_digest(values):
    """只用于核对附件是否与留档相符，不作为模型特征。"""
    return hashlib.sha256(np.ascontiguousarray(values, dtype=np.float64).view(np.uint8)).hexdigest()

@dataclass(frozen=True)
class ModelConfig:
    name: str
    family: str
    target_kind: str
    objective: str = "regression"
    n_estimators: int = 0
    learning_rate: float = 0.0
    num_leaves: int = 0
    min_child_samples: int = 0
    reg_alpha: float = 0.0
    reg_lambda: float = 0.0
    ridge_alpha: float = 0.0
    train_window_days: int | None = None


MODEL_CONFIGS: tuple[ModelConfig, ...] = (
    ModelConfig(
        "lgbm_level_l2_large",
        "lgbm",
        "level",
        objective="regression",
        n_estimators=180,
        learning_rate=0.040,
        num_leaves=63,
        min_child_samples=80,
        reg_alpha=0.20,
        reg_lambda=5.0,
    ),
    ModelConfig(
        "lgbm_level_l2_moretrees",
        "lgbm",
        "level",
        objective="regression",
        n_estimators=260,
        learning_rate=0.025,
        num_leaves=31,
        min_child_samples=60,
        reg_alpha=0.10,
        reg_lambda=2.0,
    ),
    ModelConfig(
        "lgbm_residual_l1_moretrees",
        "lgbm",
        "residual",
        objective="regression_l1",
        n_estimators=220,
        learning_rate=0.030,
        num_leaves=31,
        min_child_samples=60,
        reg_alpha=0.10,
        reg_lambda=2.0,
    ),
    ModelConfig("ridge_level_a1", "ridge", "level", ridge_alpha=1.0),
    ModelConfig("ridge_level_a10", "ridge", "level", ridge_alpha=10.0),
    ModelConfig("ridge_level_a100", "ridge", "level", ridge_alpha=100.0),
    ModelConfig("ridge_residual_a10", "ridge", "residual", ridge_alpha=10.0),
)


DEFAULT_CONFIG = {
    ("load", "lgbm"): "lgbm_level_l2_large",
    ("pv", "lgbm"): "lgbm_residual_l1_moretrees",
    ("load", "ridge"): "ridge_level_a10",
    ("pv", "ridge"): "ridge_residual_a10",
    ("load", "seasonal"): "seasonal_lag7",
    ("pv", "seasonal"): "seasonal_lag1",
}


def feature_row_from_history(history: np.ndarray, target_date: date) -> np.ndarray:
    """由目标日前的逐日曲线构造目标日特征，不读取目标日真实值。"""
    history = np.asarray(history, dtype=np.float64)
    if history.ndim != 2 or history.shape[1] != PERIODS_PER_DAY:
        raise ValueError("history 必须是 n×144 数组。")
    if len(history) < 14:
        raise ValueError("至少需要14个完整历史日。")

    slots = np.arange(PERIODS_PER_DAY, dtype=np.float64)
    slot_angle = 2.0 * np.pi * slots / PERIODS_PER_DAY
    lag1, lag2, lag3 = history[-1], history[-2], history[-3]
    lag7, lag14 = history[-7], history[-14]
    hist3, hist7, hist14 = history[-3:], history[-7:], history[-14:]
    hist28 = history[-min(28, len(history)) :]
    weekday = target_date.weekday()
    doy = target_date.timetuple().tm_yday
    weekday_angle = 2.0 * np.pi * weekday / 7.0
    doy_angle = 2.0 * np.pi * (doy - 1) / 365.25

    row = np.column_stack(
        [
            lag1,
            lag2,
            lag3,
            lag7,
            lag14,
            np.mean(hist3, axis=0),
            np.mean(hist7, axis=0),
            np.mean(hist14, axis=0),
            np.mean(hist28, axis=0),
            np.std(hist3, axis=0),
            np.std(hist7, axis=0),
            np.std(hist14, axis=0),
            np.std(hist28, axis=0),
            lag1 - lag2,
            lag7 - lag14,
            np.full(PERIODS_PER_DAY, np.mean(lag1)),
            np.full(PERIODS_PER_DAY, np.max(lag1)),
            np.full(PERIODS_PER_DAY, np.std(lag1)),
            np.full(PERIODS_PER_DAY, np.mean(lag7)),
            np.full(PERIODS_PER_DAY, np.max(lag7)),
            slots / (PERIODS_PER_DAY - 1),
            np.sin(slot_angle),
            np.cos(slot_angle),
            np.sin(2.0 * slot_angle),
            np.cos(2.0 * slot_angle),
            np.full(PERIODS_PER_DAY, weekday / 6.0),
            np.full(PERIODS_PER_DAY, float(weekday >= 5)),
            np.full(PERIODS_PER_DAY, math.sin(weekday_angle)),
            np.full(PERIODS_PER_DAY, math.cos(weekday_angle)),
            np.full(PERIODS_PER_DAY, (doy - 1) / 364.0),
            np.full(PERIODS_PER_DAY, math.sin(doy_angle)),
            np.full(PERIODS_PER_DAY, math.cos(doy_angle)),
            np.full(PERIODS_PER_DAY, target_date.month / 12.0),
        ]
    )
    if row.shape != (PERIODS_PER_DAY, 33) or not np.all(np.isfinite(row)):
        raise RuntimeError("特征行形状或数值异常。")
    return row.astype(np.float32)


def build_feature_cube(series: np.ndarray, dates: tuple[date, ...]) -> np.ndarray:
    cube = np.full((len(dates), PERIODS_PER_DAY, 33), np.nan, dtype=np.float32)
    for target_index in range(FEATURE_START_INDEX, len(dates)):
        cube[target_index] = feature_row_from_history(
            series[:target_index], dates[target_index]
        )
    if not np.all(np.isfinite(cube[FEATURE_START_INDEX:])):
        raise RuntimeError("有效预测日的特征中存在非有限值。")
    return cube


def make_model(config: ModelConfig, lgbm_threads: int) -> Any:
    if config.family == "lgbm":
        return lgb.LGBMRegressor(
            objective=config.objective,
            n_estimators=config.n_estimators,
            learning_rate=config.learning_rate,
            num_leaves=config.num_leaves,
            min_child_samples=config.min_child_samples,
            subsample=0.90,
            subsample_freq=1,
            colsample_bytree=0.90,
            reg_alpha=config.reg_alpha,
            reg_lambda=config.reg_lambda,
            max_bin=127,
            random_state=RANDOM_STATE,
            n_jobs=lgbm_threads,
            verbosity=-1,
            deterministic=True,
            force_col_wise=True,
        )
    if config.family == "ridge":
        return make_pipeline(
            StandardScaler(),
            Ridge(alpha=config.ridge_alpha, solver="cholesky"),
        )
    raise ValueError(f"未知模型族：{config.family}")


def fit_model_for_origin(
    series: np.ndarray,
    feature_cube: np.ndarray,
    config: ModelConfig,
    target_index: int,
    residual_lag: int,
    lgbm_threads: int,
) -> Any:
    train_start = FEATURE_START_INDEX
    if config.train_window_days is not None:
        train_start = max(train_start, target_index - config.train_window_days)
    train_days = np.arange(train_start, target_index, dtype=int)
    if len(train_days) == 0:
        raise RuntimeError(f"{target_index=} 没有有效训练日。")
    assert int(np.max(train_days)) < target_index
    x_train = feature_cube[train_days].reshape(-1, feature_cube.shape[-1])
    if config.target_kind == "level":
        y_train = series[train_days].reshape(-1)
    elif config.target_kind == "residual":
        if int(np.min(train_days)) - residual_lag < 0:
            raise RuntimeError("残差模型基线滞后越界。")
        y_train = (series[train_days] - series[train_days - residual_lag]).reshape(-1)
    else:
        raise ValueError(f"未知目标类型：{config.target_kind}")
    model = make_model(config, lgbm_threads)
    model.fit(x_train, y_train)
    return model


def predict_with_fitted_model(
    model: Any,
    config: ModelConfig,
    feature_row: np.ndarray,
    baseline: np.ndarray,
) -> np.ndarray:
    if config.family == "lgbm":
        raw = np.asarray(model.booster_.predict(feature_row), dtype=np.float64)
    else:
        raw = np.asarray(model.predict(feature_row), dtype=np.float64)
    if config.target_kind == "residual":
        raw = baseline + raw
    return np.clip(raw.reshape(PERIODS_PER_DAY), 0.0, None)


def fit_predict_origin(
    series: np.ndarray,
    dates: tuple[date, ...],
    feature_cube: np.ndarray,
    config: ModelConfig,
    target_index: int,
    residual_lag: int,
    lgbm_threads: int,
    include_next_day: bool,
) -> tuple[np.ndarray, np.ndarray | None]:
    """拟合 < target_index 的标签，预测 target_index 及可选递归下一日。"""
    model = fit_model_for_origin(
        series,
        feature_cube,
        config,
        target_index,
        residual_lag,
        lgbm_threads,
    )
    current_baseline = series[target_index - residual_lag]
    current = predict_with_fitted_model(
        model, config, feature_cube[target_index], current_baseline
    )
    if not include_next_day or target_index + 1 >= len(dates):
        return current, None

    augmented = np.vstack([series[:target_index], current[None, :]])
    next_feature = feature_row_from_history(augmented, dates[target_index + 1])
    next_baseline = augmented[-residual_lag]
    following = predict_with_fitted_model(model, config, next_feature, next_baseline)
    return current, following


def weighted_error_score(
    actual: np.ndarray,
    predicted: np.ndarray,
    origins: np.ndarray,
    asof_index: int,
    half_life_days: int,
) -> tuple[float, float, float]:
    if len(origins) == 0:
        return math.inf, math.inf, math.inf
    assert int(np.max(origins)) < asof_index
    error = actual[origins] - predicted[origins]
    if not np.all(np.isfinite(error)):
        return math.inf, math.inf, math.inf
    ages = asof_index - origins
    day_weights = np.power(0.5, ages / max(half_life_days, 1))
    point_weights = np.repeat(day_weights[:, None], PERIODS_PER_DAY, axis=1)
    denominator = float(np.sum(point_weights * np.abs(actual[origins])))
    wape = float(np.sum(point_weights * np.abs(error)) / max(denominator, 1.0))
    rmse = float(np.sqrt(np.sum(point_weights * error**2) / np.sum(point_weights)))
    scale = float(
        np.sum(point_weights * np.abs(actual[origins])) / np.sum(point_weights)
    )
    nrmse = rmse / max(scale, 1.0)
    score = 0.65 * wape + 0.35 * nrmse
    return score, wape, rmse


def validation_origins(
    target_index: int, oos_start_index: int, validation_days: int, stride: int
) -> np.ndarray:
    start = max(oos_start_index, target_index - validation_days)
    origins = list(range(start, target_index, stride))
    if target_index - 1 >= oos_start_index and (not origins or origins[-1] != target_index - 1):
        origins.append(target_index - 1)
    result = np.asarray(sorted(set(origins)), dtype=int)
    if len(result):
        assert int(np.max(result)) < target_index
    return result


def candidate_indices(family: str) -> list[int]:
    return [index for index, config in enumerate(MODEL_CONFIGS) if config.family == family]


def choose_trainable_config(
    component: str,
    family: str,
    series: np.ndarray,
    dates: tuple[date, ...],
    feature_cube: np.ndarray,
    hyper_predictions: np.ndarray,
    tuning_scores: np.ndarray,
    month_start_index: int,
    oos_start_index: int,
    residual_lag: int,
    args: argparse.Namespace,
) -> tuple[str, np.ndarray]:
    origins = validation_origins(
        month_start_index,
        oos_start_index,
        args.validation_days,
        args.tune_stride,
    )
    config_ids = candidate_indices(family)
    if len(origins) < args.minimum_tuning_origins:
        return DEFAULT_CONFIG[(component, family)], origins

    for config_id in config_ids:
        config = MODEL_CONFIGS[config_id]
        for origin in origins:
            if not np.all(np.isfinite(hyper_predictions[config_id, origin])):
                forecast, _ = fit_predict_origin(
                    series,
                    dates,
                    feature_cube,
                    config,
                    int(origin),
                    residual_lag,
                    args.lgbm_threads,
                    include_next_day=False,
                )
                hyper_predictions[config_id, origin] = forecast.astype(np.float32)
        score, _, _ = weighted_error_score(
            series,
            hyper_predictions[config_id],
            origins,
            month_start_index,
            args.tune_half_life,
        )
        tuning_scores[month_start_index, config_id] = score

    best_id = min(config_ids, key=lambda index: tuning_scores[month_start_index, index])
    return MODEL_CONFIGS[best_id].name, origins


def config_by_name(name: str) -> ModelConfig:
    return next(config for config in MODEL_CONFIGS if config.name == name)


def periodic_features(history, target_date, component):
    """Only a prefix is accepted; all 144 target slots are built together."""
    x = feature_row_from_history(history, target_date.date()).astype(float)
    angle = 2*np.pi*(np.arange(144)+1)/144
    harmonics = np.column_stack([v for k in range(1, 5)
                               for v in (np.sin(k*angle), np.cos(k*angle))])
    smooth1 = gaussian_filter1d(history[-1], 2, mode='nearest')
    smooth7 = gaussian_filter1d(history[-7], 2, mode='nearest')
    smoothmean = gaussian_filter1d(history[-7:].mean(0), 2, mode='nearest')
    # Seasonal baseline is explicitly computed from past observations.
    base = (0.65*smooth7+0.35*smooth1 if component == 'load'
            else 0.65*smooth1+0.35*gaussian_filter1d(history[-3:].mean(0), 2))
    interactions = np.column_stack([harmonics * z[:, None]/1000
                                   for z in (smooth1, smooth7, smoothmean)])
    weekday = np.eye(7)[target_date.dayofweek]
    calendar = np.hstack([np.tile(weekday, (144, 1)),
                         np.einsum('tj,k->tjk', harmonics[:, :4], weekday).reshape(144, -1)])
    local = np.column_stack([smooth1, smooth7, smoothmean,
                             np.gradient(smooth1), np.gradient(smooth7),
                             np.maximum(history[-14:].max(0), 1)])
    return np.column_stack([x, harmonics, interactions, calendar, local]).astype('float32'), base


def fit_periodic_predict(x, y, baseline, dates, d, model_id, threads):
    train = np.arange(14, d)
    xx = x[train].reshape(-1, x.shape[-1])
    yy = (y[train] - baseline[train]).reshape(-1)
    half_life = 56 if model_id == 0 else 84
    sw = np.repeat(2.0 ** (-(d-1-train)/half_life), 144)
    if model_id == 0:
        scaler = StandardScaler().fit(xx, sample_weight=sw)
        model = Ridge(alpha=40, solver='cholesky')
        model.fit(scaler.transform(xx), yy, sample_weight=sw)
        pred = model.predict(scaler.transform(x[d]))
        pred = gaussian_filter1d(pred, 1, mode='nearest')
    else:
        model = lgb.LGBMRegressor(objective='regression' if model_id == 1 else 'regression_l1',
                    n_estimators=240, learning_rate=0.035, num_leaves=31,
                    min_child_samples=90, reg_lambda=8, reg_alpha=0.2,
                    max_bin=127, colsample_bytree=0.9, n_jobs=threads,
                    random_state=20260911, verbosity=-1, deterministic=True,
                    force_col_wise=True)
        model.fit(xx, yy, sample_weight=sw)
        pred = model.booster_.predict(x[d], num_threads=threads)
    return np.maximum(0, baseline[d]+pred)


def build_calendar(dates, component):
    a=2*np.pi*(np.asarray(dates.dayofyear)-1)/365.25
    annual=np.column_stack([np.sin(a),np.cos(a),np.sin(2*a),np.cos(2*a)])
    if component=='pv':
        return annual
    week=np.eye(7)[dates.dayofweek]
    return np.column_stack([annual,week,
                           np.einsum('ij,ik->ijk',week,annual[:,:2]).reshape(len(dates),-1)])


def train_component(component, series, date_strings, cache_path, signature, threads=1, end_index=365):
    """从原始功率训练四个分量。每次只拟合目标日前的标签。

    仅复现正式方案所用的模型；未选入的旧季节融合、周期L2等不必运行。
    旧LightGBM/Ridge内部月度选参的候选池必须保留，不能用全年最终配置替代。
    """
    dates = pd.DatetimeIndex(date_strings)
    date_tuple = tuple(x.date() for x in dates)
    args = argparse.Namespace(validation_days=35, tune_stride=3, tune_half_life=21,
                              minimum_tuning_origins=5, lgbm_threads=threads)
    path = Path(cache_path)
    state = dict(predictions=np.full((4,365,144),np.nan),
        hyper=np.full((7,365,144),np.nan,dtype=np.float32),
        scores=np.full((365,7),np.nan), selected=np.full((365,2),'',dtype='U64'),
        calendar_bank=np.full((3,365,144),np.nan), calendar_active=np.full(365,-1,dtype=np.int16),
        done=np.zeros(365,dtype=bool))
    if path.exists():
        saved = read_npz(path)
        if str(saved.pop('signature').item()) != signature:
            raise ValueError('缓存与当前源码、输入或环境不符，请使用新的输出目录。')
        for key in state:
            if key not in saved or saved[key].shape != state[key].shape:
                raise ValueError('缓存字段或形状错误：' + key)
            state[key] = saved[key]
        completed = np.flatnonzero(state['done'])
        if len(completed) and not np.array_equal(completed,np.arange(START,completed[-1]+1)):
            raise ValueError('缓存完成日不连续。')
    print(f'[{component}] 已完成{int(state["done"].sum())}个训练日，准备特征。',flush=True)
    basic = build_feature_cube(series,date_tuple)
    feature_rows, base_rows = zip(*[periodic_features(series[:d],dates[d],component) for d in range(14,365)])
    periodic = np.full((365,144,feature_rows[0].shape[-1]),np.nan,dtype=np.float32)
    baseline = np.full_like(series,np.nan)
    periodic[14:],baseline[14:] = np.asarray(feature_rows),np.asarray(base_rows)
    calendar = build_calendar(dates,component)
    residual_lag = 7 if component == 'load' else 1
    active = [DEFAULT_CONFIG[(component,'lgbm')],DEFAULT_CONFIG[(component,'ridge')]]
    calendar_active = 1
    started = time.time()
    try:
        with threadpool_limits(limits=1, user_api='blas'):
            for d in range(START,end_index):
                if state['done'][d]:
                    active = state['selected'][d].tolist()
                    calendar_active = int(state['calendar_active'][d])
                    continue
                if d == START or dates[d].day == 1:
                    for family_index,family in enumerate(['lgbm','ridge']):
                        active[family_index],_ = choose_trainable_config(component,family,series,date_tuple,basic,
                            state['hyper'],state['scores'],d,START,residual_lag,args)
                    past = np.arange(max(START,d-56),d)
                    if len(past) >= 7:
                        calendar_active = int(np.mean((state['calendar_bank'][:,past]-series[past][None])**2,axis=(1,2)).argmin())
                for family_index,name in enumerate(active):
                    config = config_by_name(name)
                    current,_ = fit_predict_origin(series,date_tuple,basic,config,d,residual_lag,threads,False)
                    rounded = current.astype(np.float32)  # 原方案历史预测库以float32保存这两个分量
                    state['hyper'][MODEL_CONFIGS.index(config),d] = rounded
                    state['predictions'][family_index,d] = rounded
                    state['selected'][d,family_index] = name
                state['predictions'][2,d] = fit_periodic_predict(periodic,series,baseline,dates,d,2,threads)
                weights = 2.**(-(d-1-np.arange(d))/84.)
                for i,alpha in enumerate([.01,.1,1.]):
                    model = Ridge(alpha=alpha,solver='cholesky')
                    model.fit(calendar[:d],series[:d],sample_weight=weights)
                    raw = model.predict(calendar[d:d+1])[0]
                    state['calendar_bank'][i,d] = np.maximum(0,gaussian_filter1d(raw,1,mode='nearest'))
                state['predictions'][3,d] = state['calendar_bank'][calendar_active,d]
                state['calendar_active'][d] = calendar_active
                state['done'][d] = True
                if d % 7 == 0 or dates[d].day == 1 or d == end_index-1:
                    atomic_npz(path,signature=np.asarray(signature),**state)
                    print(f'[{component}] {dates[d].date()}，累计{time.time()-started:.1f}秒，已保存缓存。',flush=True)
    except BaseException:
        atomic_npz(path,signature=np.asarray(signature),**state)
        raise
    atomic_npz(path,signature=np.asarray(signature),**state)
    return str(path)


# 二、单层融合、闭环购电规划、动态规划储能及逐10分钟改约。
def simplex_least_squares(
    x: np.ndarray, y: np.ndarray, regularization: float, prior: np.ndarray
) -> np.ndarray:
    component_count = x.shape[1]
    gram = (x.T @ x) / max(len(y), 1)
    cross = (x.T @ y) / max(len(y), 1)
    prior_error = float(np.mean((x @ prior - y) ** 2))
    penalty = regularization * max(prior_error, 1.0)
    gram = gram + (penalty + 1.0e-8) * np.eye(component_count)
    cross = cross + penalty * prior
    best: np.ndarray | None = None
    best_value = math.inf
    for mask in range(1, 1 << component_count):
        active = np.asarray(
            [index for index in range(component_count) if mask & (1 << index)], dtype=int
        )
        a = gram[np.ix_(active, active)]
        b = cross[active]
        ones = np.ones(len(active))
        kkt = np.block([[2.0 * a, ones[:, None]], [ones[None, :], np.zeros((1, 1))]])
        rhs = np.concatenate([2.0 * b, np.ones(1)])
        try:
            active_weight = np.linalg.solve(kkt, rhs)[:-1]
        except np.linalg.LinAlgError:
            active_weight = np.linalg.lstsq(kkt, rhs, rcond=None)[0][:-1]
        if np.min(active_weight) < -1.0e-8:
            continue
        active_weight = np.clip(active_weight, 0.0, None)
        active_weight /= np.sum(active_weight)
        weight = np.zeros(component_count)
        weight[active] = active_weight
        value = float(
            np.mean((x @ weight - y) ** 2)
            + penalty * np.sum((weight - prior) ** 2)
        )
        if value < best_value:
            best_value = value
            best = weight
    return prior.copy() if best is None else best


def online_blend(actual, cube, dates, windows=(28, 56), groups=False):
    """Monthly choice among causal recipes, each using earlier OOS errors.

    groups=True uses six four-hour groups with shrinkage towards the global
    weights; this is a predeclared ablation, not selected with full-year data.
    """
    n, m, t = cube.shape
    pred = np.full((len(windows), n, t), np.nan)
    ws = np.full((len(windows), n, 6 if groups else 1, m), np.nan)
    prior = np.ones(m)/m
    for ri, window in enumerate(windows):
        for d in range(START, n):
            h = np.arange(max(START, d-window), d)
            w = prior
            if len(h) >= 7:
                xx = cube[h].transpose(0, 2, 1).reshape(-1, m)
                w = simplex_least_squares(xx, actual[h].reshape(-1), 0.15, prior)
            for group in range(6 if groups else 1):
                sl = slice(group*24, (group+1)*24) if groups else slice(0, 144)
                wg = w
                if groups and len(h) >= 14:
                    xx = cube[h, :, sl].transpose(0, 2, 1).reshape(-1, m)
                    local = simplex_least_squares(xx, actual[h, sl].reshape(-1), 0.3, w)
                    wg = 0.5*w+0.5*local
                ws[ri, d, group] = wg
                pred[ri, d, sl] = wg @ cube[d, :, sl]
    out = np.full_like(actual, np.nan)
    records = []
    active, cutoff = 0, -1
    for d in range(START, n):
        if d == START or dates[d].day == 1:
            h = np.arange(max(START, d-56), d)
            if len(h) >= 7:
                loss = np.mean((pred[:, h] - actual[h][None])**2, axis=(1, 2))
                active = int(loss.argmin())
            cutoff = d-1
        out[d] = pred[active, d]
        records.append(dict(date=str(dates[d].date()), window=windows[active],
                            monthly_choice_last_day=cutoff, weight_last_day=d-1,
                            weights=ws[active, d].tolist()))
    return out, records


@lru_cache(maxsize=30)
def matrix(k, t):
    n = 4*t+1+k*t
    ii = np.tile(np.arange(t), k)
    r = np.repeat(np.arange(k*t), 4)
    col = np.column_stack([ii, t+ii, 2*t+ii, 4*t+1+np.arange(k*t)]).ravel()
    a = coo_matrix((np.tile([-1., 1., -1., -1.], k*t), (r, col)),
                   shape=(k*t, n)).tocsr()
    r = np.repeat(np.arange(t), 4)
    col = np.column_stack([t+np.arange(t), 2*t+np.arange(t),
                           3*t+np.arange(t), 3*t+1+np.arange(t)]).ravel()
    eq = coo_matrix((np.tile([-ETA, 1/ETA, -1., 1.], t), (r, col)),
                    shape=(t, n)).tocsr()
    bounds = [(0., None)]*t+[(0., P_MAX)]*(2*t)+[(S_MIN, S_MAX)]*(t+1)+[(0., None)]*(k*t)
    return a, eq, bounds


def solve(prices, scenarios, weights, initial, terminal_value, terminal=None):
    scenarios = np.asarray(scenarios)
    k, t = scenarios.shape
    a, eq, original_bounds = matrix(k, t)
    bounds = original_bounds.copy()
    bounds[3*t] = (initial, initial)
    if terminal is not None:
        bounds[4*t] = (terminal, terminal)
    c = np.zeros(4*t+1+k*t)
    c[:t] = prices
    c[t:3*t] = 1e-7  # numerical tie-break, excluded from reported cash costs
    c[4*t] = -terminal_value
    c[4*t+1:] = (5*weights[:, None]*prices[None]).ravel()
    result = linprog(c, A_ub=a, b_ub=-scenarios.ravel(), A_eq=eq,
                     b_eq=np.zeros(t), bounds=bounds, method='highs')
    if not result.success:
        raise RuntimeError(result.message)
    g, charge, discharge = [np.maximum(result.x[i*t:(i+1)*t], 0) for i in range(3)]
    soc = result.x[3*t:4*t+1]
    assert np.max(np.abs(np.diff(soc)-ETA*charge+discharge/ETA)) < 1e-5
    assert np.max(np.minimum(charge, discharge)) < 1e-5
    return g, charge, discharge, soc


def scenario_set(history_net, history_pred, current_pred, window=56, half_life=14, limit=28):
    """API accepts only completed history; cannot index target-day truth."""
    d = len(history_net)
    assert history_pred.shape == history_net.shape
    indices = np.arange(max(START, d-window), d)
    if len(indices) == 0:
        return current_pred[None], np.ones(1), -1
    assert np.isfinite(history_pred[indices]).all()
    if len(indices) > limit:
        indices = indices[np.rint(np.linspace(0, len(indices)-1, limit)).astype(int)]
    weights = 2.**(-(d-indices)/half_life)
    weights /= weights.sum()
    return current_pred[None]+history_net[indices]-history_pred[indices], weights, int(indices[-1])


def weighted_quantile(values, weights, q):
    order = np.argsort(values, axis=0, kind='stable')
    sorted_values = np.take_along_axis(values, order, axis=0)
    cw = np.cumsum(np.broadcast_to(weights[:, None], values.shape)[order, np.arange(values.shape[1])], axis=0)
    index = np.argmax(cw >= q-1e-12, axis=0)
    return sorted_values[index, np.arange(values.shape[1])]


def step_action(net_now, planned_now, state, reserve):
    deficit = net_now-planned_now
    if deficit > 0:
        discharge = min(deficit, P_MAX, max(0., ETA*(state-reserve)))
        charge, emergency, spill = 0., deficit-discharge, 0.
    else:
        charge = min(-deficit, P_MAX, max(0., (S_MAX-state)/ETA))
        discharge, emergency, spill = 0., 0., -deficit-charge
    return charge, discharge, emergency, spill, state+ETA*charge-discharge/ETA


def execute(net, grid, reserves, initial):
    flows = np.zeros((5, 144)); flows[0] = grid
    soc = np.zeros(145); soc[0] = initial
    for t in range(144):
        c, b, u, w, soc[t+1] = step_action(float(net[t]), grid[t], soc[t], reserves[t])
        flows[1:, t] = c, b, u, w
    assert np.max(np.abs(flows[0]+flows[2]+flows[3]-net-flows[1]-flows[4])) < 1e-6
    assert np.max(np.abs(np.diff(soc)-ETA*flows[1]+flows[2]/ETA)) < 1e-6
    assert soc.min() >= S_MIN-1e-6 and soc.max() <= S_MAX+1e-6
    assert np.max(flows[1:3]) <= P_MAX+1e-6
    assert np.min(flows) >= -1e-6
    assert np.max(np.minimum(flows[1], flows[2])) < 1e-8
    return flows, soc


def closed_loop_loss_kernel(offsets, base_grid, scale, reserves, scenarios, weights, initial, prices, value):
    """Equivalent scalar-loop kernel, optionally compiled with Numba."""
    grid = np.zeros(144); dg = np.zeros(144)
    gradient = np.zeros(6); cost = 0.
    for t in range(144):
        raw = base_grid[t]+offsets[t//24]*scale[t]
        grid[t] = max(0., raw)
        dg[t] = scale[t] if raw > 0 else 0.
        cost += prices[t]*grid[t]
        gradient[t//24] += prices[t]*dg[t]
    for i in range(len(scenarios)):
        state = float(initial)
        js = np.zeros(6)
        for t in range(144):
            b = t//24
            h = scenarios[i,t]-grid[t]
            if h <= 0:
                room = (S_MAX-state)/ETA
                if -h <= P_MAX and -h <= room:
                    charge = -h
                    js[b] += ETA*dg[t]
                elif P_MAX <= room:
                    charge = P_MAX
                else:
                    charge = max(0., room)
                    js[:] = 0.
                state += ETA*charge
            else:
                available = max(0., ETA*(state-reserves[t]))
                if h <= P_MAX and h <= available:
                    discharge = h
                    emergency = 0.
                    js[b] += dg[t]/ETA
                elif P_MAX <= available:
                    discharge = P_MAX
                    emergency = h-P_MAX
                    gradient[b] -= weights[i]*5*prices[t]*dg[t]
                else:
                    discharge = available
                    emergency = h-available
                    gradient[b] -= weights[i]*5*prices[t]*dg[t]
                    if state > reserves[t]:
                        for j in range(6):
                            gradient[j] -= weights[i]*5*prices[t]*ETA*js[j]
                            js[j] = 0.
                state -= discharge/ETA
                cost += weights[i]*5*prices[t]*emergency
        cost -= weights[i]*value*state
        for j in range(6):
            gradient[j] -= weights[i]*value*js[j]
    return cost, gradient


def optimize_grid(base_grid, scale, reserves, scenarios, weights, initial, prices, value):
    objective = lambda z: closed_loop_loss(z, base_grid, scale, reserves, scenarios,
                                           weights, initial, prices, value)
    zero = np.zeros(6)
    before = objective(zero)[0]
    fit = minimize(objective, zero, jac=True, method='L-BFGS-B', bounds=[(-2., 2.)]*6,
                   options=dict(maxiter=60, ftol=1e-9, gtol=1e-4, maxls=25))
    accepted = fit.x if np.isfinite(fit.fun) and fit.fun < before else zero
    grid = np.maximum(0, base_grid+accepted[BLOCK]*scale)
    return grid, dict(objective_before=before, objective_after=float(objective(accepted)[0]),
                     optimizer_success=bool(fit.success), iterations=int(fit.nit), offsets=accepted.tolist())

closed_loop_loss=njit(cache=True)(closed_loop_loss_kernel)

def dp_reserves(grid,scenarios,weights,prices,value,step=50.):
    states=np.linspace(S_MIN,S_MAX,round((S_MAX-S_MIN)/step)+1)
    continuation=-value*states
    reserves=np.zeros(len(grid))
    for t in range(len(grid)-1,-1,-1):
        reserve=states[np.argmin(continuation+5*prices[t]*ETA*states)]
        reserves[t]=reserve
        s=states[None,:];h=(scenarios[:,t]-grid[t])[:,None]
        minimum=np.maximum(S_MIN,s-np.minimum(np.maximum(h,0),P_MAX)/ETA)
        deficit_end=np.maximum(minimum,np.minimum(s,reserve))
        surplus_end=np.minimum(S_MAX,s+ETA*np.minimum(np.maximum(-h,0),P_MAX))
        end=np.where(h>0,deficit_end,surplus_end)
        emergency=np.where(h>0,np.maximum(h-ETA*(s-end),0),0)
        future=np.interp(end.ravel(),states,continuation).reshape(end.shape)
        continuation=weights@(5*prices[t]*emergency+future)
    return reserves


@lru_cache(maxsize=8)
def lp_matrices(t):
    # [A,C,D,S(0..T),U,up,down], 7T+1 variables.
    k=np.arange(t);n=7*t+1
    a=coo_matrix((np.tile([-1.,1.,-1.,-1.],t),
        (np.repeat(k,4),np.column_stack([k,t+k,2*t+k,4*t+1+k]).ravel())),shape=(t,n)).tocsr()
    rows=np.r_[np.repeat(k,4),np.repeat(t+k,3)]
    cols=np.r_[np.column_stack([t+k,2*t+k,3*t+k,3*t+1+k]).ravel(),
        np.column_stack([k,5*t+1+k,6*t+1+k]).ravel()]
    vals=np.r_[np.tile([-ETA,1/ETA,-1.,1.],t),np.tile([1.,-1.,1.],t)]
    eq=coo_matrix((vals,(rows,cols)),shape=(2*t,n)).tocsr()
    bounds=[(0,None)]*t+[(0,P_MAX)]*(2*t)+[(S_MIN,S_MAX)]*(t+1)+[(0,None)]*(3*t)
    return a,eq,bounds


def initial_adjustment_plan(quantile,original,prices,initial,value):
    t=len(prices);a,eq,bounds0=lp_matrices(t);bounds=bounds0.copy()
    bounds[3*t]=(initial,initial)
    c=np.zeros(7*t+1);c[t:3*t]=1e-7;c[4*t]=-value
    c[4*t+1:5*t+1]=5*prices
    c[5*t+1:6*t+1]=1.5*prices;c[6*t+1:]=-.5*prices
    result=linprog(c,A_ub=a,b_ub=-quantile,A_eq=eq,b_eq=np.r_[np.zeros(t),original],bounds=bounds,method='highs')
    if not result.success:raise RuntimeError(result.message)
    assert np.minimum(result.x[5*t+1:6*t+1],result.x[6*t+1:]).max()<1e-6
    assert np.minimum(result.x[t:2*t],result.x[2*t:3*t]).max()<1e-5
    return np.maximum(0,result.x[:t])

@njit(cache=True)
def flexible_loss(offsets,base_grid,scale,reserves,scenarios,weights,initial,prices,value,original,blocks,adjustment):
    t_count=len(prices);n=len(offsets)
    grid=np.zeros(t_count);dg=np.zeros(t_count);gradient=np.zeros(n);cost=0.
    for t in range(t_count):
        b=blocks[t];raw=base_grid[t]+offsets[b]*scale[t]
        grid[t]=max(0.,raw);dg[t]=scale[t] if raw>0 else 0.
        if adjustment:
            delta=grid[t]-original[t]
            cost+=prices[t]*original[t]+prices[t]*(1.5*max(delta,0.)-.5*max(-delta,0.))
            slope=1.5 if delta>0 else (.5 if delta<0 else 1.)
        else:
            cost+=prices[t]*grid[t]
            slope=1.
        gradient[b]+=slope*prices[t]*dg[t]
    for i in range(len(scenarios)):
        state=float(initial);jac_state=np.zeros(n)
        for t in range(t_count):
            b=blocks[t];h=scenarios[i,t]-grid[t]
            if h<=0:
                room=(S_MAX-state)/ETA
                if -h<=P_MAX and -h<=room:
                    charge=-h;jac_state[b]+=ETA*dg[t]
                elif P_MAX<=room:charge=P_MAX
                else:
                    charge=max(0.,room);jac_state[:]=0.
                state+=ETA*charge
            else:
                available=max(0.,ETA*(state-reserves[t]))
                if h<=P_MAX and h<=available:
                    discharge=h;emergency=0.;jac_state[b]+=dg[t]/ETA
                elif P_MAX<=available:
                    discharge=P_MAX;emergency=h-P_MAX
                    gradient[b]-=weights[i]*5*prices[t]*dg[t]
                else:
                    discharge=available;emergency=h-available
                    gradient[b]-=weights[i]*5*prices[t]*dg[t]
                    if state>reserves[t]:
                        for j in range(n):
                            gradient[j]-=weights[i]*5*prices[t]*ETA*jac_state[j]
                            jac_state[j]=0.
                state-=discharge/ETA
                cost+=weights[i]*5*prices[t]*emergency
        cost-=weights[i]*value*state
        gradient-=weights[i]*value*jac_state
    return cost,gradient


def curve_refine(grid,reserve,scenarios,weights,prices,initial,value,original,begin,step,iterations=100):
    """按固定历史情景优化曲线；original=None代表0点正常计划购电。"""
    mean=weights@scenarios
    scale=np.maximum(20.,np.sqrt(weights@((scenarios-mean)**2)))
    blocks=(np.arange(begin,begin+len(prices))//step-begin//step).astype(np.int64)
    zero=np.zeros(int(blocks.max())+1)
    adjustment=original is not None
    origin=original if adjustment else np.zeros(len(prices))
    args=(scenarios,weights,initial,prices,value,origin,blocks,adjustment)
    objective=lambda z:flexible_loss(z,grid,scale,reserve,*args)
    before=float(objective(zero)[0])
    fit=minimize(objective,zero,jac=True,method='L-BFGS-B',bounds=[(-2.,2.)]*len(zero),
        options=dict(maxiter=iterations,ftol=1e-9,gtol=1e-4,maxls=30 if iterations==100 else 25))
    z=fit.x if np.isfinite(fit.fun) and fit.fun<before else zero
    changed=np.maximum(0,grid+z[blocks]*scale)
    updated=dp_reserves(changed,scenarios,weights,prices,value)
    old=float(flexible_loss(zero,changed,scale,reserve,*args)[0])
    if flexible_loss(zero,changed,scale,updated,*args)[0]<old:reserve=updated
    after=float(flexible_loss(zero,changed,scale,reserve,*args)[0])
    assert after<=before+1e-6
    return changed,reserve,dict(before=before,after=after,dimensions=len(zero),iterations=int(fit.nit),success=bool(fit.success))


def initial_plan(scenarios,weights,prices,initial,value,slot_refine=False):
    """原Q2计划器的相同计算顺序，终端库存价值由调用方显式提供。"""
    q=weighted_quantile(scenarios,weights,.8)
    grid,_,_,_=solve(prices,q[None],np.ones(1),initial,value)
    mean=weights@scenarios;scale=np.maximum(20.,np.sqrt(weights@((scenarios-mean)**2)))
    reserve=dp_reserves(grid,scenarios,weights,prices,value)
    grid,detail=optimize_grid(grid,scale,reserve,scenarios,weights,initial,prices,value)
    updated=dp_reserves(grid,scenarios,weights,prices,value)
    args=(grid,scale,scenarios,weights,initial,prices,value)
    old=closed_loop_loss(np.zeros(6),grid,scale,reserve,scenarios,weights,initial,prices,value)[0]
    new=closed_loop_loss(np.zeros(6),grid,scale,updated,scenarios,weights,initial,prices,value)[0]
    if new<old:reserve=updated
    if slot_refine:
        grid,reserve,extra=curve_refine(grid,reserve,scenarios,weights,prices,initial,value,None,0,1)
        detail['slot_refine']=extra
    return grid,reserve,detail


def adjustment_plan(scenarios,weights,prices,initial,original,begin,value,slot_refine=False):
    grid=initial_adjustment_plan(weighted_quantile(scenarios,weights,.8),original,prices,initial,value)
    reserve=dp_reserves(grid,scenarios,weights,prices,value)
    grid,reserve,detail=curve_refine(grid,reserve,scenarios,weights,prices,initial,value,original,begin,24,60)
    if slot_refine:
        grid,reserve,extra=curve_refine(grid,reserve,scenarios,weights,prices,initial,value,original,begin,1)
        detail['slot_refine']=extra
    return grid,reserve,detail

# 三、光伏发布预报、日内负载校正和条件残差情景。
def interpolate_release(day,release,raw,actual_pv):
    begin=release*36
    anchor=actual_pv[day,begin-1] if begin else (actual_pv[day-1,-1] if day else 0.)
    hours=release*6+np.arange(25)
    values=np.r_[anchor,raw[day,release]]
    endpoints=(np.arange(begin,144)+1)/6
    return np.maximum(0,np.interp(endpoints,hours,values))


def calibration_weight(day,release,actual_pv,base_pv,external):
    begin=release*36
    ids=np.arange(max(START,day-56),day)
    if len(ids)<7:return .5,len(ids),int(ids[-1]) if len(ids) else -1
    delta=external[ids,release,begin:]-base_pv[ids,begin:]
    error=actual_pv[ids,begin:]-base_pv[ids,begin:]
    weights=2.**(-(day-ids)/14)
    denominator=float(np.sum(weights[:,None]*delta*delta))
    alpha=float(np.clip(np.sum(weights[:,None]*delta*error)/denominator,0.,1.)) if denominator>1e-8 else 0.
    return alpha,len(ids),int(ids[-1])


def observed_features(day,release,load,base_load,pv):
    begin=36*release
    if not begin:return np.zeros(3)
    error=load[day,:begin]-base_load[day,:begin]
    revision=np.mean(pv[day,release,begin:begin+36]-pv[day,0,begin:begin+36])
    return np.array([error[-6:].mean(),error.mean(),revision])


def predict_load(day,release,actual,base):
    begin=36*release;result=base[day].copy();records=[]
    if not begin:return result,records
    ids=np.arange(max(15,day-56),day)
    if len(ids)<7:return result,records
    errors=actual[:day]-base[:day]
    x=np.column_stack([errors[ids,begin-6:begin].mean(axis=1),errors[ids,:begin].mean(axis=1)])
    now=np.array([(actual[day,begin-6:begin]-base[day,begin-6:begin]).mean(),
        (actual[day,:begin]-base[day,:begin]).mean()])
    weights=2.**(-(day-ids)/14)
    mean=np.average(x,axis=0,weights=weights)
    scale=np.maximum(10.,np.sqrt(np.average((x-mean)**2,axis=0,weights=weights)))
    design=np.column_stack([np.ones(len(ids)),(x-mean)/scale])
    vector=np.r_[1.,(now-mean)/scale]
    lhs=design.T@(weights[:,None]*design)+10*np.eye(3)
    for block in range(release,4):
        target=errors[ids,36*block:36*(block+1)].mean(axis=1)
        coef=np.linalg.solve(lhs,design.T@(weights*target))
        correction=float(vector@coef)
        result[36*block:36*(block+1)]=np.maximum(0,result[36*block:36*(block+1)]+correction)
        records.append(dict(day=day,release=release,block=block,history_last=int(ids[-1]),
            history_count=len(ids),correction_kw=correction,coefficients=coef.tolist(),mean=mean.tolist(),scale=scale.tolist()))
    return result,records


def historical_ids(day):
    ids=np.arange(max(15,day-56),day)
    if len(ids)>28:ids=ids[np.rint(np.linspace(0,len(ids)-1,28)).astype(int)]
    return ids


def scenarios(day,release,pred,actual,features,conditional=False):
    begin=36*release;ids=historical_ids(day)
    sc=pred[day,release,begin:][None]+actual[ids,begin:]-pred[ids,release,begin:]
    weights=2.**(-(day-ids)/14);weights/=weights.sum()
    if conditional and release:
        x=features[ids,release];now=features[day,release]
        scale=np.maximum(20.,np.sqrt(np.average((x-np.average(x,axis=0,weights=weights))**2,axis=0,weights=weights)))
        distance=np.mean(((x-now)/scale)**2,axis=1)
        near=weights*np.exp(-.5*np.minimum(distance,50));near/=near.sum()
        weights=.25*weights+.75*near
    return sc,weights,ids

def prepare_supply(data, output):
    output.mkdir(parents=True,exist_ok=True)
    predictions=[];fusion_logs=[]
    for name,actual,cube in zip(['load','pv'],[data['load'],data['pv']],data['components']):
        forecast,logs=online_blend(actual,cube.transpose(1,0,2),DATES)
        predictions.append(forecast)
        for row in logs:row['component']=name
        fusion_logs.extend(logs)
    base_load,base_pv=predictions
    external=np.full((365,4,144),np.nan);pv=np.full_like(external,np.nan);pvlogs=[]
    for day in range(365):
        for release in range(4):
            begin=36*release
            external[day,release,begin:]=interpolate_release(day,release,data['raw'],data['pv'])
    for day in range(START,365):
        for release in range(4):
            begin=36*release
            alpha,count,last=calibration_weight(day,release,data['pv'],base_pv,external)
            pv[day,release,begin:]=(1-alpha)*base_pv[day,begin:]+alpha*external[day,release,begin:]
            pvlogs.append(dict(day=day,release=release,weight=alpha,history_count=count,history_last=last,latest_observed_slot=begin-1))
    base=(base_load[:,None,:]-pv)/6;corrected=base.copy()
    features=np.zeros((365,4,3));loadbank=np.repeat(base_load[:,None,:],4,axis=1);loadlogs=[]
    for day in range(START,365):
        for release in range(4):
            features[day,release]=observed_features(day,release,data['load'],base_load,pv)
            loadbank[day,release],logs=predict_load(day,release,data['load'],base_load)
            corrected[day,release]=(loadbank[day,release]-pv[day,release])/6
            loadlogs.extend(logs)
    np.savez_compressed(output/'supply_forecasts.npz',load_prediction=base_load,pv_prediction=base_pv,
        q2_net=(base_load-base_pv)*(1/6),q3_base=base,q3_corrected=corrected,q3_features=features,
        q3_pv=pv,q3_load=loadbank)
    write_json(output/'supply_fusion_log.json',fusion_logs);write_json(output/'pv_release_log.json',pvlogs)
    write_json(output/'load_correction_log.json',loadlogs);write_json(output/'supply_provenance.json',source_hashes())
    print('原单层融合、已发布PV预报和日内负载校正已重新计算。',flush=True)


# 四、价格候选、历史单层融合及日内价格校正。
def predict_day(history,target_index,method):
    """history只含origin之前价格；target_index可为origin或origin+1。"""
    origin=len(history)
    if not origin:return np.ones(144)
    ids=np.arange(max(0,origin-56),origin)
    same=ids[ids%7==target_index%7]
    if not len(same):return history.mean(axis=0)
    if method=='weekly_persistence':return history[same[-1]].copy()
    w=2.**(-(origin-same)/28);w/=w.sum()
    average=w@history[same]
    if method=='weekday_mean' or origin<14:return np.maximum(.001,average)
    # 均价=星期效应+局部趋势；日内形状采用历史同星期的去均值曲线。
    weekday=np.eye(7)[ids%7][:,1:]
    x=np.column_stack([np.ones(len(ids)),weekday,(ids-origin)/28])
    h=2.**(-(origin-ids)/28)
    penalty=np.diag([1e-8]+[.1]*6+[.01])
    coef=np.linalg.solve(x.T@(h[:,None]*x)+penalty,x.T@(h*history[ids].mean(axis=1)))
    vector=np.r_[1.,np.eye(7)[target_index%7,1:],(target_index-origin)/28]
    mean=float(vector@coef)
    shape=w@(history[same]-history[same].mean(axis=1)[:,None])
    return np.maximum(.001,mean+shape)


def intraday_correct(day,release,history,base,current_prefix):
    """history=实际价格[:day]，当前日仅传入0..begin-1；训练标签来自之前日。"""
    begin=release*36
    result=base[day].copy()
    if not begin:return result,[]
    assert current_prefix.shape==(begin,)
    ids=np.arange(max(7,day-56),day)
    if len(ids)<7:return result,[]
    errors=history[ids]-base[ids]
    x=np.column_stack([errors[:,begin-6:begin].mean(axis=1),errors[:,:begin].mean(axis=1)])
    now_error=current_prefix-base[day,:begin]
    now=np.array([now_error[-6:].mean(),now_error.mean()])
    w=2.**(-(day-ids)/14)
    mean=np.average(x,axis=0,weights=w)
    scale=np.maximum(.005,np.sqrt(np.average((x-mean)**2,axis=0,weights=w)))
    design=np.column_stack([np.ones(len(ids)),(x-mean)/scale]);vector=np.r_[1.,(now-mean)/scale]
    lhs=design.T@(w[:,None]*design)+10*np.eye(3)
    log=[]
    for block in range(release,4):
        sl=slice(block*36,(block+1)*36)
        coef=np.linalg.solve(lhs,design.T@(w*errors[:,sl].mean(axis=1)))
        delta=float(vector@coef)
        result[sl]=np.maximum(.001,result[sl]+delta)
        log.append(dict(day=day,release=release,block=block,training_last_day=day-1,
            latest_observed_slot=begin-1,correction=delta,coefficients=coef.tolist()))
    return result,log

METHODS=["weekly_persistence","weekday_mean","weekday_trend"]

def prepare_prices(data, output):
    output.mkdir(parents=True,exist_ok=True);price=data["price"]
    base=np.full((len(METHODS),365,144),np.nan);nextday=np.full_like(base,np.nan)
    for day in range(365):
        for m,method in enumerate(METHODS):
            base[m,day]=predict_day(price[:day],day,method)
            nextday[m,day]=predict_day(price[:day],day+1,method)
    selected=np.zeros(365,dtype=int);chosen=np.empty((365,144));chosen_next=np.empty_like(chosen);logs=[]
    active=0;choice_day=0
    for day in range(365):
        if day==15 or DATES[day].day==1:
            ids=np.arange(max(7,day-56),day)
            loss=np.mean((base[:,ids]-price[ids][None])**2,axis=(1,2)) if len(ids)>=7 else None
            if loss is not None:active=int(loss.argmin())
            choice_day=day
            logs.append(dict(day=day,date=str(DATES[day].date()),selected=METHODS[active],
                validation_last_day=day-1,validation_count=len(ids),loss=None if loss is None else loss.tolist()))
        selected[day]=active;chosen[day]=base[active,day];chosen_next[day]=nextday[active,day]
    updated=np.repeat(chosen[:,None,:],4,axis=1);corrections=[]
    for day in range(15,365):
        for release in [1,2,3]:
            updated[day,release],rows=intraday_correct(day,release,price[:day],chosen[:day+1],price[day,:36*release])
            corrections.extend(rows)
    np.savez_compressed(output/'price_forecasts.npz',methods=np.array(METHODS),base_candidates=base,
        next_candidates=nextday,selected=selected,midnight=chosen,updated=updated,next_day=chosen_next)
    write_json(output/'price_selection.json',logs);write_json(output/'price_corrections.json',corrections)
    records=[]
    for start,stop,period in [(15,31,'January'),(31,365,'February-December')]:
        for name,forecast in [*zip(METHODS,base),('monthly_selected',chosen)]:
            e=forecast[start:stop]-price[start:stop]
            records.append(dict(period=period,method=name,release_hour=0,rmse=float(np.sqrt(np.mean(e**2))),mae=float(np.abs(e).mean())))
        for release in [1,2,3]:
            sl=slice(36*release,36*(release+1))
            for name,forecast in [('midnight_unchanged',chosen),('intraday_corrected',updated[:,release])]:
                e=forecast[start:stop,sl]-price[start:stop,sl]
                records.append(dict(period=period,method=name,release_hour=release*6,rmse=float(np.sqrt(np.mean(e**2))),mae=float(np.abs(e).mean())))
    table=pd.DataFrame(records);table.to_csv(output/'price_metrics.csv',index=False,encoding='utf-8-sig')
    write_json(output/'price_provenance.json',source_hashes())
    print("三条价格候选和历史月度选择已重算。",flush=True)

STATES=np.linspace(S_MIN,S_MAX,193)

@njit(cache=True)
def night_cost(prices, net, states):
    """短期完全预测近似只用因果预测曲线，不输入次日实际数据。"""
    future = np.zeros(len(states))
    for t in range(len(prices) - 1, -1, -1):
        current = np.empty_like(future)
        for i in range(len(states)):
            best = np.inf
            for j in range(len(states)):
                delta = states[j] - states[i]
                charge = max(delta, 0.) / ETA
                discharge = max(-delta, 0.) * ETA
                if charge > P_MAX + 1e-9 or discharge > P_MAX + 1e-9:
                    continue
                grid = max(0., net[t] + charge - discharge)
                best = min(best, prices[t] * grid + future[j])
            current[i] = best
        future = current
    return future - future[0]


def reserves_with_curve(grid, scenarios, weights, prices, terminal):
    future = terminal.copy()
    reserve = np.zeros(len(grid))
    for t in range(len(grid) - 1, -1, -1):
        rule = STATES[np.argmin(future + 5 * prices[t] * ETA * STATES)]
        reserve[t] = rule
        state = STATES[None, :]
        error = (scenarios[:, t] - grid[t])[:, None]
        minimum = np.maximum(S_MIN, state - np.minimum(np.maximum(error, 0), P_MAX) / ETA)
        deficit_end = np.maximum(minimum, np.minimum(state, rule))
        surplus_end = np.minimum(S_MAX, state + ETA * np.minimum(np.maximum(-error, 0), P_MAX))
        end = np.where(error > 0, deficit_end, surplus_end)
        emergency = np.where(error > 0, np.maximum(error - ETA * (state - end), 0), 0)
        continuation = np.interp(end.ravel(), STATES, future).reshape(end.shape)
        future = weights @ (5 * prices[t] * emergency + continuation)
    return reserve

def price_weight(day, history_price, candidates):
    assert len(history_price) == day
    ids = np.arange(max(7, day - 56), day)
    if len(ids) < 7:
        return np.ones(3) / 3
    x = candidates[:, ids].transpose(1, 2, 0).reshape(-1, 3) * 100
    y = history_price[ids].reshape(-1) * 100
    return simplex_least_squares(x, y, .15, np.ones(3) / 3)


def next_net(history_load, history_pv, target):
    assert len(history_load) == len(history_pv)
    origin = len(history_load)
    ids = np.arange(max(0, origin - 56), origin)
    ids = ids[ids % 7 == target % 7]
    if not len(ids):
        return np.zeros(144) if not origin else (history_load.mean(axis=0) - history_pv.mean(axis=0)) / 6
    weights = 2. ** (-(origin - ids) / 28)
    weights /= weights.sum()
    return (weights @ history_load[ids] - weights @ history_pv[ids]) / 6

def prepare_price_fusion(data, output, original):
    output.mkdir(parents=True,exist_ok=True)
    midnight = np.zeros((365, 144))
    next_day = np.zeros_like(midnight)
    weights = np.zeros((365, 3))
    future_net = np.zeros_like(midnight)
    logs = []
    for d in range(365):
        w = price_weight(d, data['price'][:d], original['base_candidates'])
        weights[d] = w
        midnight[d] = w @ original['base_candidates'][:, d]
        next_day[d] = w @ original['next_candidates'][:, d]
        future_net[d] = next_net(data['load'][:d], data['pv'][:d], d + 1)
        logs.append(dict(day=d, weight_last_day=d - 1, next_load_last_day=d - 1, next_pv_last_day=d - 1,
                         next_target_index=d + 1, weights=w.tolist()))
    updated = np.repeat(midnight[:, None, :], 4, axis=1)
    correction_logs = []
    for d in range(15, 365):
        for release in [1, 2, 3]:
            updated[d, release], rows = intraday_correct(d, release, data['price'][:d], midnight[:d + 1],
                                                       data['price'][d, :release * 36])
            correction_logs.extend(rows)
    terminal_old = np.zeros((365, len(STATES)))
    terminal_blend = np.zeros_like(terminal_old)
    for d in range(15, 365):
        terminal_old[d] = night_cost(original['next_day'][d, :36], future_net[d, :36], STATES)
        terminal_blend[d] = night_cost(next_day[d, :36], future_net[d, :36], STATES)
    np.savez_compressed(output / 'new_forecasts.npz', midnight=midnight, updated=updated, next_day=next_day,
                        weights=weights, next_net=future_net, terminal_original=terminal_old,
                        terminal_blend=terminal_blend)
    write_json(output / 'forecast_history.json', logs)
    write_json(output / 'price_correction_history.json', correction_logs)
    write_json(output / 'forecast_provenance.json', source_hashes())
    metrics = []
    for name, bank in [('original', original), ('price_fusion', dict(midnight=midnight, updated=updated))]:
        for r in range(4):
            forecast = bank['midnight'] if not r else bank['updated'][:, r]
            stop = 144 if not r else (r + 1) * 36
            diff = forecast[31:, r * 36:stop] - data['price'][31:, r * 36:stop]
            metrics.append(dict(model=name, release_hour=r * 6, rmse=float(np.sqrt(np.mean(diff ** 2))),
                                mae=float(np.mean(abs(diff)))))
    write_json(output / 'price_metrics.json', metrics)
    print('电价融合与次日凌晨库存成本曲线已按历史截止生成。', flush=True)


# 五、第四问配对价格情景、额外曲线优化与历史预期改约。
@njit(cache=True)
def loss(z,base_grid,scale,base_reserve,net,weights,initial,prices,terminal,original,
         adjustment,blocks,change_grid,change_reserve):
    ntime=len(base_grid);n=len(z)
    grid=base_grid.copy();reserves=base_reserve.copy()
    dg=np.zeros(ntime);dr=np.zeros(ntime);grad=np.zeros(n)
    ngrid=ntime if change_grid else 0
    for t in range(ntime):
        if change_grid:
            raw=base_grid[t]+z[t]*scale[t]
            grid[t]=max(0.,raw);dg[t]=scale[t] if raw>0 else 0.
        if change_reserve:
            raw=base_reserve[t]+1000*z[ngrid+blocks[t]]
            reserves[t]=min(S_MAX,max(S_MIN,raw))
            dr[t]=1000. if S_MIN<raw<S_MAX else 0.
    value=0.
    for i in range(len(net)):
        state=initial;js=np.zeros(n);cost=0.
        for t in range(ntime):
            p=prices[i,t];weight=weights[i]
            if adjustment:
                delta=grid[t]-original[t]
                cost+=p*(original[t]+1.5*max(delta,0.)-.5*max(-delta,0.))
                slope=1.5 if delta>0 else (.5 if delta<0 else 1.)
            else:
                cost+=p*grid[t];slope=1.
            if change_grid:grad[t]+=weight*slope*p*dg[t]
            h=net[i,t]-grid[t]
            if h<=0:
                room=(S_MAX-state)/ETA
                if -h<=P_MAX and -h<=room:
                    charge=-h
                    if change_grid:js[t]+=ETA*dg[t]
                elif P_MAX<=room:charge=P_MAX
                else:
                    charge=max(0.,room);js[:]=0.
                state+=ETA*charge
            else:
                available=max(0.,ETA*(state-reserves[t]))
                if h<=P_MAX and h<=available:
                    discharge=h;emergency=0.
                    if change_grid:js[t]+=dg[t]/ETA
                elif P_MAX<=available:
                    discharge=P_MAX;emergency=h-P_MAX
                    if change_grid:grad[t]-=weight*5*p*dg[t]
                else:
                    discharge=available;emergency=h-available
                    if change_grid:grad[t]-=weight*5*p*dg[t]
                    if state>reserves[t]:
                        grad-=weight*5*p*ETA*js;js[:]=0.
                        if change_reserve:
                            b=ngrid+blocks[t]
                            grad[b]+=weight*5*p*ETA*dr[t];js[b]=dr[t]
                state-=discharge/ETA;cost+=5*p*emergency
        index=min(len(terminal)-2,max(0,int((state-S_MIN)/50)))
        slope=(terminal[index+1]-terminal[index])/50
        cost+=terminal[index]+slope*(state-S_MIN-index*50)
        grad+=weights[i]*slope*js
        value+=weights[i]*cost
    return value,grad


def optimize(grid,reserve,net,weights,initial,prices,terminal,original,begin,
             change_grid=True,change_reserve=False,rounds=1):
    mean=weights@net;scale=np.maximum(20.,np.sqrt(weights@((net-mean)**2)))
    blocks=(np.arange(begin,begin+len(grid))//24-begin//24).astype(np.int64)
    ngrid=len(grid) if change_grid else 0
    nreserve=int(blocks.max())+1 if change_reserve else 0
    n=ngrid+nreserve
    if not n:return grid,reserve,[]
    bounds=[(-2.,2.)]*ngrid+[(-1.5,1.5)]*nreserve
    zero=np.zeros(n);records=[]
    origin=np.zeros(len(grid)) if original is None else original
    for turn in range(rounds):
        args=(grid,scale,reserve,net,weights,initial,prices,terminal,origin,
              original is not None,blocks,change_grid,change_reserve)
        objective=lambda z:loss(z,*args)
        before=float(objective(zero)[0])
        fitted=minimize(objective,zero,jac=True,method='L-BFGS-B',bounds=bounds,
            options=dict(maxiter=100,ftol=1e-9,gtol=1e-4,maxls=30))
        z=fitted.x if np.isfinite(fitted.fun) and fitted.fun<before else zero
        after=float(objective(z)[0]);assert after<=before+1e-6
        if change_grid:grid=np.maximum(0,grid+z[:ngrid]*scale)
        if change_reserve:reserve=np.clip(reserve+1000*z[ngrid+blocks],S_MIN,S_MAX)
        records.append(dict(before=before,after=after,iterations=int(fitted.nit),
                            success=bool(fitted.success),grid_dimensions=ngrid,reserve_dimensions=nreserve))
    return grid,reserve,records


def paired_prices(point,weights,historical_actual,historical_forecast,fraction):
    error=historical_actual-historical_forecast
    error-=weights@error
    negative=np.maximum(0,-error.min(axis=0))
    contraction=np.minimum(1.,.95*point/np.maximum(negative,1e-12))
    result=point[None]+fraction*error*contraction
    assert result.min()>0 and np.max(np.abs(weights@result-point))<1e-10
    return np.ascontiguousarray(result)


@njit(cache=True)
def future_cost(prices,net,terminal_value=0.):
    future=-terminal_value*STATES
    step=50.;count=len(STATES)
    for t in range(len(prices)-1,-1,-1):
        current=np.empty(count)
        for i in range(count):
            low=max(0,i-int(P_MAX/ETA/step))
            high=min(count-1,i+int(P_MAX*ETA/step))
            best=np.inf
            for j in range(low,high+1):
                delta=STATES[j]-STATES[i]
                charge=max(delta,0.)/ETA;discharge=max(-delta,0.)*ETA
                grid=max(0.,net[t]+charge-discharge)
                cost=prices[t]*grid+future[j]
                if cost<best:best=cost
            current[i]=best
        future=current
    return future-future[0]

def price_view(day,release,actual,bank,config):
    if config.get('constant_price') is not None:return np.asarray(config['constant_price'],float)
    if config.get('known_price'):return actual[day].copy()
    return bank['updated'][day,release].copy() if config.get('price_update',True) else bank['midnight'][day].copy()


def stock_value(day,price0,bank,config):
    kind=config.get('terminal','current_min')
    if kind=='current_min':return float(price0[:36].min()/ETA)
    if kind=='next_q25':return float(np.quantile(bank['next_day'][day,:36],.25)/ETA)
    raise ValueError(kind)


def make_stage(day,release,state,original,data,supply,pbank,config):
    begin=36*release
    price=price_view(day,release,data['price'],pbank,config)
    value=stock_value(day,price_view(day,0,data['price'],pbank,config),pbank,config)
    # 保留原两问的浮点运算顺序；数学上均为功率乘1/6小时。
    actual=(data['load']-data['pv'])*(1/6) if config['question']==2 else (data['load']-data['pv'])/6
    if config['question']==2:
        sc,w,last=scenario_set(actual[:day],supply['q2_net'][:day],supply['q2_net'][day])
        sc=sc[:,begin:]
    else:
        pred=supply['q3_corrected' if config.get('nowcast') else 'q3_base']
        sc,w,ids=scenarios(day,release,pred,actual,supply['q3_features'],config.get('conditional',False))
        last=int(ids[-1])
    if not release:
        grid,reserve,detail=initial_plan(sc,w,price,state,value,config.get('midnight_refine',False))
    elif config['question']==2:
        grid=original[begin:].copy();reserve=dp_reserves(grid,sc,w,price[begin:],value);detail={}
    else:
        grid,reserve,detail=adjustment_plan(sc,w,price[begin:],state,original[begin:],begin,value,config.get('refine',False))
    detail.update(day=day,release=release,scenario_last_day=last,latest_observed_slot=begin-1,
        terminal_value=value,scenario_count=len(w),effective_scenarios=float(1/(w@w)),
        future_price_information='day_ahead_known_assumption' if config.get('known_price') else 'historical_forecast')
    assert last<day
    return grid,reserve,price,detail

BASE_CONFIG={2: {'question': 2, 'terminal': 'next_q25', 'midnight_refine': True, 'battery_update': True}, 3: {'question': 3, 'terminal': 'next_q25', 'nowcast': True, 'conditional': True, 'refine': True, 'releases': [1, 2, 3], 'midnight_refine': True}}

class Planner:
    def __init__(self,ctx,question,options):
        self.data,self.supply,self.bank0,self.optional=ctx
        self.question=question;self.options=options
        self.config=BASE_CONFIG[question].copy()
        self.bank=({k:self.optional[k] for k in ['midnight','updated','next_day']}
                   if options.get('price_fusion') else self.bank0)
        net=self.data['load']-self.data['pv']
        self.actual=net*(1/6) if question==2 else net/6
        self.pred=self.supply['q2_net'] if question==2 else self.supply['q3_corrected']
        self.curves={}

    def stage_data(self,d,r):
        ids=historical_ids(d);begin=36*r
        if self.question==2:
            sc,w,last=scenario_set(self.actual[:d],self.pred[:d],self.pred[d]);sc=sc[:,begin:]
        else:
            sc,w,ids=scenarios(d,r,self.pred,self.actual,self.supply['q3_features'],True)
        return sc,w,ids

    def anticipate(self,d,state,grid,reserve,price,value,sc,w,ids,fraction):
        # Every hypothetical forecast release is a historical vintage revision,
        # added to the current midnight forecast, never target-day future data.
        selected=np.unique(np.rint(np.linspace(0,len(ids)-1,min(7,len(ids)))).astype(int))
        ws=w[selected];ws/=ws.sum();contracts=[]
        for j in selected:
            historical=ids[j];book=grid.copy();rule=reserve.copy();s=state
            for r in range(4):
                begin=36*r
                if r:
                    hypoth=self.pred[d,0,begin:]+self.pred[historical,r,begin:]-self.pred[historical,0,begin:]
                    residual=self.actual[ids,begin:]-self.pred[ids,r,begin:]
                    future_sc=hypoth[None]+residual
                    ph=np.maximum(.001,price[begin:]+self.bank['updated'][historical,r,begin:]-self.bank['midnight'][historical,begin:])
                    quantile=weighted_quantile(future_sc,w,.8)
                    planned=initial_adjustment_plan(quantile,grid[begin:],ph,s,value)
                    book[begin:begin+36]=planned[:36]
                    rule[begin:]=dp_reserves(planned,future_sc,w,ph,value)
                for t in range(begin,begin+36):
                    *_,s=step_action(float(sc[j,t]),float(book[t]),s,float(rule[t]))
            contracts.append(book)
        contracts=np.array(contracts)
        target=weighted_quantile(contracts,ws,.5)
        changed=(1-fraction)*grid+fraction*target
        changed[:36]=grid[:36]
        expected=(1-fraction)*grid+fraction*(ws@contracts);expected[:36]=changed[:36]
        rule=dp_reserves(expected,sc,w,price,value)
        return changed,rule,dict(recourse_fraction=fraction,recourse_scenarios=len(selected),
            recourse_history_last=int(ids.max()),hypothetical_releases_from_completed_days=True)

    def stage(self,d,r,state,original):
        grid,reserve,price,details=make_stage(d,r,state,original,self.data,self.supply,self.bank,self.config)
        opt=self.options;begin=36*r
        if not opt:return grid,reserve,price,details
        sc,w,ids=self.stage_data(d,r);value=details['terminal_value']
        change_grid=(r==0 or self.question==3)
        origin=original[begin:] if r and self.question==3 else None
        if 'terminal_scale' in opt:
            value*=opt['terminal_scale']
            if not r:grid,reserve,extra=initial_plan(sc,w,price,state,value,True)
            elif self.question==3:grid,reserve,extra=adjustment_plan(sc,w,price[begin:],state,origin,begin,value,True)
            else:reserve=dp_reserves(grid,sc,w,price[begin:],value);extra={}
            details.update(extra,terminal_scale=opt['terminal_scale'],terminal_value=value)
        terminal=-value*STATES
        if opt.get('night_hours'):
            if d not in self.curves:
                stop=opt['night_hours']*6
                self.curves[d]=future_cost(self.bank['next_day'][d,:stop],self.optional['next_net'][d,:stop])
            fraction=opt.get('night_fraction',1.)
            terminal=(1-fraction)*(-value*(STATES-STATES[0]))+fraction*self.curves[d]
            reserve=reserves_with_curve(grid,sc,w,price[begin:],terminal)
            details.update(night_hours=opt['night_hours'],night_fraction=fraction,night_last_day=d-1)
        paired=opt.get('paired',0.)
        ps=paired_prices(price[begin:],w,self.data['price'][ids,begin:],self.bank['updated'][ids,r,begin:],paired)
        if opt.get('path') or opt.get('night_hours'):
            grid,reserve,records=optimize(grid,reserve,sc,w,state,ps,terminal,origin,begin,
                change_grid=change_grid,change_reserve=opt.get('reserve',False),rounds=opt.get('rounds',1))
            if change_grid and not opt.get('reserve'):
                candidate=reserves_with_curve(grid,sc,w,price[begin:],terminal)
                args=(sc,w,state,ps,terminal,np.zeros(len(grid)) if origin is None else origin,origin is not None,
                      np.zeros(len(grid),dtype=np.int64),False,False)
                objective=lambda rr:loss(np.empty(0),grid,np.ones(len(grid)),rr,*args)[0]
                if objective(candidate)<objective(reserve):reserve=candidate
            details.update(path_optimization=records,paired_price_fraction=paired,
                scenario_price_mean_error=float(np.max(abs(w@ps-price[begin:]))))
        if not r and opt.get('recourse'):
            assert self.question==3
            grid,reserve,extra=self.anticipate(d,state,grid,reserve,price,value,sc,w,ids,opt['recourse'])
            details.update(extra)
        details.update(scenario_last_day=int(ids.max()),price_residual_last_day=int(ids.max()))
        return grid,reserve,price,details

    def day(self,d,state):
        grid,rule,price,detail=self.stage(d,0,state,None)
        original=grid.copy();book=grid.copy();f=np.zeros((5,144));s=np.zeros(145);s[0]=state
        trades=np.zeros((4,2,144));bookings=np.zeros((4,144));aux=np.full((4,144),np.nan)
        prices=np.full((4,144),np.nan);prices[0]=price;aux[0]=original;logs=[detail]
        for r in range(4):
            begin=36*r
            if r:
                plan,reserve,price,detail=self.stage(d,r,float(s[begin]),original)
                aux[r,begin:]=plan;prices[r]=price;rule[begin:]=reserve;logs.append(detail)
                if self.question==3:
                    delta=plan[:36]-book[begin:begin+36]
                    trades[r,0,begin:begin+36]=np.maximum(delta,0)
                    trades[r,1,begin:begin+36]=np.maximum(-delta,0)
                    book[begin:begin+36]=plan[:36]
                    detail.update(commit_begin=begin,commit_stop=begin+36)
            bookings[r]=book
            for t in range(begin,begin+36):
                f[0,t]=book[t]
                f[1,t],f[2,t],f[3,t],f[4,t],s[t+1]=step_action(float(self.actual[d,t]),float(book[t]),float(s[t]),float(rule[t]))
        return dict(original_plan=original,flows=f,soc=s,reserves=rule,trades=trades,
                    bookings=bookings,auxiliary=aux,planning_prices=prices),logs


# 六、连续1月预运行：2月不重置电量。
def warmup(data,supply,pbank,output):
    """统一采用新价格下Q2基础策略预运行，不使用2月或未发布信息。"""
    state=6000.;flows=np.zeros((31,5,144));soc=np.zeros((31,145));logs=[];actual=(data['load']-data['pv'])*(1/6)
    for day in range(31):
        if day<START:
            grid=np.maximum(actual[day-1],0) if day else np.zeros(144)
            urgent=np.maximum(actual[day]-grid,0)
            flows[day]=np.stack([grid,np.zeros(144),np.zeros(144),urgent,grid+urgent-actual[day]])
            soc[day]=state
        else:
            grid,reserve,_,detail=make_stage(day,0,state,None,data,supply,pbank,dict(question=2,terminal='current_min'))
            flows[day],soc[day]=execute(actual[day],grid,reserve,state);logs.append(detail)
        state=float(soc[day,-1])
    np.savez_compressed(output/'january_warmup.npz',flows=flows,soc=soc,dates=DATES[:31].strftime('%Y-%m-%d').to_numpy(dtype='U10'))
    write_json(output/'january_warmup.json',dict(initial_soc=6000.,february_initial_soc=state,policy='q42_price_transfer',
        no_free_reset=True,decisions=logs,source_hashes=source_hashes()))
    print(f'新价格背景下，2月1日期初SOC={state:.12f} kWh。',flush=True)
    return state


# 七、附件五原格式导表及逐格核验。
def minutes(value):
    if hasattr(value,'hour'):return value.hour*60+value.minute
    text=str(value);h,m=map(int,text.replace('+1','').split(':')[:2])
    return h*60+m+(1440 if '+1' in text else 0)


def read_daily(path,sheets):
    wb=load_workbook(path,read_only=True,data_only=True);result=[]
    try:
        for sheet in wb.worksheets[:sheets]:
            rows=list(sheet.values)
            assert [minutes(x) for x in rows[0][1:]]==list(range(10,1441,10))
            rows=[r for r in rows[1:] if r[0] is not None]
            assert pd.DatetimeIndex([r[0] for r in rows]).equals(DATES)
            x=np.array([r[1:] for r in rows],float)
            assert x.shape==(365,144) and np.isfinite(x).all() and x.min()>=0
            result.append(x)
    finally:wb.close()
    assert len(result)==sheets
    return result


def clock(t):return f'{t//6}:{t%6*10:02d}'

def events(values,prices):
    result=[];t=0
    while t<144:
        if values[t]<=1e-6:t+=1;continue
        start=t
        while t<144 and values[t]>1e-6:t+=1
        result.append(dict(interval=f'{clock(start)}-{clock(t)}',quantity=float(values[start:t].sum()),
                           cost=float(5*prices[start:t]@values[start:t])))
    return result

def export_case(template,destination,a,summary,question):
    assert len(a['soc'])==334
    wb=load_workbook(template)
    expected=['计划购电量']+(['调整购电量'] if question==3 else [])+['充放电量','紧急购电量']
    assert wb.sheetnames==expected
    plan=wb['计划购电量'];store=wb['充放电量'];urgent=wb['紧急购电量']
    def prototypes(sheet,rows,columns):
        return [([(copy.copy(sheet.cell(r,c)._style),sheet.cell(r,c).value) for c in range(1,columns+1)],copy.copy(sheet.row_dimensions[r])) for r in rows]
    sr=prototypes(store,range(2,8),6);er=prototypes(urgent,range(2,5),3)
    store.delete_rows(2,store.max_row-1);urgent.delete_rows(2,urgent.max_row-1)
    def setrow(sheet,row,prototype):
        cells,dim=prototype;dim=copy.copy(dim);dim.index=row;sheet.row_dimensions[row]=dim
        for c,(style,value) in enumerate(cells,1):sheet.cell(row,c)._style=copy.copy(style);sheet.cell(row,c).value=value
    erow=2;event_rows=[];paper1=[];paper2=[];paper3=[];mapping=[]
    paper_dates={'2025-03-20','2025-06-21','2025-09-23','2025-12-21'}
    for i,date in enumerate(a['dates']):
        dt=pd.Timestamp(str(date)).to_pydatetime();row=i+2
        assert pd.Timestamp(plan.cell(row,1).value)==pd.Timestamp(dt)
        f=a['flows'][i];g=a['original_plan'][i];p=a['actual_prices'][i]
        pc=float(p@g);up=float(1.5*p@np.maximum(f[0]-g,0));refund=float(.5*p@np.maximum(g-f[0],0));ec=float(5*p@f[3])
        for sheet,quantities,cost in [(plan,g,pc)]+([(wb['调整购电量'],f[0],pc+up-refund)] if question==3 else []):
            assert sheet['B1'].value=='0:10-0:20' and sheet['EO1'].value=='0:00-0:10+1'
            for t,q in enumerate(quantities):sheet.cell(row,t+2,float(q)).number_format='0.00'
            sheet.cell(row,146,float(quantities.sum())).number_format='0.00'
            sheet.cell(row,147,cost).number_format='0.00'
        for block in range(6):
            r=2+6*i+block;sl=slice(24*block,24*(block+1));setrow(store,r,sr[block])
            store.cell(r,1).value=dt if block==0 else None
            store.cell(r,3,float(f[1,sl].sum())).number_format='0.00'
            store.cell(r,4,float(f[2,sl].sum())).number_format='0.00'
            store.cell(r,6).value=float(a['soc'][i,0]) if block==0 else (float(a['soc'][i,-1]) if block==1 else None)
            if block<2:store.cell(r,6).number_format='0.00'
            if str(date) in paper_dates:
                paper2.append(dict(question=f'4-{question}',date=str(date),interval=f'{clock(block*24)}-{clock((block+1)*24)}',
                    charge_kwh=float(f[1,sl].sum()),discharge_kwh=float(f[2,sl].sum()),soc_0=float(a['soc'][i,0]),soc_24=float(a['soc'][i,-1])))
        day_events=events(f[3],p) or [dict(interval='无',quantity=0.,cost=0.)]
        count=max(3,len(day_events))
        for j in range(count):
            setrow(urgent,erow,er[0 if j==0 else (2 if j==count-1 else 1)])
            urgent.cell(erow,1).value=dt if j==0 else None
            urgent.cell(erow,2).value=day_events[j]['interval'] if j<len(day_events) else None
            urgent.cell(erow,3).value=day_events[j]['quantity'] if j<len(day_events) else None
            urgent.cell(erow,3).number_format='0.00';erow+=1
        for event in day_events:
            record=dict(question=f'4-{question}',date=str(date),**event);event_rows.append(record)
            if str(date) in paper_dates:paper3.append(record)
        if str(date) in paper_dates:
            record=dict(question=f'4-{question}',date=str(date),original_daily_kwh=float(g.sum()),final_daily_kwh=float(f[0].sum()),
                        planned_cost=pc,up_cost=up,refund=refund,contract_cost=pc+up-refund,emergency_cost=ec,total_cost=pc+up-refund+ec)
            for t in [60,72,84,96,108,120]:
                record[f'original_{clock(t)}-{clock(t+1)}']=float(g[t]);record[f'final_{clock(t)}-{clock(t+1)}']=float(f[0,t])
            paper1.append(record)
    for t in range(144):mapping.append(dict(excel_column=t+2,template_header=plan.cell(1,t+2).value,model_interval=f'{clock(t)}-{clock(t+1)}'))
    wb.save(destination);wb.close()
    for label,records in [('紧急购电事件',event_rows),('指定日期_表1',paper1),('指定日期_表2',paper2),('指定日期_表3',paper3),('时间映射',mapping)]:
        pd.DataFrame(records).to_csv(destination.parent/f'result4-{question}_{label}.csv',index=False,encoding='utf-8-sig')
    verify_excel(template,destination,a,summary,question)
    print(str(destination),flush=True)


def verify_excel(template,destination,a,summary,question):
    src=load_workbook(template,data_only=False);wb=load_workbook(destination,data_only=True)
    try:
        assert wb.sheetnames==src.sheetnames
        for s,t in zip(src,wb):
            assert [c.value for c in s[1]]==[c.value for c in t[1]]
            assert list(s.merged_cells.ranges)==list(t.merged_cells.ranges)
            assert s.freeze_panes==t.freeze_panes
            for c,dim in s.column_dimensions.items():assert t.column_dimensions[c].width==dim.width
            assert not any(c.data_type=='e' for row in t for c in row)
        for name,expected in [('计划购电量',a['original_plan'])]+([('调整购电量',a['flows'][:,0])] if question==3 else []):
            s=wb[name];assert (s.max_row,s.max_column)==(335,147)
            actual=np.array([[s.cell(r,c).value for c in range(2,146)] for r in range(2,336)])
            assert abs(actual-expected).max()<1e-8
            assert pd.DatetimeIndex([s.cell(r,1).value for r in range(2,336)]).equals(DATES[31:])
            for i in range(334):assert abs(s.cell(i+2,146).value-expected[i].sum())<1e-7
        store=wb['充放电量'];assert (store.max_row,store.max_column)==(2005,6)
        for i in range(334):
            row=2+6*i
            assert pd.Timestamp(store.cell(row,1).value)==DATES[i+31]
            assert all(store.cell(row+j,1).value is None for j in range(1,6))
            assert abs(store.cell(row,6).value-a['soc'][i,0])<1e-7
            assert abs(store.cell(row+1,6).value-a['soc'][i,-1])<1e-7
            for j in range(6):
                assert abs(store.cell(row+j,3).value-a['flows'][i,1,24*j:24*(j+1)].sum())<1e-7
                assert abs(store.cell(row+j,4).value-a['flows'][i,2,24*j:24*(j+1)].sum())<1e-7
        urgent=wb['紧急购电量'];quantities=sum(float(r[2] or 0) for r in urgent.iter_rows(min_row=2,values_only=True))
        assert abs(quantities-a['flows'][:,3].sum())<1e-6
        dates=[r[0] for r in urgent.iter_rows(min_row=2,values_only=True) if r[0] is not None]
        assert pd.DatetimeIndex(dates).equals(DATES[31:])
        planned=sum(wb['计划购电量'].cell(r,147).value for r in range(2,336))
        assert abs(planned-summary['planned_cost'])<1e-6
        contract=planned if question==2 else sum(wb['调整购电量'].cell(r,147).value for r in range(2,336))
        assert abs(contract+summary['emergency_cost']-summary['total_cost'])<1e-6
        write_json(destination.parent/f'result4-{question}_excel_audit.json',dict(all_values_reconciled=True,
            original_headers_styles_and_date_groups_preserved=True,sheets=wb.sheetnames,
            shapes={s.title:[s.max_row,s.max_column] for s in wb},contract_cost=contract,emergency_kwh=quantities,
            adjustment_sheet_contains_final_contract_quantities=True if question==3 else None))
    finally:src.close();wb.close()


# 八、显式输入、运行入口和独立核验。
def source_hashes():
    return dict(_RUN_INPUTS)


def read_data(attachment2,attachment3,attachment4):
    load,pv=read_daily(attachment2,2)
    price=read_daily(attachment4,1)[0]
    if price.min()<=0:raise ValueError('附件4电价必须为正。')
    wb=load_workbook(attachment3,read_only=True,data_only=True)
    raw=np.full((365,4,24),np.nan);seen=set();current=None
    try:
        for row in list(wb.active.values)[1:]:
            if row[0] is not None and str(row[0]).strip():current=pd.Timestamp(row[0])
            if current is None:raise ValueError('附件3缺少日期。')
            d=(current-DATES[0]).days
            h=row[1].hour if hasattr(row[1],'hour') else int(str(row[1]).split(':')[0])
            if h not in (0,6,12,18) or (d,h) in seen:raise ValueError('附件3发布时间错误或重复。')
            seen.add((d,h));raw[d,h//6]=np.array(row[2:26],float)
    finally:wb.close()
    if len(seen)!=1460 or not np.isfinite(raw).all() or raw.min()<0:raise ValueError('附件3数据不完整。')
    return dict(dates=DATES,load=load,pv=pv,price=price,raw=raw)


def archived_components(path,data):
    saved=read_npz(path)
    if list(saved['component_names'])!=COMPONENT_NAMES:raise ValueError('基础模型名称或顺序不一致。')
    if not np.array_equal(saved['dates'],DATES.strftime('%Y-%m-%d').to_numpy(dtype='U10')):
        raise ValueError('预测留档日期与附件不一致。')
    for key in ['load','pv']:
        if str(saved[key+'_sha256'].item())!=array_digest(data[key]):
            raise ValueError('预测留档与附件2不符，请使用retrain模式。')
    # 原四分量不使用电价；留档内price_sha256对应问题2固定电价，不能错误地与附件4比较。
    x=saved['component_predictions_kw']
    if x.shape!=(2,4,365,144) or not np.isfinite(x[:,:,START:]).all() or x[:,:,START:].min()<0:
        raise ValueError('四分量预测形状或有效性错误。')
    return x


def make_components(args,data,output,versions):
    if args.mode=='archive':
        archive=required_input(args.archive,'基础预测留档')
        _RUN_INPUTS['archive']=digest(archive)
        print('显式读取基础四分量；重新计算后续全部过程。',flush=True)
        return archived_components(archive,data)
    if str(args.archive).strip():raise ValueError('retrain不使用--archive，请移除此参数。')
    # 两个单文件共用的训练段哈希，不受题号与调度段的差异影响。
    import inspect
    training_text='\n'.join(inspect.getsource(f) for f in [ModelConfig,feature_row_from_history,
        build_feature_cube,make_model,fit_model_for_origin,predict_with_fitted_model,fit_predict_origin,
        weighted_error_score,validation_origins,candidate_indices,choose_trainable_config,config_by_name,
        periodic_features,fit_periodic_predict,build_calendar,train_component])
    signature=hashlib.sha256((training_text+repr(MODEL_CONFIGS)+repr(DEFAULT_CONFIG)+str(args.threads)
        +json.dumps(versions,sort_keys=True)+array_digest(data['load'])+array_digest(data['pv'])).encode()).hexdigest()
    cache=output/'训练缓存';cache.mkdir(exist_ok=True)
    jobs=[(name,data[name],DATES.strftime('%Y-%m-%d').to_numpy(dtype='U10'),str(cache/f'{name}.npz'),
           signature,args.threads,args.training_end_index) for name in ['load','pv']]
    if args.workers==1:
        for job in jobs:train_component(*job)
    else:
        with ProcessPoolExecutor(max_workers=2) as pool:
            futures=[pool.submit(train_component,*job) for job in jobs]
            for f in futures:f.result()
    if args.training_end_index<365:
        print('短程训练完成；不导出正式年度结果。原命令去掉--training-end-index可继续。',flush=True)
        return None
    states=[read_npz(cache/f'{name}.npz') for name in ['load','pv']]
    assert all(x['done'][START:].all() for x in states)
    rows=[]
    for component,state in zip(['load','pv'],states):
        for d in range(START,365):
            rows.append(dict(component=component,date=str(DATES[d].date()),last_training_day=d-1,
                parameter_choice_day=START if d<31 else d-DATES[d].day+1,
                lgb_config=state['selected'][d,0],ridge_config=state['selected'][d,1],
                calendar_alpha=[.01,.1,1.][int(state['calendar_active'][d])]))
    pd.DataFrame(rows).to_csv(output/'训练与选参记录.csv',index=False,encoding='utf-8-sig')
    return np.stack([x['predictions'] for x in states])


def audit_arrays(a,data,summary):
    f,s,g,p,trades=[a[k] for k in ['flows','soc','original_plan','actual_prices','trades']]
    n=len(s);net=(data['load'][31:31+n]-data['pv'][31:31+n])/6
    errors=dict(balance=float(np.max(abs(f[:,0]+f[:,2]+f[:,3]-net-f[:,1]-f[:,4]))),
        state_equation=float(np.max(abs(np.diff(s,axis=1)-ETA*f[:,1]+f[:,2]/ETA))),
        cross_day=float(np.max(abs(s[1:,0]-s[:-1,-1]))) if n>1 else 0.)
    assert max(errors.values())<1e-7
    assert s.min()>=S_MIN-1e-7 and s.max()<=S_MAX+1e-7
    assert f.min()>=-1e-7 and f[:,1:3].max()<=P_MAX+1e-7
    assert np.minimum(f[:,1],f[:,2]).max()<1e-8
    assert np.array_equal(p,data['price'][31:31+n])
    up=np.maximum(f[:,0]-g,0);down=np.maximum(g-f[:,0],0)
    assert np.max(abs(trades[:,:,0].sum(1)-up))<1e-7
    assert np.max(abs(trades[:,:,1].sum(1)-down))<1e-7
    assert not np.count_nonzero(trades[:,0])
    for r in [1,2,3]:
        assert not np.count_nonzero(trades[:,r,:,:r*36])
        assert not np.count_nonzero(trades[:,r,:,(r+1)*36:])
    if QUESTION==2:assert np.array_equal(f[:,0],g) and not np.count_nonzero(trades)
    fee=float(np.sum(p*g)+1.5*np.sum(p*up)-.5*np.sum(p*down)+5*np.sum(p*f[:,3]))
    assert abs(fee-summary['total_cost'])<1e-6
    return dict(passed=True,days=n,errors=errors,cash_recomputed=fee,
        trade_times=[6,12,18] if QUESTION==3 else [],current_interval_feedback_approximation=True,
        independent_holdout=False,full_year_model_selection=True)


def replay(data,supply,prices,fused,initial,output,end=365):
    options=dict(price_fusion=True,path=True,paired=1. if QUESTION==2 else 0.)
    if QUESTION==3:options['recourse']=.5
    planner=Planner((data,supply,prices,fused),QUESTION,options)
    state=initial;arrays={};rows=[];logs=[];started=time.perf_counter()
    for d in range(31,end):
        result,records=planner.day(d,state);logs.extend(records)
        for k,x in result.items():arrays.setdefault(k,[]).append(x)
        f=result['flows'];g=result['original_plan'];s=result['soc'];p=data['price'][d]
        pc=float(p@g);up=float(1.5*p@np.maximum(f[0]-g,0));refund=float(.5*p@np.maximum(g-f[0],0));ec=float(5*p@f[3])
        state=float(s[-1]);rows.append(dict(date=str(DATES[d].date()),planned_cost=pc,up_cost=up,refund=refund,
            emergency_cost=ec,total_cost=pc+up-refund+ec,soc_start=float(s[0]),soc_end=state,
            grid_kwh=float(f[0].sum()),emergency_kwh=float(f[3].sum()),spill_kwh=float(f[4].sum()),
            emergency_intervals=int((f[3]>1e-6).sum())))
        if (d-31)%30==0 or d==end-1:print(f'4-{QUESTION} {DATES[d].date()} {time.perf_counter()-started:.1f}s',flush=True)
    df=pd.DataFrame(rows);a={k:np.array(v) for k,v in arrays.items()}
    summary={k:float(df[k].sum()) for k in ['planned_cost','up_cost','refund','emergency_cost','total_cost','grid_kwh','emergency_kwh','spill_kwh']}
    v=float(np.quantile(prices['next_day'][31,:36],.25)/.9)
    summary.update(question=QUESTION,options=options,days=end-31,initial_soc=initial,final_soc=state,
        inventory_adjusted_cost=summary['total_cost']+v*(initial-state),inventory_comparison_value=v,
        seconds=time.perf_counter()-started,emergency_days=int((df.emergency_intervals>0).sum()))
    a.update(dates=DATES[31:end].strftime('%Y-%m-%d').to_numpy(dtype='U10'),actual_prices=data['price'][31:end])
    atomic_npz(output/'arrays.npz',**a)
    df.to_csv(output/'daily.csv',index=False,encoding='utf-8-sig')
    write_json(output/'summary.json',summary);write_json(output/'decisions.json',logs)
    write_json(output/'物理与费用核验.json',audit_arrays(a,data,summary))
    return a,summary


def main():
    p=argparse.ArgumentParser(description=f'第四问4-{QUESTION}独立Python程序')
    p.add_argument('--attachment2',default=ATTACHMENT2_PATH)
    p.add_argument('--attachment3',default=ATTACHMENT3_PATH)
    p.add_argument('--attachment4',default=ATTACHMENT4_PATH)
    p.add_argument('--template',default=RESULT_TEMPLATE_PATH)
    p.add_argument('--output-dir',default=OUTPUT_DIRECTORY)
    p.add_argument('--mode',choices=['archive','retrain'],default=RUN_MODE)
    p.add_argument('--archive',default=ARCHIVED_PREDICTIONS_PATH)
    p.add_argument('--workers',type=int,choices=[1,2],default=TRAINING_WORKERS)
    p.add_argument('--threads',type=int,default=LIGHTGBM_THREADS)
    p.add_argument('--training-end-index',type=int,default=365,help='仅重训：小于365时只做阶段训练')
    p.add_argument('--end-index',type=int,default=365,help='调度截止索引；小于365不导出正式表')
    args=p.parse_args();started=time.perf_counter()
    if not 16<=args.training_end_index<=365 or not 32<=args.end_index<=365 or args.threads<1:
        p.error('索引或线程参数超出范围。')
    if args.mode=='archive' and args.training_end_index!=365:p.error('archive不进行阶段训练。')
    paths=[required_input(getattr(args,f'attachment{i}'),f'附件{i}') for i in (2,3,4)]
    template=required_input(args.template,'结果空白模板')
    output=Path(args.output_dir).expanduser().resolve() if args.output_dir.strip() else Path(__file__).resolve().parent/f'4-{QUESTION}_{args.mode}_输出'
    if output/f'result4-{QUESTION}.xlsx' in paths+[template]:raise ValueError('输出不能覆盖输入。')
    wb=load_workbook(template,read_only=True,data_only=False)
    try:
        expected=['计划购电量']+(['调整购电量'] if QUESTION==3 else [])+['充放电量','紧急购电量']
        if wb.sheetnames!=expected or wb.worksheets[0]['B1'].value!='0:10-0:20' or wb.worksheets[0]['EO1'].value!='0:00-0:10+1':
            raise ValueError('模板与题号或附件五原表头不一致。')
    finally:wb.close()
    versions={name:metadata.version(name) for name in ['numpy','pandas','scipy','scikit-learn','lightgbm','numba','llvmlite','openpyxl','threadpoolctl']}
    versions['python']=sys.version
    input_files=list(zip(['attachment2','attachment3','attachment4'],paths))
    input_files.extend([('template',template),('source',Path(__file__))])
    for label,path in input_files:_RUN_INPUTS[label]=digest(path)
    output.mkdir(exist_ok=True,parents=True)
    provenance=output/'运行来源与环境.json'
    previous=json.loads(provenance.read_text(encoding='utf-8')) if provenance.exists() else None
    if previous and (previous['run_mode']!=args.mode or previous['question']!=QUESTION or previous['inputs']!=_RUN_INPUTS):
        raise ValueError('输出目录属于其他模式、题号或输入，请换一个目录。')
    initial_inputs=dict(_RUN_INPUTS)
    write_json(provenance,dict(question=QUESTION,run_mode=args.mode,versions=versions,inputs=initial_inputs,status='running'))
    data=read_data(*paths)
    data['components']=make_components(args,data,output,versions)
    if data['components'] is None:return
    atomic_npz(output/'本次基础预测.npz',components=data['components'])
    with threadpool_limits(1):
        prepare_supply(data,output);supply=read_npz(output/'supply_forecasts.npz')
        prepare_prices(data,output);prices=read_npz(output/'price_forecasts.npz')
        prepare_price_fusion(data,output,prices);fused=read_npz(output/'new_forecasts.npz')
        initial=warmup(data,supply,prices,output)
        a,summary=replay(data,supply,prices,fused,initial,output,args.end_index)
    summary.update(run_mode=args.mode,base_models_retrained=args.mode=='retrain',independent_holdout=False,
                   full_year_configuration_selection=True)
    write_json(output/'summary.json',summary)
    if args.end_index==365:export_case(template,output/f'result4-{QUESTION}.xlsx',a,summary,QUESTION)
    write_json(provenance,dict(question=QUESTION,run_mode=args.mode,versions=versions,inputs=initial_inputs,
        all_inputs=dict(_RUN_INPUTS),status='complete' if args.end_index==365 else 'partial',
        no_external_project_modules=True,embedded_final_results=False,
        elapsed_seconds=time.perf_counter()-started))
    print(f'4-{QUESTION} 总费用={summary["total_cost"]:.8f}元；结果已生成。',flush=True)


if __name__=='__main__':
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    main()
