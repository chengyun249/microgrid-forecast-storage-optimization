#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""第二问正式单层融合方案：从附件训练、连续调度、按附件五模板导出。

独立运行文件，不导入项目中的其他 Python 脚本，不内嵌历史预测或答案。
先填写下方三项路径，或使用命令行 --attachment1/--attachment2/--template。
默认从原始数据重新训练；另支持--mode archive显式读取预测留档并完整复算调度。
模型口径：00:10功率代表00:00—00:10；购电在0点固定，电池允许当前区间反馈。
模板表头原样保留，按144列原顺序填值；表头位移问题详见配套Markdown。
"""
from __future__ import annotations

# ==================== 一、用户配置：附件路径 ====================
ATTACHMENT1_PATH = r""       # 附件1.xlsx：仅读取固定电价
ATTACHMENT2_PATH = r""       # 附件2.xlsx：全年负载和光伏
RESULT2_TEMPLATE_PATH = r""  # 附件5文件夹中的空白result2.xlsx
OUTPUT_DIRECTORY = ""        # 留空：在脚本旁按运行模式分别生成输出文件夹
RUN_MODE = "retrain"          # retrain完整训练；archive使用显式指定的原四分量预测
ARCHIVED_PREDICTIONS_PATH = r""  # 仅archive模式使用
TRAINING_WORKERS = 2          # 负载/PV独立训练；内存不足时改1
LIGHTGBM_THREADS = 1          # 每个训练进程的LightGBM线程数

import argparse
from concurrent.futures import ProcessPoolExecutor
import copy
from dataclasses import dataclass
from datetime import date, datetime
from functools import lru_cache
import hashlib
import importlib.metadata as metadata
import json
import math
import os
from pathlib import Path
import sys
import time
from typing import Any

import lightgbm as lgb
import numpy as np
import pandas as pd
from numba import njit
from openpyxl import load_workbook
from scipy.ndimage import gaussian_filter1d
from scipy.optimize import linprog, minimize
from scipy.sparse import coo_matrix
from sklearn.linear_model import Ridge
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from threadpoolctl import threadpool_limits

PERIODS_PER_DAY = 144
FEATURE_START_INDEX = 14
START = 15
FORMAL = 31
DT = 1 / 6
RANDOM_STATE = 2026
ETA = 0.9
S_MIN, S_MAX, P_MAX = 1200., 10800., 5000. / 6
BLOCK = np.arange(144) // 24
COMPONENT_NAMES = ['legacy_daily_lgb', 'legacy_daily_ridge', 'periodic_lgb_l1', 'calendar_curve']
EXPECTED_COST = 13475274.43452325  # 仅供结果核对，不参与任何预测、选参或规划


def write_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def atomic_npz(path, **data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + '.tmp')
    with temporary.open('wb') as stream:
        np.savez_compressed(stream, **data)
    os.replace(temporary, path)


def read_npz(path):
    with np.load(path, allow_pickle=False) as z:
        return {name: z[name].copy() for name in z.files}


def read_inputs(attachment1, attachment2):
    """只读取附件1电价与附件2两个工作表，并严格校验日期和144列。"""
    workbook = load_workbook(attachment2, read_only=True, data_only=True)
    arrays, dates, headers = [], [], []
    try:
        if len(workbook.worksheets) < 2:
            raise ValueError('附件2必须包含负载、光伏两个工作表。')
        for sheet in workbook.worksheets[:2]:
            rows = list(sheet.values)
            headers.append(rows[0][1:])
            records = [row for row in rows[1:] if row[0] is not None]
            dates.append(pd.DatetimeIndex([row[0] for row in records]))
            arrays.append(np.array([row[1:] for row in records], float))
    finally:
        workbook.close()
    expected = pd.date_range('2025-01-01', '2025-12-31')
    if not all(x.equals(expected) for x in dates):
        raise ValueError('附件2日期必须完整、依序覆盖2025年365天。')
    if headers[0] != headers[1] or any(x.shape != (365, 144) for x in arrays):
        raise ValueError('附件2的两个工作表必须使用相同的144个时间列。')
    if not all(np.isfinite(x).all() and (x >= 0).all() for x in arrays):
        raise ValueError('附件2存在空值、非数字或负功率。')
    workbook = load_workbook(attachment1, read_only=True, data_only=True)
    try:
        rows = [row for row in list(workbook.active.values)[1:] if row[0] is not None]
        prices = np.array([row[1] for row in rows], float)
    finally:
        workbook.close()
    def minutes(value):
        if hasattr(value, 'hour'):
            return value.hour * 60 + value.minute
        raw = str(value).replace('+1', '')
        hour, minute = map(int, raw.split(':')[:2])
        return hour * 60 + minute + (1440 if '+1' in str(value) else 0)
    expected_minutes = list(range(10, 1441, 10))
    if [minutes(row[0]) for row in rows] != expected_minutes or [minutes(x) for x in headers[0]] != expected_minutes:
        raise ValueError('附件时间列应从00:10依次到0:00+1，每10分钟一个值。')
    if prices.shape != (144,) or not np.isfinite(prices).all() or (prices <= 0).any():
        raise ValueError('附件1电价须为144个正数。')
    return expected, arrays[0], arrays[1], prices


# ==================== 二、两个历史模型：逐日训练、月初选参 ====================

@dataclass(frozen=True)
class ModelConfig:
    name: str
    family: str
    target_kind: str
    objective: str = "regression"
    n_estimators: int = 0
    learning_rate: float = 0.0
    num_leaves: int = 0
    min_child_samples: int = 0
    reg_alpha: float = 0.0
    reg_lambda: float = 0.0
    ridge_alpha: float = 0.0
    train_window_days: int | None = None


MODEL_CONFIGS: tuple[ModelConfig, ...] = (
    ModelConfig(
        "lgbm_level_l2_large",
        "lgbm",
        "level",
        objective="regression",
        n_estimators=180,
        learning_rate=0.040,
        num_leaves=63,
        min_child_samples=80,
        reg_alpha=0.20,
        reg_lambda=5.0,
    ),
    ModelConfig(
        "lgbm_level_l2_moretrees",
        "lgbm",
        "level",
        objective="regression",
        n_estimators=260,
        learning_rate=0.025,
        num_leaves=31,
        min_child_samples=60,
        reg_alpha=0.10,
        reg_lambda=2.0,
    ),
    ModelConfig(
        "lgbm_residual_l1_moretrees",
        "lgbm",
        "residual",
        objective="regression_l1",
        n_estimators=220,
        learning_rate=0.030,
        num_leaves=31,
        min_child_samples=60,
        reg_alpha=0.10,
        reg_lambda=2.0,
    ),
    ModelConfig("ridge_level_a1", "ridge", "level", ridge_alpha=1.0),
    ModelConfig("ridge_level_a10", "ridge", "level", ridge_alpha=10.0),
    ModelConfig("ridge_level_a100", "ridge", "level", ridge_alpha=100.0),
    ModelConfig("ridge_residual_a10", "ridge", "residual", ridge_alpha=10.0),
)


DEFAULT_CONFIG = {
    ("load", "lgbm"): "lgbm_level_l2_large",
    ("pv", "lgbm"): "lgbm_residual_l1_moretrees",
    ("load", "ridge"): "ridge_level_a10",
    ("pv", "ridge"): "ridge_residual_a10",
    ("load", "seasonal"): "seasonal_lag7",
    ("pv", "seasonal"): "seasonal_lag1",
}


def feature_row_from_history(history: np.ndarray, target_date: date) -> np.ndarray:
    """由目标日前的逐日曲线构造目标日特征，不读取目标日真实值。"""
    history = np.asarray(history, dtype=np.float64)
    if history.ndim != 2 or history.shape[1] != PERIODS_PER_DAY:
        raise ValueError("history 必须是 n×144 数组。")
    if len(history) < 14:
        raise ValueError("至少需要14个完整历史日。")

    slots = np.arange(PERIODS_PER_DAY, dtype=np.float64)
    slot_angle = 2.0 * np.pi * slots / PERIODS_PER_DAY
    lag1, lag2, lag3 = history[-1], history[-2], history[-3]
    lag7, lag14 = history[-7], history[-14]
    hist3, hist7, hist14 = history[-3:], history[-7:], history[-14:]
    hist28 = history[-min(28, len(history)) :]
    weekday = target_date.weekday()
    doy = target_date.timetuple().tm_yday
    weekday_angle = 2.0 * np.pi * weekday / 7.0
    doy_angle = 2.0 * np.pi * (doy - 1) / 365.25

    row = np.column_stack(
        [
            lag1,
            lag2,
            lag3,
            lag7,
            lag14,
            np.mean(hist3, axis=0),
            np.mean(hist7, axis=0),
            np.mean(hist14, axis=0),
            np.mean(hist28, axis=0),
            np.std(hist3, axis=0),
            np.std(hist7, axis=0),
            np.std(hist14, axis=0),
            np.std(hist28, axis=0),
            lag1 - lag2,
            lag7 - lag14,
            np.full(PERIODS_PER_DAY, np.mean(lag1)),
            np.full(PERIODS_PER_DAY, np.max(lag1)),
            np.full(PERIODS_PER_DAY, np.std(lag1)),
            np.full(PERIODS_PER_DAY, np.mean(lag7)),
            np.full(PERIODS_PER_DAY, np.max(lag7)),
            slots / (PERIODS_PER_DAY - 1),
            np.sin(slot_angle),
            np.cos(slot_angle),
            np.sin(2.0 * slot_angle),
            np.cos(2.0 * slot_angle),
            np.full(PERIODS_PER_DAY, weekday / 6.0),
            np.full(PERIODS_PER_DAY, float(weekday >= 5)),
            np.full(PERIODS_PER_DAY, math.sin(weekday_angle)),
            np.full(PERIODS_PER_DAY, math.cos(weekday_angle)),
            np.full(PERIODS_PER_DAY, (doy - 1) / 364.0),
            np.full(PERIODS_PER_DAY, math.sin(doy_angle)),
            np.full(PERIODS_PER_DAY, math.cos(doy_angle)),
            np.full(PERIODS_PER_DAY, target_date.month / 12.0),
        ]
    )
    if row.shape != (PERIODS_PER_DAY, 33) or not np.all(np.isfinite(row)):
        raise RuntimeError("特征行形状或数值异常。")
    return row.astype(np.float32)


def build_feature_cube(series: np.ndarray, dates: tuple[date, ...]) -> np.ndarray:
    cube = np.full((len(dates), PERIODS_PER_DAY, 33), np.nan, dtype=np.float32)
    for target_index in range(FEATURE_START_INDEX, len(dates)):
        cube[target_index] = feature_row_from_history(
            series[:target_index], dates[target_index]
        )
    if not np.all(np.isfinite(cube[FEATURE_START_INDEX:])):
        raise RuntimeError("有效预测日的特征中存在非有限值。")
    return cube


def make_model(config: ModelConfig, lgbm_threads: int) -> Any:
    if config.family == "lgbm":
        return lgb.LGBMRegressor(
            objective=config.objective,
            n_estimators=config.n_estimators,
            learning_rate=config.learning_rate,
            num_leaves=config.num_leaves,
            min_child_samples=config.min_child_samples,
            subsample=0.90,
            subsample_freq=1,
            colsample_bytree=0.90,
            reg_alpha=config.reg_alpha,
            reg_lambda=config.reg_lambda,
            max_bin=127,
            random_state=RANDOM_STATE,
            n_jobs=lgbm_threads,
            verbosity=-1,
            deterministic=True,
            force_col_wise=True,
        )
    if config.family == "ridge":
        return make_pipeline(
            StandardScaler(),
            Ridge(alpha=config.ridge_alpha, solver="cholesky"),
        )
    raise ValueError(f"未知模型族：{config.family}")


def fit_model_for_origin(
    series: np.ndarray,
    feature_cube: np.ndarray,
    config: ModelConfig,
    target_index: int,
    residual_lag: int,
    lgbm_threads: int,
) -> Any:
    train_start = FEATURE_START_INDEX
    if config.train_window_days is not None:
        train_start = max(train_start, target_index - config.train_window_days)
    train_days = np.arange(train_start, target_index, dtype=int)
    if len(train_days) == 0:
        raise RuntimeError(f"{target_index=} 没有有效训练日。")
    assert int(np.max(train_days)) < target_index
    x_train = feature_cube[train_days].reshape(-1, feature_cube.shape[-1])
    if config.target_kind == "level":
        y_train = series[train_days].reshape(-1)
    elif config.target_kind == "residual":
        if int(np.min(train_days)) - residual_lag < 0:
            raise RuntimeError("残差模型基线滞后越界。")
        y_train = (series[train_days] - series[train_days - residual_lag]).reshape(-1)
    else:
        raise ValueError(f"未知目标类型：{config.target_kind}")
    model = make_model(config, lgbm_threads)
    model.fit(x_train, y_train)
    return model


def predict_with_fitted_model(
    model: Any,
    config: ModelConfig,
    feature_row: np.ndarray,
    baseline: np.ndarray,
) -> np.ndarray:
    if config.family == "lgbm":
        raw = np.asarray(model.booster_.predict(feature_row), dtype=np.float64)
    else:
        raw = np.asarray(model.predict(feature_row), dtype=np.float64)
    if config.target_kind == "residual":
        raw = baseline + raw
    return np.clip(raw.reshape(PERIODS_PER_DAY), 0.0, None)


def fit_predict_origin(
    series: np.ndarray,
    dates: tuple[date, ...],
    feature_cube: np.ndarray,
    config: ModelConfig,
    target_index: int,
    residual_lag: int,
    lgbm_threads: int,
    include_next_day: bool,
) -> tuple[np.ndarray, np.ndarray | None]:
    """拟合 < target_index 的标签，预测 target_index 及可选递归下一日。"""
    model = fit_model_for_origin(
        series,
        feature_cube,
        config,
        target_index,
        residual_lag,
        lgbm_threads,
    )
    current_baseline = series[target_index - residual_lag]
    current = predict_with_fitted_model(
        model, config, feature_cube[target_index], current_baseline
    )
    if not include_next_day or target_index + 1 >= len(dates):
        return current, None

    augmented = np.vstack([series[:target_index], current[None, :]])
    next_feature = feature_row_from_history(augmented, dates[target_index + 1])
    next_baseline = augmented[-residual_lag]
    following = predict_with_fitted_model(model, config, next_feature, next_baseline)
    return current, following


def weighted_error_score(
    actual: np.ndarray,
    predicted: np.ndarray,
    origins: np.ndarray,
    asof_index: int,
    half_life_days: int,
) -> tuple[float, float, float]:
    if len(origins) == 0:
        return math.inf, math.inf, math.inf
    assert int(np.max(origins)) < asof_index
    error = actual[origins] - predicted[origins]
    if not np.all(np.isfinite(error)):
        return math.inf, math.inf, math.inf
    ages = asof_index - origins
    day_weights = np.power(0.5, ages / max(half_life_days, 1))
    point_weights = np.repeat(day_weights[:, None], PERIODS_PER_DAY, axis=1)
    denominator = float(np.sum(point_weights * np.abs(actual[origins])))
    wape = float(np.sum(point_weights * np.abs(error)) / max(denominator, 1.0))
    rmse = float(np.sqrt(np.sum(point_weights * error**2) / np.sum(point_weights)))
    scale = float(
        np.sum(point_weights * np.abs(actual[origins])) / np.sum(point_weights)
    )
    nrmse = rmse / max(scale, 1.0)
    score = 0.65 * wape + 0.35 * nrmse
    return score, wape, rmse


def validation_origins(
    target_index: int, oos_start_index: int, validation_days: int, stride: int
) -> np.ndarray:
    start = max(oos_start_index, target_index - validation_days)
    origins = list(range(start, target_index, stride))
    if target_index - 1 >= oos_start_index and (not origins or origins[-1] != target_index - 1):
        origins.append(target_index - 1)
    result = np.asarray(sorted(set(origins)), dtype=int)
    if len(result):
        assert int(np.max(result)) < target_index
    return result


def candidate_indices(family: str) -> list[int]:
    return [index for index, config in enumerate(MODEL_CONFIGS) if config.family == family]


def choose_trainable_config(
    component: str,
    family: str,
    series: np.ndarray,
    dates: tuple[date, ...],
    feature_cube: np.ndarray,
    hyper_predictions: np.ndarray,
    tuning_scores: np.ndarray,
    month_start_index: int,
    oos_start_index: int,
    residual_lag: int,
    args: argparse.Namespace,
) -> tuple[str, np.ndarray]:
    origins = validation_origins(
        month_start_index,
        oos_start_index,
        args.validation_days,
        args.tune_stride,
    )
    config_ids = candidate_indices(family)
    if len(origins) < args.minimum_tuning_origins:
        return DEFAULT_CONFIG[(component, family)], origins

    for config_id in config_ids:
        config = MODEL_CONFIGS[config_id]
        for origin in origins:
            if not np.all(np.isfinite(hyper_predictions[config_id, origin])):
                forecast, _ = fit_predict_origin(
                    series,
                    dates,
                    feature_cube,
                    config,
                    int(origin),
                    residual_lag,
                    args.lgbm_threads,
                    include_next_day=False,
                )
                hyper_predictions[config_id, origin] = forecast.astype(np.float32)
        score, _, _ = weighted_error_score(
            series,
            hyper_predictions[config_id],
            origins,
            month_start_index,
            args.tune_half_life,
        )
        tuning_scores[month_start_index, config_id] = score

    best_id = min(config_ids, key=lambda index: tuning_scores[month_start_index, index])
    return MODEL_CONFIGS[best_id].name, origins


def config_by_name(name: str) -> ModelConfig:
    return next(config for config in MODEL_CONFIGS if config.name == name)


def simplex_least_squares(
    x: np.ndarray, y: np.ndarray, regularization: float, prior: np.ndarray
) -> np.ndarray:
    component_count = x.shape[1]
    gram = (x.T @ x) / max(len(y), 1)
    cross = (x.T @ y) / max(len(y), 1)
    prior_error = float(np.mean((x @ prior - y) ** 2))
    penalty = regularization * max(prior_error, 1.0)
    gram = gram + (penalty + 1.0e-8) * np.eye(component_count)
    cross = cross + penalty * prior
    best: np.ndarray | None = None
    best_value = math.inf
    for mask in range(1, 1 << component_count):
        active = np.asarray(
            [index for index in range(component_count) if mask & (1 << index)], dtype=int
        )
        a = gram[np.ix_(active, active)]
        b = cross[active]
        ones = np.ones(len(active))
        kkt = np.block([[2.0 * a, ones[:, None]], [ones[None, :], np.zeros((1, 1))]])
        rhs = np.concatenate([2.0 * b, np.ones(1)])
        try:
            active_weight = np.linalg.solve(kkt, rhs)[:-1]
        except np.linalg.LinAlgError:
            active_weight = np.linalg.lstsq(kkt, rhs, rcond=None)[0][:-1]
        if np.min(active_weight) < -1.0e-8:
            continue
        active_weight = np.clip(active_weight, 0.0, None)
        active_weight /= np.sum(active_weight)
        weight = np.zeros(component_count)
        weight[active] = active_weight
        value = float(
            np.mean((x @ weight - y) ** 2)
            + penalty * np.sum((weight - prior) ** 2)
        )
        if value < best_value:
            best_value = value
            best = weight
    return prior.copy() if best is None else best



# ==================== 三、周期模型与单层融合 ====================

def periodic_features(history, target_date, component):
    """Only a prefix is accepted; all 144 target slots are built together."""
    x = feature_row_from_history(history, target_date.date()).astype(float)
    angle = 2*np.pi*(np.arange(144)+1)/144
    harmonics = np.column_stack([v for k in range(1, 5)
                               for v in (np.sin(k*angle), np.cos(k*angle))])
    smooth1 = gaussian_filter1d(history[-1], 2, mode='nearest')
    smooth7 = gaussian_filter1d(history[-7], 2, mode='nearest')
    smoothmean = gaussian_filter1d(history[-7:].mean(0), 2, mode='nearest')
    # Seasonal baseline is explicitly computed from past observations.
    base = (0.65*smooth7+0.35*smooth1 if component == 'load'
            else 0.65*smooth1+0.35*gaussian_filter1d(history[-3:].mean(0), 2))
    interactions = np.column_stack([harmonics * z[:, None]/1000
                                   for z in (smooth1, smooth7, smoothmean)])
    weekday = np.eye(7)[target_date.dayofweek]
    calendar = np.hstack([np.tile(weekday, (144, 1)),
                         np.einsum('tj,k->tjk', harmonics[:, :4], weekday).reshape(144, -1)])
    local = np.column_stack([smooth1, smooth7, smoothmean,
                             np.gradient(smooth1), np.gradient(smooth7),
                             np.maximum(history[-14:].max(0), 1)])
    return np.column_stack([x, harmonics, interactions, calendar, local]).astype('float32'), base


def fit_periodic_predict(x, y, baseline, dates, d, model_id, threads):
    train = np.arange(14, d)
    xx = x[train].reshape(-1, x.shape[-1])
    yy = (y[train] - baseline[train]).reshape(-1)
    half_life = 56 if model_id == 0 else 84
    sw = np.repeat(2.0 ** (-(d-1-train)/half_life), 144)
    if model_id == 0:
        scaler = StandardScaler().fit(xx, sample_weight=sw)
        model = Ridge(alpha=40, solver='cholesky')
        model.fit(scaler.transform(xx), yy, sample_weight=sw)
        pred = model.predict(scaler.transform(x[d]))
        pred = gaussian_filter1d(pred, 1, mode='nearest')
    else:
        model = lgb.LGBMRegressor(objective='regression' if model_id == 1 else 'regression_l1',
                    n_estimators=240, learning_rate=0.035, num_leaves=31,
                    min_child_samples=90, reg_lambda=8, reg_alpha=0.2,
                    max_bin=127, colsample_bytree=0.9, n_jobs=threads,
                    random_state=20260911, verbosity=-1, deterministic=True,
                    force_col_wise=True)
        model.fit(xx, yy, sample_weight=sw)
        pred = model.booster_.predict(x[d], num_threads=threads)
    return np.maximum(0, baseline[d]+pred)


def online_blend(actual, cube, dates, windows=(28, 56), groups=False):
    """Monthly choice among causal recipes, each using earlier OOS errors.

    groups=True uses six four-hour groups with shrinkage towards the global
    weights; this is a predeclared ablation, not selected with full-year data.
    """
    n, m, t = cube.shape
    pred = np.full((len(windows), n, t), np.nan)
    ws = np.full((len(windows), n, 6 if groups else 1, m), np.nan)
    prior = np.ones(m)/m
    for ri, window in enumerate(windows):
        for d in range(START, n):
            h = np.arange(max(START, d-window), d)
            w = prior
            if len(h) >= 7:
                xx = cube[h].transpose(0, 2, 1).reshape(-1, m)
                w = simplex_least_squares(xx, actual[h].reshape(-1), 0.15, prior)
            for group in range(6 if groups else 1):
                sl = slice(group*24, (group+1)*24) if groups else slice(0, 144)
                wg = w
                if groups and len(h) >= 14:
                    xx = cube[h, :, sl].transpose(0, 2, 1).reshape(-1, m)
                    local = simplex_least_squares(xx, actual[h, sl].reshape(-1), 0.3, w)
                    wg = 0.5*w+0.5*local
                ws[ri, d, group] = wg
                pred[ri, d, sl] = wg @ cube[d, :, sl]
    out = np.full_like(actual, np.nan)
    records = []
    active, cutoff = 0, -1
    for d in range(START, n):
        if d == START or dates[d].day == 1:
            h = np.arange(max(START, d-56), d)
            if len(h) >= 7:
                loss = np.mean((pred[:, h] - actual[h][None])**2, axis=(1, 2))
                active = int(loss.argmin())
            cutoff = d-1
        out[d] = pred[active, d]
        records.append(dict(date=str(dates[d].date()), window=windows[active],
                            monthly_choice_last_day=cutoff, weight_last_day=d-1,
                            weights=ws[active, d].tolist()))
    return out, records



# ==================== 四、日历曲线与完整基础训练 ====================

def build_calendar(dates, component):
    a=2*np.pi*(np.asarray(dates.dayofyear)-1)/365.25
    annual=np.column_stack([np.sin(a),np.cos(a),np.sin(2*a),np.cos(2*a)])
    if component=='pv':
        return annual
    week=np.eye(7)[dates.dayofweek]
    return np.column_stack([annual,week,
                           np.einsum('ij,ik->ijk',week,annual[:,:2]).reshape(len(dates),-1)])




def train_component(component, series, date_strings, cache_path, signature, threads=1, end_index=365):
    """从原始功率训练四个分量。每次只拟合目标日前的标签。

    仅复现正式方案所用的模型；未选入的旧季节融合、周期L2等不必运行。
    旧LightGBM/Ridge内部月度选参的候选池必须保留，不能用全年最终配置替代。
    """
    dates = pd.DatetimeIndex(date_strings)
    date_tuple = tuple(x.date() for x in dates)
    args = argparse.Namespace(validation_days=35, tune_stride=3, tune_half_life=21,
                              minimum_tuning_origins=5, lgbm_threads=threads)
    path = Path(cache_path)
    state = dict(predictions=np.full((4,365,144),np.nan),
        hyper=np.full((7,365,144),np.nan,dtype=np.float32),
        scores=np.full((365,7),np.nan), selected=np.full((365,2),'',dtype='U64'),
        calendar_bank=np.full((3,365,144),np.nan), calendar_active=np.full(365,-1,dtype=np.int16),
        done=np.zeros(365,dtype=bool))
    if path.exists():
        saved = read_npz(path)
        if str(saved.pop('signature').item()) != signature:
            raise ValueError('缓存与当前源码、输入或环境不符，请使用新的输出目录。')
        for key in state:
            if key not in saved or saved[key].shape != state[key].shape:
                raise ValueError('缓存字段或形状错误：' + key)
            state[key] = saved[key]
        completed = np.flatnonzero(state['done'])
        if len(completed) and not np.array_equal(completed,np.arange(START,completed[-1]+1)):
            raise ValueError('缓存完成日不连续。')
    print(f'[{component}] 已完成{int(state["done"].sum())}个训练日，准备特征。',flush=True)
    basic = build_feature_cube(series,date_tuple)
    feature_rows, base_rows = zip(*[periodic_features(series[:d],dates[d],component) for d in range(14,365)])
    periodic = np.full((365,144,feature_rows[0].shape[-1]),np.nan,dtype=np.float32)
    baseline = np.full_like(series,np.nan)
    periodic[14:],baseline[14:] = np.asarray(feature_rows),np.asarray(base_rows)
    calendar = build_calendar(dates,component)
    residual_lag = 7 if component == 'load' else 1
    active = [DEFAULT_CONFIG[(component,'lgbm')],DEFAULT_CONFIG[(component,'ridge')]]
    calendar_active = 1
    started = time.time()
    try:
        with threadpool_limits(limits=1, user_api='blas'):
            for d in range(START,end_index):
                if state['done'][d]:
                    active = state['selected'][d].tolist()
                    calendar_active = int(state['calendar_active'][d])
                    continue
                if d == START or dates[d].day == 1:
                    for family_index,family in enumerate(['lgbm','ridge']):
                        active[family_index],_ = choose_trainable_config(component,family,series,date_tuple,basic,
                            state['hyper'],state['scores'],d,START,residual_lag,args)
                    past = np.arange(max(START,d-56),d)
                    if len(past) >= 7:
                        calendar_active = int(np.mean((state['calendar_bank'][:,past]-series[past][None])**2,axis=(1,2)).argmin())
                for family_index,name in enumerate(active):
                    config = config_by_name(name)
                    current,_ = fit_predict_origin(series,date_tuple,basic,config,d,residual_lag,threads,False)
                    rounded = current.astype(np.float32)  # 原方案历史预测库以float32保存这两个分量
                    state['hyper'][MODEL_CONFIGS.index(config),d] = rounded
                    state['predictions'][family_index,d] = rounded
                    state['selected'][d,family_index] = name
                state['predictions'][2,d] = fit_periodic_predict(periodic,series,baseline,dates,d,2,threads)
                weights = 2.**(-(d-1-np.arange(d))/84.)
                for i,alpha in enumerate([.01,.1,1.]):
                    model = Ridge(alpha=alpha,solver='cholesky')
                    model.fit(calendar[:d],series[:d],sample_weight=weights)
                    raw = model.predict(calendar[d:d+1])[0]
                    state['calendar_bank'][i,d] = np.maximum(0,gaussian_filter1d(raw,1,mode='nearest'))
                state['predictions'][3,d] = state['calendar_bank'][calendar_active,d]
                state['calendar_active'][d] = calendar_active
                state['done'][d] = True
                if d % 7 == 0 or dates[d].day == 1 or d == end_index-1:
                    atomic_npz(path,signature=np.asarray(signature),**state)
                    print(f'[{component}] {dates[d].date()}，累计{time.time()-started:.1f}秒，已保存缓存。',flush=True)
    except BaseException:
        atomic_npz(path,signature=np.asarray(signature),**state)
        raise
    atomic_npz(path,signature=np.asarray(signature),**state)
    return str(path)


# ==================== 五、历史情景与电池反馈控制 ====================

@lru_cache(maxsize=30)
def matrix(k, t):
    n = 4*t+1+k*t
    ii = np.tile(np.arange(t), k)
    r = np.repeat(np.arange(k*t), 4)
    col = np.column_stack([ii, t+ii, 2*t+ii, 4*t+1+np.arange(k*t)]).ravel()
    a = coo_matrix((np.tile([-1., 1., -1., -1.], k*t), (r, col)),
                   shape=(k*t, n)).tocsr()
    r = np.repeat(np.arange(t), 4)
    col = np.column_stack([t+np.arange(t), 2*t+np.arange(t),
                           3*t+np.arange(t), 3*t+1+np.arange(t)]).ravel()
    eq = coo_matrix((np.tile([-ETA, 1/ETA, -1., 1.], t), (r, col)),
                    shape=(t, n)).tocsr()
    bounds = [(0., None)]*t+[(0., P_MAX)]*(2*t)+[(S_MIN, S_MAX)]*(t+1)+[(0., None)]*(k*t)
    return a, eq, bounds


def solve(prices, scenarios, weights, initial, terminal_value, terminal=None):
    scenarios = np.asarray(scenarios)
    k, t = scenarios.shape
    a, eq, original_bounds = matrix(k, t)
    bounds = original_bounds.copy()
    bounds[3*t] = (initial, initial)
    if terminal is not None:
        bounds[4*t] = (terminal, terminal)
    c = np.zeros(4*t+1+k*t)
    c[:t] = prices
    c[t:3*t] = 1e-7  # numerical tie-break, excluded from reported cash costs
    c[4*t] = -terminal_value
    c[4*t+1:] = (5*weights[:, None]*prices[None]).ravel()
    result = linprog(c, A_ub=a, b_ub=-scenarios.ravel(), A_eq=eq,
                     b_eq=np.zeros(t), bounds=bounds, method='highs')
    if not result.success:
        raise RuntimeError(result.message)
    g, charge, discharge = [np.maximum(result.x[i*t:(i+1)*t], 0) for i in range(3)]
    soc = result.x[3*t:4*t+1]
    assert np.max(np.abs(np.diff(soc)-ETA*charge+discharge/ETA)) < 1e-5
    assert np.max(np.minimum(charge, discharge)) < 1e-5
    return g, charge, discharge, soc


def scenario_set(history_net, history_pred, current_pred, window=56, half_life=14, limit=28):
    """API accepts only completed history; cannot index target-day truth."""
    d = len(history_net)
    assert history_pred.shape == history_net.shape
    indices = np.arange(max(START, d-window), d)
    if len(indices) == 0:
        return current_pred[None], np.ones(1), -1
    assert np.isfinite(history_pred[indices]).all()
    if len(indices) > limit:
        indices = indices[np.rint(np.linspace(0, len(indices)-1, limit)).astype(int)]
    weights = 2.**(-(d-indices)/half_life)
    weights /= weights.sum()
    return current_pred[None]+history_net[indices]-history_pred[indices], weights, int(indices[-1])


def weighted_quantile(values, weights, q):
    order = np.argsort(values, axis=0, kind='stable')
    sorted_values = np.take_along_axis(values, order, axis=0)
    cw = np.cumsum(np.broadcast_to(weights[:, None], values.shape)[order, np.arange(values.shape[1])], axis=0)
    index = np.argmax(cw >= q-1e-12, axis=0)
    return sorted_values[index, np.arange(values.shape[1])]


def dp_reserves(grid_purchase, scenarios, weights, prices, terminal_value, step=50.):
    """Bellman recursion for marginal independent demands; continuous execution.

    Charging uses available surplus only. On a deficit, convex interpolated
    continuation value gives a common reserve threshold, known at midnight.
    Full-day residual correlation is NOT represented in this DP approximation.
    """
    states = np.linspace(S_MIN, S_MAX, round((S_MAX-S_MIN)/step)+1)
    value = -terminal_value*states
    reserves = np.zeros(144)
    for t in range(143, -1, -1):
        reserve = states[np.argmin(value+5*prices[t]*ETA*states)]
        reserves[t] = reserve
        deficit = scenarios[:, t]-grid_purchase[t]
        s = states[None, :]
        h = deficit[:, None]
        minimum_end = np.maximum(S_MIN, s-np.minimum(np.maximum(h, 0), P_MAX)/ETA)
        end_deficit = np.maximum(minimum_end, np.minimum(s, reserve))
        end_surplus = np.minimum(S_MAX, s+ETA*np.minimum(np.maximum(-h, 0), P_MAX))
        end = np.where(h > 0, end_deficit, end_surplus)
        emergency = np.where(h > 0, np.maximum(h-ETA*(s-end), 0), 0)
        continuation = np.interp(end.ravel(), states, value).reshape(end.shape)
        value = weights @ (5*prices[t]*emergency+continuation)
    return reserves


def step_action(net_now, planned_now, state, reserve):
    deficit = net_now-planned_now
    if deficit > 0:
        discharge = min(deficit, P_MAX, max(0., ETA*(state-reserve)))
        charge, emergency, spill = 0., deficit-discharge, 0.
    else:
        charge = min(-deficit, P_MAX, max(0., (S_MAX-state)/ETA))
        discharge, emergency, spill = 0., 0., -deficit-charge
    return charge, discharge, emergency, spill, state+ETA*charge-discharge/ETA


def execute(net, grid, reserves, initial):
    flows = np.zeros((5, 144)); flows[0] = grid
    soc = np.zeros(145); soc[0] = initial
    for t in range(144):
        c, b, u, w, soc[t+1] = step_action(float(net[t]), grid[t], soc[t], reserves[t])
        flows[1:, t] = c, b, u, w
    assert np.max(np.abs(flows[0]+flows[2]+flows[3]-net-flows[1]-flows[4])) < 1e-6
    assert np.max(np.abs(np.diff(soc)-ETA*flows[1]+flows[2]/ETA)) < 1e-6
    assert soc.min() >= S_MIN-1e-6 and soc.max() <= S_MAX+1e-6
    assert np.max(flows[1:3]) <= P_MAX+1e-6
    assert np.min(flows) >= -1e-6
    assert np.max(np.minimum(flows[1], flows[2])) < 1e-8
    return flows, soc


def closed_loop_loss_kernel(offsets, base_grid, scale, reserves, scenarios, weights, initial, prices, value):
    """Equivalent scalar-loop kernel, optionally compiled with Numba."""
    grid = np.zeros(144); dg = np.zeros(144)
    gradient = np.zeros(6); cost = 0.
    for t in range(144):
        raw = base_grid[t]+offsets[t//24]*scale[t]
        grid[t] = max(0., raw)
        dg[t] = scale[t] if raw > 0 else 0.
        cost += prices[t]*grid[t]
        gradient[t//24] += prices[t]*dg[t]
    for i in range(len(scenarios)):
        state = float(initial)
        js = np.zeros(6)
        for t in range(144):
            b = t//24
            h = scenarios[i,t]-grid[t]
            if h <= 0:
                room = (S_MAX-state)/ETA
                if -h <= P_MAX and -h <= room:
                    charge = -h
                    js[b] += ETA*dg[t]
                elif P_MAX <= room:
                    charge = P_MAX
                else:
                    charge = max(0., room)
                    js[:] = 0.
                state += ETA*charge
            else:
                available = max(0., ETA*(state-reserves[t]))
                if h <= P_MAX and h <= available:
                    discharge = h
                    emergency = 0.
                    js[b] += dg[t]/ETA
                elif P_MAX <= available:
                    discharge = P_MAX
                    emergency = h-P_MAX
                    gradient[b] -= weights[i]*5*prices[t]*dg[t]
                else:
                    discharge = available
                    emergency = h-available
                    gradient[b] -= weights[i]*5*prices[t]*dg[t]
                    if state > reserves[t]:
                        for j in range(6):
                            gradient[j] -= weights[i]*5*prices[t]*ETA*js[j]
                            js[j] = 0.
                state -= discharge/ETA
                cost += weights[i]*5*prices[t]*emergency
        cost -= weights[i]*value*state
        for j in range(6):
            gradient[j] -= weights[i]*value*js[j]
    return cost, gradient


def optimize_grid(base_grid, scale, reserves, scenarios, weights, initial, prices, value):
    objective = lambda z: closed_loop_loss(z, base_grid, scale, reserves, scenarios,
                                           weights, initial, prices, value)
    zero = np.zeros(6)
    before = objective(zero)[0]
    fit = minimize(objective, zero, jac=True, method='L-BFGS-B', bounds=[(-2., 2.)]*6,
                   options=dict(maxiter=60, ftol=1e-9, gtol=1e-4, maxls=25))
    accepted = fit.x if np.isfinite(fit.fun) and fit.fun < before else zero
    grid = np.maximum(0, base_grid+accepted[BLOCK]*scale)
    return grid, dict(objective_before=before, objective_after=float(objective(accepted)[0]),
                     optimizer_success=bool(fit.success), iterations=int(fit.nit), offsets=accepted.tolist())




closed_loop_loss = njit(cache=True)(closed_loop_loss_kernel)


def make_main_plan(scenarios, weights, prices, initial):
    """正式选定的flat_closed_loop_dp分支，剔除未采用的研究候选。"""
    value = float(prices[:36].min() / ETA)
    q = weighted_quantile(scenarios, weights, .8)
    grid, _, _, _ = solve(prices, q[None], np.ones(1), initial, value)
    mean = weights @ scenarios
    scale = np.maximum(20., np.sqrt(weights @ ((scenarios - mean) ** 2)))
    reserves = dp_reserves(grid, scenarios, weights, prices, value)
    grid, details = optimize_grid(grid, scale, reserves, scenarios, weights, initial, prices, value)
    revised = dp_reserves(grid, scenarios, weights, prices, value)
    old_cost = closed_loop_loss(np.zeros(6), grid, scale, reserves, scenarios, weights, initial, prices, value)[0]
    new_cost = closed_loop_loss(np.zeros(6), grid, scale, revised, scenarios, weights, initial, prices, value)[0]
    if new_cost < old_cost:
        reserves = revised
    return grid, reserves, details


def replay(dates, load, pv, prices, component_predictions, output):
    """每日融合→0时固定购电与反馈规则→逐区间执行。SOC全年连续。"""
    prediction, logs = [], []
    for component, actual, cube in zip(['load','pv'],[load,pv],component_predictions):
        forecast, records = online_blend(actual,cube.transpose(1,0,2),dates)
        prediction.append(forecast)
        for record in records:
            day = dates.get_loc(record['date'])
            assert record['weight_last_day'] < day and record['monthly_choice_last_day'] < day
            weights = np.asarray(record['weights']).ravel()
            assert weights.min() >= -1e-10 and abs(weights.sum()-1) < 1e-10
            logs.append(dict(component=component,date=record['date'],window_days=record['window'],
                history_last_index=record['weight_last_day'],monthly_choice_last_index=record['monthly_choice_last_day'],
                **dict(zip(COMPONENT_NAMES,weights.tolist()))))
    prediction = np.asarray(prediction)
    actual_net, predicted_net = (load-pv)*DT, (prediction[0]-prediction[1])*DT
    flows = np.zeros((365,5,144)); soc = np.zeros((365,145)); reserves = np.full((365,144),np.nan)
    records = []; state = 6000.; started = time.time()
    for d in range(365):
        if d < START:
            # 附录给定1月1日电量6000；前15日闲置电池，1月16日起连续预热。
            grid = np.maximum(actual_net[d-1],0) if d else np.zeros(144)
            emergency = np.maximum(actual_net[d]-grid,0)
            flows[d] = np.stack([grid,np.zeros(144),np.zeros(144),emergency,grid+emergency-actual_net[d]])
            soc[d] = state; last = d-1; extra = {}
        else:
            scenarios, weights, last = scenario_set(actual_net[:d],predicted_net[:d],predicted_net[d])
            grid, rule, extra = make_main_plan(scenarios,weights,prices,state)
            reserves[d] = rule
            # 当天实际值只在不可改的grid和rule生成后进入执行，不参与当天0时规划。
            flows[d],soc[d] = execute(actual_net[d],grid,rule,state)
            assert last < d
        state = float(soc[d,-1]); f = flows[d]
        planned_cost, emergency_cost = float(prices@f[0]),float(5*prices@f[3])
        records.append(dict(date=str(dates[d].date()),formal=d>=FORMAL,planned_cost=planned_cost,
            emergency_cost=emergency_cost,total_cost=planned_cost+emergency_cost,
            inventory_adjusted_cost=planned_cost+emergency_cost+prices[:36].min()/ETA*(soc[d,0]-soc[d,-1]),
            grid_kwh=float(f[0].sum()),charge_kwh=float(f[1].sum()),discharge_kwh=float(f[2].sum()),
            emergency_kwh=float(f[3].sum()),spill_kwh=float(f[4].sum()),soc_start=float(soc[d,0]),
            soc_end=state,emergency_intervals=int((f[3]>1e-6).sum()),residual_latest_index=last,**extra))
        if d%35 == 0 or d == 364:
            print(f'[调度] {dates[d].date()}，累计{time.time()-started:.1f}秒。',flush=True)
    audit = dict(balance_max_kwh=float(np.abs(flows[:,0]+flows[:,2]+flows[:,3]-actual_net-flows[:,1]-flows[:,4]).max()),
        state_equation_max_kwh=float(np.abs(np.diff(soc,axis=1)-ETA*flows[:,1]+flows[:,2]/ETA).max()),
        cross_day_soc_max_kwh=float(np.abs(soc[1:,0]-soc[:-1,-1]).max()),
        soc_min_kwh=float(soc.min()),soc_max_kwh=float(soc.max()),
        charge_discharge_max_kwh=float(flows[:,1:3].max()),
        simultaneous_charge_discharge_max_kwh=float(np.minimum(flows[:,1],flows[:,2]).max()),
        full_replanning_days=350,formal_days=334)
    assert max(audit[k] for k in ['balance_max_kwh','state_equation_max_kwh','cross_day_soc_max_kwh']) < 1e-6
    assert S_MIN-1e-6 <= soc.min() and soc.max() <= S_MAX+1e-6
    assert flows.min() >= -1e-6 and flows[:,1:3].max() <= P_MAX+1e-6
    assert audit['simultaneous_charge_discharge_max_kwh'] < 1e-8
    result = dict(dates=dates.strftime('%Y-%m-%d').to_numpy(dtype='U10'),load_kw=load,pv_kw=pv,
        price_yuan_per_kwh=prices,predicted_load_kw=prediction[0],predicted_pv_kw=prediction[1],
        component_names=np.array(COMPONENT_NAMES),component_predictions_kw=component_predictions,
        flows=flows,flow_names=np.array(['grid','charge','discharge','emergency','spill']),soc=soc,reserves=reserves)
    atomic_npz(output/'计算明细.npz',**result)
    daily = pd.DataFrame(records)
    daily.to_csv(output/'每日计算结果.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame(logs).to_csv(output/'每日融合权重.csv',index=False,encoding='utf-8-sig')
    formal = daily[daily.formal]
    summary = {k:float(formal[k].sum()) for k in ['planned_cost','emergency_cost','total_cost','inventory_adjusted_cost',
        'grid_kwh','charge_kwh','discharge_kwh','emergency_kwh','spill_kwh']}
    summary.update(policy='flat_closed_loop_dp',days=334,initial_soc=float(soc[FORMAL,0]),final_soc=float(soc[-1,-1]),
        emergency_days=int((formal.emergency_intervals>0).sum()),emergency_intervals=int(formal.emergency_intervals.sum()),
        archived_total_cost_reference=EXPECTED_COST,difference_from_archived_total_cost=float(summary['total_cost']-EXPECTED_COST))
    write_json(output/'费用汇总.json',summary)
    write_json(output/'物理与因果检查.json',audit)
    return result,summary


# ==================== 六、按附件五原模板填写Excel ====================

def clock_label(slot):
    return f'{slot//6}:{slot%6*10:02d}'


def emergency_events(values):
    """将连续紧急区间合并。下游使用原高精度电量，显示舍入不参与求和。"""
    events = []; t = 0
    while t < 144:
        if values[t] <= 1e-6:
            t += 1
            continue
        start = t
        while t < 144 and values[t] > 1e-6:
            t += 1
        events.append((f'{clock_label(start)}-{clock_label(t)}',float(values[start:t].sum())))
    return events


def export_excel(template, destination, result):
    """复制模板的工作表、表头、列宽与行样式；展开省略号代表的日期。"""
    template, destination = Path(template).resolve(),Path(destination).resolve()
    if template == destination:
        raise ValueError('输出文件不能覆盖附件五原模板。')
    workbook = load_workbook(template)
    expected_sheets = ['计划购电量','充放电量','紧急购电量']
    if workbook.sheetnames != expected_sheets:
        raise ValueError('模板应包含且依次为：计划购电量、充放电量、紧急购电量。')
    plan,storage,emergency = [workbook[name] for name in expected_sheets]
    if plan.max_column != 147 or plan.max_row != 335:
        raise ValueError('计划购电量模板应为335行、147列。')
    if [storage.cell(1,c).value for c in range(1,7)] != ['日期','时间段','充电量','放电量','时刻','储电量']:
        raise ValueError('充放电量表头不符合附件五。')
    if [emergency.cell(1,c).value for c in range(1,4)] != ['日期','购电时间段','购电量']:
        raise ValueError('紧急购电量表头不符合附件五。')
    if str(plan.cell(1,2).value) != '0:10-0:20' or str(plan.cell(1,145).value) != '0:00-0:10+1':
        raise ValueError('需要附件五原始result2模板；不要使用改写过时间表头的旧结果表。')
    def prototypes(sheet,rows,cols):
        return [([dict(value=sheet.cell(r,c).value,style=copy.copy(sheet.cell(r,c)._style)) for c in range(1,cols+1)],
                 copy.copy(sheet.row_dimensions[r])) for r in rows]
    storage_rows = prototypes(storage,range(2,8),6)
    emergency_rows = prototypes(emergency,range(2,5),3)
    storage.delete_rows(2,storage.max_row-1)
    emergency.delete_rows(2,emergency.max_row-1)
    def copy_row(sheet,row,prototype):
        cells,dimensions = prototype
        dimensions = copy.copy(dimensions);dimensions.index = row
        sheet.row_dimensions[row] = dimensions
        for column,saved in enumerate(cells,1):
            cell = sheet.cell(row,column)
            cell._style = copy.copy(saved['style'])
            cell.value = saved['value']
    events_rows = []; emergency_row = 2
    for day in range(FORMAL,365):
        dt = pd.Timestamp(str(result['dates'][day])).to_pydatetime()
        row = day-FORMAL+2
        if pd.Timestamp(plan.cell(row,1).value) != pd.Timestamp(dt):
            raise ValueError('模板计划表日期错位：第'+str(row)+'行。')
        grid,charge,discharge,urgent,spill = result['flows'][day]
        for slot in range(144):
            plan.cell(row,slot+2,float(grid[slot])).number_format = '0.00'
        plan.cell(row,146,float(grid.sum())).number_format = '0.00'
        plan.cell(row,147,float(result['price_yuan_per_kwh']@grid)).number_format = '0.00'
        for block in range(6):
            storage_row = 2+(day-FORMAL)*6+block
            copy_row(storage,storage_row,storage_rows[block])
            storage.cell(storage_row,1).value = dt if block == 0 else None
            sl = slice(24*block,24*(block+1))
            storage.cell(storage_row,3,float(charge[sl].sum())).number_format = '0.00'
            storage.cell(storage_row,4,float(discharge[sl].sum())).number_format = '0.00'
            storage.cell(storage_row,6).value = float(result['soc'][day,0]) if block == 0 else (float(result['soc'][day,-1]) if block == 1 else None)
            if block < 2:
                storage.cell(storage_row,6).number_format = '0.00'
        events = emergency_events(urgent)
        if not events:
            events = [('无',0.)]
        count = max(3,len(events))  # 保留模板每日至少三行的日期分组形式
        for i in range(count):
            prototype = emergency_rows[0 if i == 0 else (2 if i == count-1 else 1)]
            copy_row(emergency,emergency_row,prototype)
            emergency.cell(emergency_row,1).value = dt if i == 0 else None
            emergency.cell(emergency_row,2).value = events[i][0] if i < len(events) else None
            emergency.cell(emergency_row,3).value = events[i][1] if i < len(events) else None
            emergency.cell(emergency_row,3).number_format = '0.00'
            emergency_row += 1
        for interval,quantity in events:
            events_rows.append(dict(date=str(dt.date()),interval=interval,emergency_kwh=quantity))
    destination.parent.mkdir(parents=True,exist_ok=True)
    workbook.save(destination)
    workbook.close()
    pd.DataFrame(events_rows).to_csv(destination.parent/'紧急购电事件.csv',index=False,encoding='utf-8-sig')
    mapping = [dict(column_index=slot+2,template_header=plan.cell(1,slot+2).value,
        model_interval=f'{clock_label(slot)}-{clock_label(slot+1)}',attachment2_endpoint=clock_label(slot+1)) for slot in range(144)]
    pd.DataFrame(mapping).to_csv(destination.parent/'模板与模型时间映射.csv',index=False,encoding='utf-8-sig')
    verify_excel(template,destination,result)


def verify_excel(template,destination,result):
    """重读保存的Excel，验证表头布局、日期分组、原精度数值与总量。"""
    original = load_workbook(template,data_only=False)
    workbook = load_workbook(destination,data_only=True)
    try:
        assert workbook.sheetnames == original.sheetnames
        for source,target in zip(original,workbook):
            assert [c.value for c in source[1]] == [c.value for c in target[1]]
            assert source.freeze_panes == target.freeze_panes
            assert list(source.merged_cells.ranges) == list(target.merged_cells.ranges)
            for key,dimension in source.column_dimensions.items():
                assert target.column_dimensions[key].width == dimension.width
        plan,storage,emergency = workbook.worksheets
        assert plan.max_row == 335 and plan.max_column == 147 and storage.max_row == 2005
        grid = np.array([[plan.cell(r,c).value for c in range(2,146)] for r in range(2,336)],float)
        assert np.abs(grid-result['flows'][FORMAL:,0]).max() < 1e-8
        for i,day in enumerate(range(FORMAL,365)):
            assert abs(plan.cell(i+2,146).value-grid[i].sum()) < 1e-7
            assert abs(plan.cell(i+2,147).value-result['price_yuan_per_kwh']@grid[i]) < 1e-7
            row = 2+6*i
            assert pd.Timestamp(storage.cell(row,1).value) == pd.Timestamp(str(result['dates'][day]))
            assert all(storage.cell(row+j,1).value is None for j in range(1,6))
            assert abs(storage.cell(row,6).value-result['soc'][day,0]) < 1e-7
            assert abs(storage.cell(row+1,6).value-result['soc'][day,-1]) < 1e-7
            for block in range(6):
                for column,channel in [(3,1),(4,2)]:
                    value = float(result['flows'][day,channel,block*24:(block+1)*24].sum())
                    assert abs(storage.cell(row+block,column).value-value) < 1e-7
        filled_dates = [row[0] for row in emergency.iter_rows(min_row=2,values_only=True) if row[0] is not None]
        assert pd.DatetimeIndex(filled_dates).equals(pd.date_range('2025-02-01','2025-12-31'))
        urgent_total = sum(float(row[2] or 0) for row in emergency.iter_rows(min_row=2,values_only=True))
        assert abs(urgent_total-float(result['flows'][FORMAL:,3].sum())) < 1e-6
        for sheet in workbook:
            assert not any(c.data_type == 'e' for row in sheet for c in row)
        write_json(destination.parent/'Excel核验.json',dict(sheets=workbook.sheetnames,
            plan_shape=[335,147],storage_shape=[2005,6],emergency_rows=emergency.max_row,
            template_headers_preserved=True,date_grouping_matches_template=True,
            all_plan_and_storage_values_reconciled=True,emergency_total_kwh=urgent_total,
            template_time_labels_preserved_by_request=True))
    finally:
        original.close();workbook.close()


# ==================== 七、命令行入口与缓存管理 ====================

def required_input(value,label):
    if not str(value).strip():
        raise ValueError(label+'路径尚未填写。请编辑脚本顶部相关路径，或使用命令行参数。')
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(label+'文件不存在，请检查配置或命令行参数。')
    return path


def array_digest(values):
    """只用于核对附件是否与留档相符，不作为模型特征。"""
    return hashlib.sha256(np.ascontiguousarray(values, dtype=np.float64).view(np.uint8)).hexdigest()


def load_archived_components(path, dates, load, pv, prices):
    """读取显式提供的样本外四分量；不读取最终购电计划、SOC或费用。"""
    saved = read_npz(path)
    needed = {'dates', 'component_names', 'component_predictions_kw', 'provenance_json',
              'load_sha256', 'pv_sha256', 'price_sha256'}
    if set(saved) != needed:
        raise ValueError('预测留档字段不符；请使用本次附带的原交付四分量预测.npz。')
    if list(saved['component_names']) != COMPONENT_NAMES:
        raise ValueError('预测留档的四分量名称或顺序不符。')
    if not np.array_equal(saved['dates'], dates.strftime('%Y-%m-%d').to_numpy(dtype='U10')):
        raise ValueError('预测留档日期与附件不符。')
    for name, values in [('load',load),('pv',pv),('price',prices)]:
        if str(saved[name+'_sha256'].item()) != array_digest(values):
            raise ValueError('预测留档与当前附件的'+name+'数据不同，不能混用；请选择retrain模式。')
    components = saved['component_predictions_kw']
    if components.shape != (2,4,365,144) or not np.isfinite(components[:,:,START:]).all():
        raise ValueError('预测留档维度或有效预测不完整。')
    if (components[:,:,START:] < 0).any():
        raise ValueError('预测留档包含负功率。')
    # 留档内的来源信息可能包含原机器路径或其他个人信息，因此不复制到新输出。
    provenance = dict(archive_file_sha256=digest(path),
                      archived_provenance_sha256=hashlib.sha256(
                          str(saved['provenance_json'].item()).encode('utf-8')).hexdigest(),
                      base_models_retrained_in_this_run=False,
                      validation_scope='本次校验输入匹配；逐日融合与调度重新计算。基础预测来源及因果性依赖原研究留档审计。')
    return components, provenance


def main():
    parser = argparse.ArgumentParser(description='第二问单文件：完整重训或原预测留档复算→闭环调度→附件五格式result2.xlsx。')
    parser.add_argument('--attachment1',default=ATTACHMENT1_PATH,help='附件1.xlsx路径')
    parser.add_argument('--attachment2',default=ATTACHMENT2_PATH,help='附件2.xlsx路径')
    parser.add_argument('--template',default=RESULT2_TEMPLATE_PATH,help='附件5中的原始result2.xlsx路径')
    parser.add_argument('--output-dir',default=OUTPUT_DIRECTORY,help='输出目录；留空时在脚本旁按模式建立输出目录')
    parser.add_argument('--mode',choices=['retrain','archive'],default=RUN_MODE,help='retrain从附件完整训练；archive读取明确提供的基础预测，重新融合和规划')
    parser.add_argument('--archive',default=None,help='仅archive模式：原交付四分量预测.npz路径；不填写时使用代码顶部配置')
    parser.add_argument('--workers',type=int,choices=[1,2],default=TRAINING_WORKERS,help='独立训练进程数')
    parser.add_argument('--threads',type=int,default=LIGHTGBM_THREADS,help='每进程LightGBM线程数')
    parser.add_argument('--end-index',type=int,default=365,help='retrain训练截止索引；默认365完成全年，小于365不导出结果表')
    args = parser.parse_args()
    started = time.time()
    attachment1 = required_input(args.attachment1,'附件1')
    attachment2 = required_input(args.attachment2,'附件2')
    template = required_input(args.template,'附件五result2模板')
    if args.mode == 'retrain' and args.archive is not None:
        raise ValueError('retrain模式不使用--archive；请移除此参数，或改用--mode archive。')
    archive_path = args.archive if args.archive is not None else ARCHIVED_PREDICTIONS_PATH
    archive = required_input(archive_path,'四分量预测留档') if args.mode == 'archive' else None
    if args.threads < 1 or not 16 <= args.end_index <= 365:
        raise ValueError('--threads至少为1，--end-index须在16—365之间。')
    if args.mode == 'archive' and args.end_index != 365:
        raise ValueError('archive模式复算完整年度，不使用--end-index短程训练参数。')
    default_folder = '第二问重训输出' if args.mode == 'retrain' else '第二问留档复算输出'
    output = Path(args.output_dir).expanduser().resolve() if str(args.output_dir).strip() else Path(__file__).resolve().parent/default_folder
    if (output/'result2.xlsx').resolve() in [attachment1,attachment2,template,archive]:
        raise ValueError('输出目录不能导致覆盖输入附件。')
    if (output/'运行来源与环境.json').exists():
        previous = json.loads((output/'运行来源与环境.json').read_text(encoding='utf-8'))
        if previous.get('run_mode',args.mode) != args.mode:
            raise ValueError('该输出目录用于另一种运行模式，请另设输出目录，避免混淆结果。')
    # 提前检查模板，避免完成长时间训练后才发现传入旧版结果表。
    workbook = load_workbook(template,read_only=True,data_only=False)
    try:
        if workbook.sheetnames != ['计划购电量','充放电量','紧急购电量']:
            raise ValueError('模板工作表名称或顺序与附件五不符。')
        plan = workbook.worksheets[0]
        if (plan.max_row,plan.max_column) != (335,147) or plan['B1'].value != '0:10-0:20' or plan['EO1'].value != '0:00-0:10+1':
            raise ValueError('请使用附件五原始空白result2.xlsx模板，保持其原表头。')
    finally:
        workbook.close()
    dates,load,pv,prices = read_inputs(attachment1,attachment2)
    output.mkdir(parents=True,exist_ok=True)
    packages = ['numpy','pandas','scipy','scikit-learn','lightgbm','numba','llvmlite','openpyxl','threadpoolctl']
    versions = {name:metadata.version(name) for name in packages}
    versions['python'] = sys.version
    signature = None
    if args.mode == 'archive':
        print('读取显式指定的原四分量样本外预测；本次不重训基础模型，随后完整重算融合、计划与执行。',flush=True)
        components,provenance = load_archived_components(archive,dates,load,pv,prices)
        write_json(output/'预测留档信息.json',provenance)
    else:
        fingerprint = hashlib.sha256()
        fingerprint.update(digest(__file__).encode());fingerprint.update(json.dumps(versions,sort_keys=True).encode())
        fingerprint.update(str(args.threads).encode())
        for x in [load,pv]:fingerprint.update(np.ascontiguousarray(x).view(np.uint8))
        signature = fingerprint.hexdigest()
        cache = output/'训练缓存';cache.mkdir(exist_ok=True)
        jobs = [(name,series,dates.strftime('%Y-%m-%d').to_numpy(dtype='U10'),str(cache/f'{name}.npz'),signature,args.threads,args.end_index)
                for name,series in [('load',load),('pv',pv)]]
        print('开始从原始附件训练；已完成的相同输入缓存会自动恢复。',flush=True)
        if args.workers == 1:
            for job in jobs:train_component(*job)
        else:
            with ProcessPoolExecutor(max_workers=2) as pool:
                futures = [pool.submit(train_component,*job) for job in jobs]
                for future in futures:future.result()
        if args.end_index < 365:
            print('阶段训练检查完成；尚未训练完整年度，因此不生成正式Excel。重新运行默认365可继续。',flush=True)
            return
        states = [read_npz(cache/f'{name}.npz') for name in ['load','pv']]
        assert all(s['done'][START:].all() and np.isfinite(s['predictions'][:,START:]).all() for s in states)
        components = np.stack([s['predictions'] for s in states])
        training_log = []
        for component,state in zip(['load','pv'],states):
            for day in range(START,365):
                month_first = day-dates[day].day+1
                choice_day = START if day < FORMAL else month_first
                history = validation_origins(choice_day,START,35,3)
                training_log.append(dict(component=component,date=str(dates[day].date()),max_training_index=day-1,
                    parameter_choice_index=choice_day,choice_history_last=int(history[-1]) if len(history) else -1,
                    lgb_config=state['selected'][day,0],ridge_config=state['selected'][day,1],
                    calendar_alpha=[.01,.1,1.][int(state['calendar_active'][day])]))
        pd.DataFrame(training_log).to_csv(output/'训练与选参记录.csv',index=False,encoding='utf-8-sig')
    with threadpool_limits(1):
        result,summary = replay(dates,load,pv,prices,components,output)
    summary.update(run_mode=args.mode,base_models_retrained_in_this_run=args.mode=='retrain')
    write_json(output/'费用汇总.json',summary)
    audit = json.loads((output/'物理与因果检查.json').read_text(encoding='utf-8'))
    audit.update(run_mode=args.mode,grid_fixed_all_day=True,forecast_weights_use_completed_history=True,
                 base_models_retrained_in_this_run=args.mode=='retrain',
                 base_prediction_audit='本文件逐日前缀训练' if args.mode=='retrain' else '沿用原研究留档审计，本次未重训基础模型')
    write_json(output/'物理与因果检查.json',audit)
    print('按附件五原始布局写入result2.xlsx；区间列按原顺序映射，详见时间映射CSV。',flush=True)
    export_excel(template,output/'result2.xlsx',result)
    input_files = [('attachment1',attachment1),('attachment2',attachment2),('template',template)]
    if archive is not None:
        input_files.append(('archive',archive))
    write_json(output/'运行来源与环境.json',dict(versions=versions,source_sha256=digest(__file__),
        input_sha256={label:digest(path) for label,path in input_files},
        run_mode=args.mode,base_models_retrained_in_this_run=args.mode=='retrain',
        training_signature=signature,workers=args.workers,threads=args.threads,
        no_embedded_predictions=True,no_external_project_modules=True,elapsed_seconds=time.time()-started))
    print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)
    print('result2.xlsx生成完成。',flush=True)


if __name__ == '__main__':
    from multiprocessing import freeze_support
    freeze_support()
    if hasattr(sys.stdout,'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    try:
        main()
    except (ValueError,FileNotFoundError,PermissionError) as exc:
        print('运行失败：'+str(exc),file=sys.stderr)
        raise SystemExit(2)
