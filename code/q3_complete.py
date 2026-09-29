#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""第三问：1301万元主方案和1309万元基准的独立Python实现。

仅依赖下列公开Python库和显式输入的数据文件，不导入旧项目模块。
retrain：只用原附件从零训练；archive：显式读取基础四分量预测，重算其余全部过程。
当前区间反馈、右端点时标及同年度开发选模限制详见环境配置与复现说明.md。
"""
from __future__ import annotations

# 填写四项路径，或通过命令行参数提供。留空不会搜索旧项目或偷偷读取缓存。
ATTACHMENT1_PATH = r""
ATTACHMENT2_PATH = r""
ATTACHMENT3_PATH = r""
RESULT3_TEMPLATE_PATH = r""
OUTPUT_DIRECTORY = ""
RUN_MODE = "retrain"
ARCHIVED_PREDICTIONS_PATH = r""
TRAINING_WORKERS = 2
LIGHTGBM_THREADS = 1

# 一、第二问共用的四分量训练、历史融合及0点闭环控制。
import argparse
from concurrent.futures import ProcessPoolExecutor
import copy
from dataclasses import dataclass
from datetime import date, datetime
from functools import lru_cache
import hashlib
import importlib.metadata as metadata
import json
import math
import os
from pathlib import Path
import sys
import time
from typing import Any

import lightgbm as lgb
import numpy as np
import pandas as pd
from numba import njit
from openpyxl import load_workbook
from scipy.ndimage import gaussian_filter1d
from scipy.optimize import linprog, minimize
from scipy.sparse import coo_matrix
from sklearn.linear_model import Ridge
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from threadpoolctl import threadpool_limits

PERIODS_PER_DAY = 144
FEATURE_START_INDEX = 14
START = 15
FORMAL = 31
DT = 1 / 6
RANDOM_STATE = 2026
ETA = 0.9
S_MIN, S_MAX, P_MAX = 1200., 10800., 5000. / 6
BLOCK = np.arange(144) // 24
COMPONENT_NAMES = ['legacy_daily_lgb', 'legacy_daily_ridge', 'periodic_lgb_l1', 'calendar_curve']


def write_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def atomic_npz(path, **data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + '.tmp')
    with temporary.open('wb') as stream:
        np.savez_compressed(stream, **data)
    os.replace(temporary, path)


def read_npz(path):
    with np.load(path, allow_pickle=False) as z:
        return {name: z[name].copy() for name in z.files}


def read_inputs(attachment1, attachment2):
    """只读取附件1电价与附件2两个工作表，并严格校验日期和144列。"""
    workbook = load_workbook(attachment2, read_only=True, data_only=True)
    arrays, dates, headers = [], [], []
    try:
        if len(workbook.worksheets) < 2:
            raise ValueError('附件2必须包含负载、光伏两个工作表。')
        for sheet in workbook.worksheets[:2]:
            rows = list(sheet.values)
            headers.append(rows[0][1:])
            records = [row for row in rows[1:] if row[0] is not None]
            dates.append(pd.DatetimeIndex([row[0] for row in records]))
            arrays.append(np.array([row[1:] for row in records], float))
    finally:
        workbook.close()
    expected = pd.date_range('2025-01-01', '2025-12-31')
    if not all(x.equals(expected) for x in dates):
        raise ValueError('附件2日期必须完整、依序覆盖2025年365天。')
    if headers[0] != headers[1] or any(x.shape != (365, 144) for x in arrays):
        raise ValueError('附件2的两个工作表必须使用相同的144个时间列。')
    if not all(np.isfinite(x).all() and (x >= 0).all() for x in arrays):
        raise ValueError('附件2存在空值、非数字或负功率。')
    workbook = load_workbook(attachment1, read_only=True, data_only=True)
    try:
        rows = [row for row in list(workbook.active.values)[1:] if row[0] is not None]
        prices = np.array([row[1] for row in rows], float)
    finally:
        workbook.close()
    def minutes(value):
        if hasattr(value, 'hour'):
            return value.hour * 60 + value.minute
        raw = str(value).replace('+1', '')
        hour, minute = map(int, raw.split(':')[:2])
        return hour * 60 + minute + (1440 if '+1' in str(value) else 0)
    expected_minutes = list(range(10, 1441, 10))
    if [minutes(row[0]) for row in rows] != expected_minutes or [minutes(x) for x in headers[0]] != expected_minutes:
        raise ValueError('附件时间列应从00:10依次到0:00+1，每10分钟一个值。')
    if prices.shape != (144,) or not np.isfinite(prices).all() or (prices <= 0).any():
        raise ValueError('附件1电价须为144个正数。')
    return expected, arrays[0], arrays[1], prices


# ==================== 二、两个历史模型：逐日训练、月初选参 ====================

@dataclass(frozen=True)
class ModelConfig:
    name: str
    family: str
    target_kind: str
    objective: str = "regression"
    n_estimators: int = 0
    learning_rate: float = 0.0
    num_leaves: int = 0
    min_child_samples: int = 0
    reg_alpha: float = 0.0
    reg_lambda: float = 0.0
    ridge_alpha: float = 0.0
    train_window_days: int | None = None


MODEL_CONFIGS: tuple[ModelConfig, ...] = (
    ModelConfig(
        "lgbm_level_l2_large",
        "lgbm",
        "level",
        objective="regression",
        n_estimators=180,
        learning_rate=0.040,
        num_leaves=63,
        min_child_samples=80,
        reg_alpha=0.20,
        reg_lambda=5.0,
    ),
    ModelConfig(
        "lgbm_level_l2_moretrees",
        "lgbm",
        "level",
        objective="regression",
        n_estimators=260,
        learning_rate=0.025,
        num_leaves=31,
        min_child_samples=60,
        reg_alpha=0.10,
        reg_lambda=2.0,
    ),
    ModelConfig(
        "lgbm_residual_l1_moretrees",
        "lgbm",
        "residual",
        objective="regression_l1",
        n_estimators=220,
        learning_rate=0.030,
        num_leaves=31,
        min_child_samples=60,
        reg_alpha=0.10,
        reg_lambda=2.0,
    ),
    ModelConfig("ridge_level_a1", "ridge", "level", ridge_alpha=1.0),
    ModelConfig("ridge_level_a10", "ridge", "level", ridge_alpha=10.0),
    ModelConfig("ridge_level_a100", "ridge", "level", ridge_alpha=100.0),
    ModelConfig("ridge_residual_a10", "ridge", "residual", ridge_alpha=10.0),
)


DEFAULT_CONFIG = {
    ("load", "lgbm"): "lgbm_level_l2_large",
    ("pv", "lgbm"): "lgbm_residual_l1_moretrees",
    ("load", "ridge"): "ridge_level_a10",
    ("pv", "ridge"): "ridge_residual_a10",
    ("load", "seasonal"): "seasonal_lag7",
    ("pv", "seasonal"): "seasonal_lag1",
}


def feature_row_from_history(history: np.ndarray, target_date: date) -> np.ndarray:
    """由目标日前的逐日曲线构造目标日特征，不读取目标日真实值。"""
    history = np.asarray(history, dtype=np.float64)
    if history.ndim != 2 or history.shape[1] != PERIODS_PER_DAY:
        raise ValueError("history 必须是 n×144 数组。")
    if len(history) < 14:
        raise ValueError("至少需要14个完整历史日。")

    slots = np.arange(PERIODS_PER_DAY, dtype=np.float64)
    slot_angle = 2.0 * np.pi * slots / PERIODS_PER_DAY
    lag1, lag2, lag3 = history[-1], history[-2], history[-3]
    lag7, lag14 = history[-7], history[-14]
    hist3, hist7, hist14 = history[-3:], history[-7:], history[-14:]
    hist28 = history[-min(28, len(history)) :]
    weekday = target_date.weekday()
    doy = target_date.timetuple().tm_yday
    weekday_angle = 2.0 * np.pi * weekday / 7.0
    doy_angle = 2.0 * np.pi * (doy - 1) / 365.25

    row = np.column_stack(
        [
            lag1,
            lag2,
            lag3,
            lag7,
            lag14,
            np.mean(hist3, axis=0),
            np.mean(hist7, axis=0),
            np.mean(hist14, axis=0),
            np.mean(hist28, axis=0),
            np.std(hist3, axis=0),
            np.std(hist7, axis=0),
            np.std(hist14, axis=0),
            np.std(hist28, axis=0),
            lag1 - lag2,
            lag7 - lag14,
            np.full(PERIODS_PER_DAY, np.mean(lag1)),
            np.full(PERIODS_PER_DAY, np.max(lag1)),
            np.full(PERIODS_PER_DAY, np.std(lag1)),
            np.full(PERIODS_PER_DAY, np.mean(lag7)),
            np.full(PERIODS_PER_DAY, np.max(lag7)),
            slots / (PERIODS_PER_DAY - 1),
            np.sin(slot_angle),
            np.cos(slot_angle),
            np.sin(2.0 * slot_angle),
            np.cos(2.0 * slot_angle),
            np.full(PERIODS_PER_DAY, weekday / 6.0),
            np.full(PERIODS_PER_DAY, float(weekday >= 5)),
            np.full(PERIODS_PER_DAY, math.sin(weekday_angle)),
            np.full(PERIODS_PER_DAY, math.cos(weekday_angle)),
            np.full(PERIODS_PER_DAY, (doy - 1) / 364.0),
            np.full(PERIODS_PER_DAY, math.sin(doy_angle)),
            np.full(PERIODS_PER_DAY, math.cos(doy_angle)),
            np.full(PERIODS_PER_DAY, target_date.month / 12.0),
        ]
    )
    if row.shape != (PERIODS_PER_DAY, 33) or not np.all(np.isfinite(row)):
        raise RuntimeError("特征行形状或数值异常。")
    return row.astype(np.float32)


def build_feature_cube(series: np.ndarray, dates: tuple[date, ...]) -> np.ndarray:
    cube = np.full((len(dates), PERIODS_PER_DAY, 33), np.nan, dtype=np.float32)
    for target_index in range(FEATURE_START_INDEX, len(dates)):
        cube[target_index] = feature_row_from_history(
            series[:target_index], dates[target_index]
        )
    if not np.all(np.isfinite(cube[FEATURE_START_INDEX:])):
        raise RuntimeError("有效预测日的特征中存在非有限值。")
    return cube


def make_model(config: ModelConfig, lgbm_threads: int) -> Any:
    if config.family == "lgbm":
        return lgb.LGBMRegressor(
            objective=config.objective,
            n_estimators=config.n_estimators,
            learning_rate=config.learning_rate,
            num_leaves=config.num_leaves,
            min_child_samples=config.min_child_samples,
            subsample=0.90,
            subsample_freq=1,
            colsample_bytree=0.90,
            reg_alpha=config.reg_alpha,
            reg_lambda=config.reg_lambda,
            max_bin=127,
            random_state=RANDOM_STATE,
            n_jobs=lgbm_threads,
            verbosity=-1,
            deterministic=True,
            force_col_wise=True,
        )
    if config.family == "ridge":
        return make_pipeline(
            StandardScaler(),
            Ridge(alpha=config.ridge_alpha, solver="cholesky"),
        )
    raise ValueError(f"未知模型族：{config.family}")


def fit_model_for_origin(
    series: np.ndarray,
    feature_cube: np.ndarray,
    config: ModelConfig,
    target_index: int,
    residual_lag: int,
    lgbm_threads: int,
) -> Any:
    train_start = FEATURE_START_INDEX
    if config.train_window_days is not None:
        train_start = max(train_start, target_index - config.train_window_days)
    train_days = np.arange(train_start, target_index, dtype=int)
    if len(train_days) == 0:
        raise RuntimeError(f"{target_index=} 没有有效训练日。")
    assert int(np.max(train_days)) < target_index
    x_train = feature_cube[train_days].reshape(-1, feature_cube.shape[-1])
    if config.target_kind == "level":
        y_train = series[train_days].reshape(-1)
    elif config.target_kind == "residual":
        if int(np.min(train_days)) - residual_lag < 0:
            raise RuntimeError("残差模型基线滞后越界。")
        y_train = (series[train_days] - series[train_days - residual_lag]).reshape(-1)
    else:
        raise ValueError(f"未知目标类型：{config.target_kind}")
    model = make_model(config, lgbm_threads)
    model.fit(x_train, y_train)
    return model


def predict_with_fitted_model(
    model: Any,
    config: ModelConfig,
    feature_row: np.ndarray,
    baseline: np.ndarray,
) -> np.ndarray:
    if config.family == "lgbm":
        raw = np.asarray(model.booster_.predict(feature_row), dtype=np.float64)
    else:
        raw = np.asarray(model.predict(feature_row), dtype=np.float64)
    if config.target_kind == "residual":
        raw = baseline + raw
    return np.clip(raw.reshape(PERIODS_PER_DAY), 0.0, None)


def fit_predict_origin(
    series: np.ndarray,
    dates: tuple[date, ...],
    feature_cube: np.ndarray,
    config: ModelConfig,
    target_index: int,
    residual_lag: int,
    lgbm_threads: int,
    include_next_day: bool,
) -> tuple[np.ndarray, np.ndarray | None]:
    """拟合 < target_index 的标签，预测 target_index 及可选递归下一日。"""
    model = fit_model_for_origin(
        series,
        feature_cube,
        config,
        target_index,
        residual_lag,
        lgbm_threads,
    )
    current_baseline = series[target_index - residual_lag]
    current = predict_with_fitted_model(
        model, config, feature_cube[target_index], current_baseline
    )
    if not include_next_day or target_index + 1 >= len(dates):
        return current, None

    augmented = np.vstack([series[:target_index], current[None, :]])
    next_feature = feature_row_from_history(augmented, dates[target_index + 1])
    next_baseline = augmented[-residual_lag]
    following = predict_with_fitted_model(model, config, next_feature, next_baseline)
    return current, following


def weighted_error_score(
    actual: np.ndarray,
    predicted: np.ndarray,
    origins: np.ndarray,
    asof_index: int,
    half_life_days: int,
) -> tuple[float, float, float]:
    if len(origins) == 0:
        return math.inf, math.inf, math.inf
    assert int(np.max(origins)) < asof_index
    error = actual[origins] - predicted[origins]
    if not np.all(np.isfinite(error)):
        return math.inf, math.inf, math.inf
    ages = asof_index - origins
    day_weights = np.power(0.5, ages / max(half_life_days, 1))
    point_weights = np.repeat(day_weights[:, None], PERIODS_PER_DAY, axis=1)
    denominator = float(np.sum(point_weights * np.abs(actual[origins])))
    wape = float(np.sum(point_weights * np.abs(error)) / max(denominator, 1.0))
    rmse = float(np.sqrt(np.sum(point_weights * error**2) / np.sum(point_weights)))
    scale = float(
        np.sum(point_weights * np.abs(actual[origins])) / np.sum(point_weights)
    )
    nrmse = rmse / max(scale, 1.0)
    score = 0.65 * wape + 0.35 * nrmse
    return score, wape, rmse


def validation_origins(
    target_index: int, oos_start_index: int, validation_days: int, stride: int
) -> np.ndarray:
    start = max(oos_start_index, target_index - validation_days)
    origins = list(range(start, target_index, stride))
    if target_index - 1 >= oos_start_index and (not origins or origins[-1] != target_index - 1):
        origins.append(target_index - 1)
    result = np.asarray(sorted(set(origins)), dtype=int)
    if len(result):
        assert int(np.max(result)) < target_index
    return result


def candidate_indices(family: str) -> list[int]:
    return [index for index, config in enumerate(MODEL_CONFIGS) if config.family == family]


def choose_trainable_config(
    component: str,
    family: str,
    series: np.ndarray,
    dates: tuple[date, ...],
    feature_cube: np.ndarray,
    hyper_predictions: np.ndarray,
    tuning_scores: np.ndarray,
    month_start_index: int,
    oos_start_index: int,
    residual_lag: int,
    args: argparse.Namespace,
) -> tuple[str, np.ndarray]:
    origins = validation_origins(
        month_start_index,
        oos_start_index,
        args.validation_days,
        args.tune_stride,
    )
    config_ids = candidate_indices(family)
    if len(origins) < args.minimum_tuning_origins:
        return DEFAULT_CONFIG[(component, family)], origins

    for config_id in config_ids:
        config = MODEL_CONFIGS[config_id]
        for origin in origins:
            if not np.all(np.isfinite(hyper_predictions[config_id, origin])):
                forecast, _ = fit_predict_origin(
                    series,
                    dates,
                    feature_cube,
                    config,
                    int(origin),
                    residual_lag,
                    args.lgbm_threads,
                    include_next_day=False,
                )
                hyper_predictions[config_id, origin] = forecast.astype(np.float32)
        score, _, _ = weighted_error_score(
            series,
            hyper_predictions[config_id],
            origins,
            month_start_index,
            args.tune_half_life,
        )
        tuning_scores[month_start_index, config_id] = score

    best_id = min(config_ids, key=lambda index: tuning_scores[month_start_index, index])
    return MODEL_CONFIGS[best_id].name, origins


def config_by_name(name: str) -> ModelConfig:
    return next(config for config in MODEL_CONFIGS if config.name == name)


def simplex_least_squares(
    x: np.ndarray, y: np.ndarray, regularization: float, prior: np.ndarray
) -> np.ndarray:
    component_count = x.shape[1]
    gram = (x.T @ x) / max(len(y), 1)
    cross = (x.T @ y) / max(len(y), 1)
    prior_error = float(np.mean((x @ prior - y) ** 2))
    penalty = regularization * max(prior_error, 1.0)
    gram = gram + (penalty + 1.0e-8) * np.eye(component_count)
    cross = cross + penalty * prior
    best: np.ndarray | None = None
    best_value = math.inf
    for mask in range(1, 1 << component_count):
        active = np.asarray(
            [index for index in range(component_count) if mask & (1 << index)], dtype=int
        )
        a = gram[np.ix_(active, active)]
        b = cross[active]
        ones = np.ones(len(active))
        kkt = np.block([[2.0 * a, ones[:, None]], [ones[None, :], np.zeros((1, 1))]])
        rhs = np.concatenate([2.0 * b, np.ones(1)])
        try:
            active_weight = np.linalg.solve(kkt, rhs)[:-1]
        except np.linalg.LinAlgError:
            active_weight = np.linalg.lstsq(kkt, rhs, rcond=None)[0][:-1]
        if np.min(active_weight) < -1.0e-8:
            continue
        active_weight = np.clip(active_weight, 0.0, None)
        active_weight /= np.sum(active_weight)
        weight = np.zeros(component_count)
        weight[active] = active_weight
        value = float(
            np.mean((x @ weight - y) ** 2)
            + penalty * np.sum((weight - prior) ** 2)
        )
        if value < best_value:
            best_value = value
            best = weight
    return prior.copy() if best is None else best



# ==================== 三、周期模型与单层融合 ====================

def periodic_features(history, target_date, component):
    """Only a prefix is accepted; all 144 target slots are built together."""
    x = feature_row_from_history(history, target_date.date()).astype(float)
    angle = 2*np.pi*(np.arange(144)+1)/144
    harmonics = np.column_stack([v for k in range(1, 5)
                               for v in (np.sin(k*angle), np.cos(k*angle))])
    smooth1 = gaussian_filter1d(history[-1], 2, mode='nearest')
    smooth7 = gaussian_filter1d(history[-7], 2, mode='nearest')
    smoothmean = gaussian_filter1d(history[-7:].mean(0), 2, mode='nearest')
    # Seasonal baseline is explicitly computed from past observations.
    base = (0.65*smooth7+0.35*smooth1 if component == 'load'
            else 0.65*smooth1+0.35*gaussian_filter1d(history[-3:].mean(0), 2))
    interactions = np.column_stack([harmonics * z[:, None]/1000
                                   for z in (smooth1, smooth7, smoothmean)])
    weekday = np.eye(7)[target_date.dayofweek]
    calendar = np.hstack([np.tile(weekday, (144, 1)),
                         np.einsum('tj,k->tjk', harmonics[:, :4], weekday).reshape(144, -1)])
    local = np.column_stack([smooth1, smooth7, smoothmean,
                             np.gradient(smooth1), np.gradient(smooth7),
                             np.maximum(history[-14:].max(0), 1)])
    return np.column_stack([x, harmonics, interactions, calendar, local]).astype('float32'), base


def fit_periodic_predict(x, y, baseline, dates, d, model_id, threads):
    train = np.arange(14, d)
    xx = x[train].reshape(-1, x.shape[-1])
    yy = (y[train] - baseline[train]).reshape(-1)
    half_life = 56 if model_id == 0 else 84
    sw = np.repeat(2.0 ** (-(d-1-train)/half_life), 144)
    if model_id == 0:
        scaler = StandardScaler().fit(xx, sample_weight=sw)
        model = Ridge(alpha=40, solver='cholesky')
        model.fit(scaler.transform(xx), yy, sample_weight=sw)
        pred = model.predict(scaler.transform(x[d]))
        pred = gaussian_filter1d(pred, 1, mode='nearest')
    else:
        model = lgb.LGBMRegressor(objective='regression' if model_id == 1 else 'regression_l1',
                    n_estimators=240, learning_rate=0.035, num_leaves=31,
                    min_child_samples=90, reg_lambda=8, reg_alpha=0.2,
                    max_bin=127, colsample_bytree=0.9, n_jobs=threads,
                    random_state=20260911, verbosity=-1, deterministic=True,
                    force_col_wise=True)
        model.fit(xx, yy, sample_weight=sw)
        pred = model.booster_.predict(x[d], num_threads=threads)
    return np.maximum(0, baseline[d]+pred)


def online_blend(actual, cube, dates, windows=(28, 56), groups=False):
    """Monthly choice among causal recipes, each using earlier OOS errors.

    groups=True uses six four-hour groups with shrinkage towards the global
    weights; this is a predeclared ablation, not selected with full-year data.
    """
    n, m, t = cube.shape
    pred = np.full((len(windows), n, t), np.nan)
    ws = np.full((len(windows), n, 6 if groups else 1, m), np.nan)
    prior = np.ones(m)/m
    for ri, window in enumerate(windows):
        for d in range(START, n):
            h = np.arange(max(START, d-window), d)
            w = prior
            if len(h) >= 7:
                xx = cube[h].transpose(0, 2, 1).reshape(-1, m)
                w = simplex_least_squares(xx, actual[h].reshape(-1), 0.15, prior)
            for group in range(6 if groups else 1):
                sl = slice(group*24, (group+1)*24) if groups else slice(0, 144)
                wg = w
                if groups and len(h) >= 14:
                    xx = cube[h, :, sl].transpose(0, 2, 1).reshape(-1, m)
                    local = simplex_least_squares(xx, actual[h, sl].reshape(-1), 0.3, w)
                    wg = 0.5*w+0.5*local
                ws[ri, d, group] = wg
                pred[ri, d, sl] = wg @ cube[d, :, sl]
    out = np.full_like(actual, np.nan)
    records = []
    active, cutoff = 0, -1
    for d in range(START, n):
        if d == START or dates[d].day == 1:
            h = np.arange(max(START, d-56), d)
            if len(h) >= 7:
                loss = np.mean((pred[:, h] - actual[h][None])**2, axis=(1, 2))
                active = int(loss.argmin())
            cutoff = d-1
        out[d] = pred[active, d]
        records.append(dict(date=str(dates[d].date()), window=windows[active],
                            monthly_choice_last_day=cutoff, weight_last_day=d-1,
                            weights=ws[active, d].tolist()))
    return out, records



# ==================== 四、日历曲线与完整基础训练 ====================

def build_calendar(dates, component):
    a=2*np.pi*(np.asarray(dates.dayofyear)-1)/365.25
    annual=np.column_stack([np.sin(a),np.cos(a),np.sin(2*a),np.cos(2*a)])
    if component=='pv':
        return annual
    week=np.eye(7)[dates.dayofweek]
    return np.column_stack([annual,week,
                           np.einsum('ij,ik->ijk',week,annual[:,:2]).reshape(len(dates),-1)])




def train_component(component, series, date_strings, cache_path, signature, threads=1, end_index=365):
    """从原始功率训练四个分量。每次只拟合目标日前的标签。

    仅复现正式方案所用的模型；未选入的旧季节融合、周期L2等不必运行。
    旧LightGBM/Ridge内部月度选参的候选池必须保留，不能用全年最终配置替代。
    """
    dates = pd.DatetimeIndex(date_strings)
    date_tuple = tuple(x.date() for x in dates)
    args = argparse.Namespace(validation_days=35, tune_stride=3, tune_half_life=21,
                              minimum_tuning_origins=5, lgbm_threads=threads)
    path = Path(cache_path)
    state = dict(predictions=np.full((4,365,144),np.nan),
        hyper=np.full((7,365,144),np.nan,dtype=np.float32),
        scores=np.full((365,7),np.nan), selected=np.full((365,2),'',dtype='U64'),
        calendar_bank=np.full((3,365,144),np.nan), calendar_active=np.full(365,-1,dtype=np.int16),
        done=np.zeros(365,dtype=bool))
    if path.exists():
        saved = read_npz(path)
        if str(saved.pop('signature').item()) != signature:
            raise ValueError('缓存与当前源码、输入或环境不符，请使用新的输出目录。')
        for key in state:
            if key not in saved or saved[key].shape != state[key].shape:
                raise ValueError('缓存字段或形状错误：' + key)
            state[key] = saved[key]
        completed = np.flatnonzero(state['done'])
        if len(completed) and not np.array_equal(completed,np.arange(START,completed[-1]+1)):
            raise ValueError('缓存完成日不连续。')
    print(f'[{component}] 已完成{int(state["done"].sum())}个训练日，准备特征。',flush=True)
    basic = build_feature_cube(series,date_tuple)
    feature_rows, base_rows = zip(*[periodic_features(series[:d],dates[d],component) for d in range(14,365)])
    periodic = np.full((365,144,feature_rows[0].shape[-1]),np.nan,dtype=np.float32)
    baseline = np.full_like(series,np.nan)
    periodic[14:],baseline[14:] = np.asarray(feature_rows),np.asarray(base_rows)
    calendar = build_calendar(dates,component)
    residual_lag = 7 if component == 'load' else 1
    active = [DEFAULT_CONFIG[(component,'lgbm')],DEFAULT_CONFIG[(component,'ridge')]]
    calendar_active = 1
    started = time.time()
    try:
        with threadpool_limits(limits=1, user_api='blas'):
            for d in range(START,end_index):
                if state['done'][d]:
                    active = state['selected'][d].tolist()
                    calendar_active = int(state['calendar_active'][d])
                    continue
                if d == START or dates[d].day == 1:
                    for family_index,family in enumerate(['lgbm','ridge']):
                        active[family_index],_ = choose_trainable_config(component,family,series,date_tuple,basic,
                            state['hyper'],state['scores'],d,START,residual_lag,args)
                    past = np.arange(max(START,d-56),d)
                    if len(past) >= 7:
                        calendar_active = int(np.mean((state['calendar_bank'][:,past]-series[past][None])**2,axis=(1,2)).argmin())
                for family_index,name in enumerate(active):
                    config = config_by_name(name)
                    current,_ = fit_predict_origin(series,date_tuple,basic,config,d,residual_lag,threads,False)
                    rounded = current.astype(np.float32)  # 原方案历史预测库以float32保存这两个分量
                    state['hyper'][MODEL_CONFIGS.index(config),d] = rounded
                    state['predictions'][family_index,d] = rounded
                    state['selected'][d,family_index] = name
                state['predictions'][2,d] = fit_periodic_predict(periodic,series,baseline,dates,d,2,threads)
                weights = 2.**(-(d-1-np.arange(d))/84.)
                for i,alpha in enumerate([.01,.1,1.]):
                    model = Ridge(alpha=alpha,solver='cholesky')
                    model.fit(calendar[:d],series[:d],sample_weight=weights)
                    raw = model.predict(calendar[d:d+1])[0]
                    state['calendar_bank'][i,d] = np.maximum(0,gaussian_filter1d(raw,1,mode='nearest'))
                state['predictions'][3,d] = state['calendar_bank'][calendar_active,d]
                state['calendar_active'][d] = calendar_active
                state['done'][d] = True
                if d % 7 == 0 or dates[d].day == 1 or d == end_index-1:
                    atomic_npz(path,signature=np.asarray(signature),**state)
                    print(f'[{component}] {dates[d].date()}，累计{time.time()-started:.1f}秒，已保存缓存。',flush=True)
    except BaseException:
        atomic_npz(path,signature=np.asarray(signature),**state)
        raise
    atomic_npz(path,signature=np.asarray(signature),**state)
    return str(path)


# ==================== 五、历史情景与电池反馈控制 ====================

@lru_cache(maxsize=30)
def matrix(k, t):
    n = 4*t+1+k*t
    ii = np.tile(np.arange(t), k)
    r = np.repeat(np.arange(k*t), 4)
    col = np.column_stack([ii, t+ii, 2*t+ii, 4*t+1+np.arange(k*t)]).ravel()
    a = coo_matrix((np.tile([-1., 1., -1., -1.], k*t), (r, col)),
                   shape=(k*t, n)).tocsr()
    r = np.repeat(np.arange(t), 4)
    col = np.column_stack([t+np.arange(t), 2*t+np.arange(t),
                           3*t+np.arange(t), 3*t+1+np.arange(t)]).ravel()
    eq = coo_matrix((np.tile([-ETA, 1/ETA, -1., 1.], t), (r, col)),
                    shape=(t, n)).tocsr()
    bounds = [(0., None)]*t+[(0., P_MAX)]*(2*t)+[(S_MIN, S_MAX)]*(t+1)+[(0., None)]*(k*t)
    return a, eq, bounds


def solve(prices, scenarios, weights, initial, terminal_value, terminal=None):
    scenarios = np.asarray(scenarios)
    k, t = scenarios.shape
    a, eq, original_bounds = matrix(k, t)
    bounds = original_bounds.copy()
    bounds[3*t] = (initial, initial)
    if terminal is not None:
        bounds[4*t] = (terminal, terminal)
    c = np.zeros(4*t+1+k*t)
    c[:t] = prices
    c[t:3*t] = 1e-7  # numerical tie-break, excluded from reported cash costs
    c[4*t] = -terminal_value
    c[4*t+1:] = (5*weights[:, None]*prices[None]).ravel()
    result = linprog(c, A_ub=a, b_ub=-scenarios.ravel(), A_eq=eq,
                     b_eq=np.zeros(t), bounds=bounds, method='highs')
    if not result.success:
        raise RuntimeError(result.message)
    g, charge, discharge = [np.maximum(result.x[i*t:(i+1)*t], 0) for i in range(3)]
    soc = result.x[3*t:4*t+1]
    assert np.max(np.abs(np.diff(soc)-ETA*charge+discharge/ETA)) < 1e-5
    assert np.max(np.minimum(charge, discharge)) < 1e-5
    return g, charge, discharge, soc


def scenario_set(history_net, history_pred, current_pred, window=56, half_life=14, limit=28):
    """API accepts only completed history; cannot index target-day truth."""
    d = len(history_net)
    assert history_pred.shape == history_net.shape
    indices = np.arange(max(START, d-window), d)
    if len(indices) == 0:
        return current_pred[None], np.ones(1), -1
    assert np.isfinite(history_pred[indices]).all()
    if len(indices) > limit:
        indices = indices[np.rint(np.linspace(0, len(indices)-1, limit)).astype(int)]
    weights = 2.**(-(d-indices)/half_life)
    weights /= weights.sum()
    return current_pred[None]+history_net[indices]-history_pred[indices], weights, int(indices[-1])


def weighted_quantile(values, weights, q):
    order = np.argsort(values, axis=0, kind='stable')
    sorted_values = np.take_along_axis(values, order, axis=0)
    cw = np.cumsum(np.broadcast_to(weights[:, None], values.shape)[order, np.arange(values.shape[1])], axis=0)
    index = np.argmax(cw >= q-1e-12, axis=0)
    return sorted_values[index, np.arange(values.shape[1])]


def q2_dp_reserves(grid_purchase, scenarios, weights, prices, terminal_value, step=50.):
    """Bellman recursion for marginal independent demands; continuous execution.

    Charging uses available surplus only. On a deficit, convex interpolated
    continuation value gives a common reserve threshold, known at midnight.
    Full-day residual correlation is NOT represented in this DP approximation.
    """
    states = np.linspace(S_MIN, S_MAX, round((S_MAX-S_MIN)/step)+1)
    value = -terminal_value*states
    reserves = np.zeros(144)
    for t in range(143, -1, -1):
        reserve = states[np.argmin(value+5*prices[t]*ETA*states)]
        reserves[t] = reserve
        deficit = scenarios[:, t]-grid_purchase[t]
        s = states[None, :]
        h = deficit[:, None]
        minimum_end = np.maximum(S_MIN, s-np.minimum(np.maximum(h, 0), P_MAX)/ETA)
        end_deficit = np.maximum(minimum_end, np.minimum(s, reserve))
        end_surplus = np.minimum(S_MAX, s+ETA*np.minimum(np.maximum(-h, 0), P_MAX))
        end = np.where(h > 0, end_deficit, end_surplus)
        emergency = np.where(h > 0, np.maximum(h-ETA*(s-end), 0), 0)
        continuation = np.interp(end.ravel(), states, value).reshape(end.shape)
        value = weights @ (5*prices[t]*emergency+continuation)
    return reserves


def step_action(net_now, planned_now, state, reserve):
    deficit = net_now-planned_now
    if deficit > 0:
        discharge = min(deficit, P_MAX, max(0., ETA*(state-reserve)))
        charge, emergency, spill = 0., deficit-discharge, 0.
    else:
        charge = min(-deficit, P_MAX, max(0., (S_MAX-state)/ETA))
        discharge, emergency, spill = 0., 0., -deficit-charge
    return charge, discharge, emergency, spill, state+ETA*charge-discharge/ETA


def execute(net, grid, reserves, initial):
    flows = np.zeros((5, 144)); flows[0] = grid
    soc = np.zeros(145); soc[0] = initial
    for t in range(144):
        c, b, u, w, soc[t+1] = step_action(float(net[t]), grid[t], soc[t], reserves[t])
        flows[1:, t] = c, b, u, w
    assert np.max(np.abs(flows[0]+flows[2]+flows[3]-net-flows[1]-flows[4])) < 1e-6
    assert np.max(np.abs(np.diff(soc)-ETA*flows[1]+flows[2]/ETA)) < 1e-6
    assert soc.min() >= S_MIN-1e-6 and soc.max() <= S_MAX+1e-6
    assert np.max(flows[1:3]) <= P_MAX+1e-6
    assert np.min(flows) >= -1e-6
    assert np.max(np.minimum(flows[1], flows[2])) < 1e-8
    return flows, soc


def closed_loop_loss_kernel(offsets, base_grid, scale, reserves, scenarios, weights, initial, prices, value):
    """Equivalent scalar-loop kernel, optionally compiled with Numba."""
    grid = np.zeros(144); dg = np.zeros(144)
    gradient = np.zeros(6); cost = 0.
    for t in range(144):
        raw = base_grid[t]+offsets[t//24]*scale[t]
        grid[t] = max(0., raw)
        dg[t] = scale[t] if raw > 0 else 0.
        cost += prices[t]*grid[t]
        gradient[t//24] += prices[t]*dg[t]
    for i in range(len(scenarios)):
        state = float(initial)
        js = np.zeros(6)
        for t in range(144):
            b = t//24
            h = scenarios[i,t]-grid[t]
            if h <= 0:
                room = (S_MAX-state)/ETA
                if -h <= P_MAX and -h <= room:
                    charge = -h
                    js[b] += ETA*dg[t]
                elif P_MAX <= room:
                    charge = P_MAX
                else:
                    charge = max(0., room)
                    js[:] = 0.
                state += ETA*charge
            else:
                available = max(0., ETA*(state-reserves[t]))
                if h <= P_MAX and h <= available:
                    discharge = h
                    emergency = 0.
                    js[b] += dg[t]/ETA
                elif P_MAX <= available:
                    discharge = P_MAX
                    emergency = h-P_MAX
                    gradient[b] -= weights[i]*5*prices[t]*dg[t]
                else:
                    discharge = available
                    emergency = h-available
                    gradient[b] -= weights[i]*5*prices[t]*dg[t]
                    if state > reserves[t]:
                        for j in range(6):
                            gradient[j] -= weights[i]*5*prices[t]*ETA*js[j]
                            js[j] = 0.
                state -= discharge/ETA
                cost += weights[i]*5*prices[t]*emergency
        cost -= weights[i]*value*state
        for j in range(6):
            gradient[j] -= weights[i]*value*js[j]
    return cost, gradient


def optimize_grid(base_grid, scale, reserves, scenarios, weights, initial, prices, value):
    objective = lambda z: q2_closed_loop_loss(z, base_grid, scale, reserves, scenarios,
                                           weights, initial, prices, value)
    zero = np.zeros(6)
    before = objective(zero)[0]
    fit = minimize(objective, zero, jac=True, method='L-BFGS-B', bounds=[(-2., 2.)]*6,
                   options=dict(maxiter=60, ftol=1e-9, gtol=1e-4, maxls=25))
    accepted = fit.x if np.isfinite(fit.fun) and fit.fun < before else zero
    grid = np.maximum(0, base_grid+accepted[BLOCK]*scale)
    return grid, dict(objective_before=before, objective_after=float(objective(accepted)[0]),
                     optimizer_success=bool(fit.success), iterations=int(fit.nit), offsets=accepted.tolist())




q2_closed_loop_loss = njit(cache=True)(closed_loop_loss_kernel)


def make_main_plan(scenarios, weights, prices, initial):
    """正式选定的flat_closed_loop_dp分支，剔除未采用的研究候选。"""
    value = float(prices[:36].min() / ETA)
    q = weighted_quantile(scenarios, weights, .8)
    grid, _, _, _ = solve(prices, q[None], np.ones(1), initial, value)
    mean = weights @ scenarios
    scale = np.maximum(20., np.sqrt(weights @ ((scenarios - mean) ** 2)))
    reserves = q2_dp_reserves(grid, scenarios, weights, prices, value)
    grid, details = optimize_grid(grid, scale, reserves, scenarios, weights, initial, prices, value)
    revised = q2_dp_reserves(grid, scenarios, weights, prices, value)
    old_cost = q2_closed_loop_loss(np.zeros(6), grid, scale, reserves, scenarios, weights, initial, prices, value)[0]
    new_cost = q2_closed_loop_loss(np.zeros(6), grid, scale, revised, scenarios, weights, initial, prices, value)[0]
    if new_cost < old_cost:
        reserves = revised
    return grid, reserves, details




# 二、第三问日内改约的结算、动态规划储备及闭环优化。

def settlement(grid,original,prices):
    up=np.maximum(grid-original,0);down=np.maximum(original-grid,0)
    return float(prices@original+1.5*prices@up-.5*prices@down)


def dp_reserves(grid,scenarios,weights,prices,value,step=50.):
    states=np.linspace(S_MIN,S_MAX,round((S_MAX-S_MIN)/step)+1)
    continuation=-value*states
    reserves=np.zeros(len(grid))
    for t in range(len(grid)-1,-1,-1):
        reserve=states[np.argmin(continuation+5*prices[t]*ETA*states)]
        reserves[t]=reserve
        s=states[None,:];h=(scenarios[:,t]-grid[t])[:,None]
        minimum=np.maximum(S_MIN,s-np.minimum(np.maximum(h,0),P_MAX)/ETA)
        deficit_end=np.maximum(minimum,np.minimum(s,reserve))
        surplus_end=np.minimum(S_MAX,s+ETA*np.minimum(np.maximum(-h,0),P_MAX))
        end=np.where(h>0,deficit_end,surplus_end)
        emergency=np.where(h>0,np.maximum(h-ETA*(s-end),0),0)
        future=np.interp(end.ravel(),states,continuation).reshape(end.shape)
        continuation=weights@(5*prices[t]*emergency+future)
    return reserves


@lru_cache(maxsize=8)
def lp_matrices(t):
    # [A,C,D,S(0..T),U,up,down], 7T+1 variables.
    k=np.arange(t);n=7*t+1
    a=coo_matrix((np.tile([-1.,1.,-1.,-1.],t),
        (np.repeat(k,4),np.column_stack([k,t+k,2*t+k,4*t+1+k]).ravel())),shape=(t,n)).tocsr()
    rows=np.r_[np.repeat(k,4),np.repeat(t+k,3)]
    cols=np.r_[np.column_stack([t+k,2*t+k,3*t+k,3*t+1+k]).ravel(),
        np.column_stack([k,5*t+1+k,6*t+1+k]).ravel()]
    vals=np.r_[np.tile([-ETA,1/ETA,-1.,1.],t),np.tile([1.,-1.,1.],t)]
    eq=coo_matrix((vals,(rows,cols)),shape=(2*t,n)).tocsr()
    bounds=[(0,None)]*t+[(0,P_MAX)]*(2*t)+[(S_MIN,S_MAX)]*(t+1)+[(0,None)]*(3*t)
    return a,eq,bounds


def initial_adjustment_plan(quantile,original,prices,initial,value):
    t=len(prices);a,eq,bounds0=lp_matrices(t);bounds=bounds0.copy()
    bounds[3*t]=(initial,initial)
    c=np.zeros(7*t+1);c[t:3*t]=1e-7;c[4*t]=-value
    c[4*t+1:5*t+1]=5*prices
    c[5*t+1:6*t+1]=1.5*prices;c[6*t+1:]=-.5*prices
    result=linprog(c,A_ub=a,b_ub=-quantile,A_eq=eq,b_eq=np.r_[np.zeros(t),original],bounds=bounds,method='highs')
    if not result.success:raise RuntimeError(result.message)
    assert np.minimum(result.x[5*t+1:6*t+1],result.x[6*t+1:]).max()<1e-6
    assert np.minimum(result.x[t:2*t],result.x[2*t:3*t]).max()<1e-5
    return np.maximum(0,result.x[:t])


@njit(cache=True)
def closed_loop_loss(offsets,base_grid,scale,reserves,scenarios,weights,initial,prices,value,original,blocks):
    t_count=len(prices);n=len(offsets)
    grid=np.zeros(t_count);dg=np.zeros(t_count);gradient=np.zeros(n);cost=0.
    for t in range(t_count):
        b=blocks[t];raw=base_grid[t]+offsets[b]*scale[t]
        grid[t]=max(0.,raw);dg[t]=scale[t] if raw>0 else 0.
        delta=grid[t]-original[t]
        cost+=prices[t]*original[t]+prices[t]*(1.5*max(delta,0.)-.5*max(-delta,0.))
        slope=1.5 if delta>0 else (.5 if delta<0 else 1.)
        gradient[b]+=slope*prices[t]*dg[t]
    for i in range(len(scenarios)):
        state=float(initial);jac_state=np.zeros(n)
        for t in range(t_count):
            b=blocks[t];h=scenarios[i,t]-grid[t]
            if h<=0:
                room=(S_MAX-state)/ETA
                if -h<=P_MAX and -h<=room:
                    charge=-h;jac_state[b]+=ETA*dg[t]
                elif P_MAX<=room:charge=P_MAX
                else:
                    charge=max(0.,room);jac_state[:]=0.
                state+=ETA*charge
            else:
                available=max(0.,ETA*(state-reserves[t]))
                if h<=P_MAX and h<=available:
                    discharge=h;emergency=0.;jac_state[b]+=dg[t]/ETA
                elif P_MAX<=available:
                    discharge=P_MAX;emergency=h-P_MAX
                    gradient[b]-=weights[i]*5*prices[t]*dg[t]
                else:
                    discharge=available;emergency=h-available
                    gradient[b]-=weights[i]*5*prices[t]*dg[t]
                    if state>reserves[t]:
                        for j in range(n):
                            gradient[j]-=weights[i]*5*prices[t]*ETA*jac_state[j]
                            jac_state[j]=0.
                state-=discharge/ETA
                cost+=weights[i]*5*prices[t]*emergency
        cost-=weights[i]*value*state
        gradient-=weights[i]*value*jac_state
    return cost,gradient


def make_adjustment(scenarios,weights,prices,initial,original,start_slot,value,weighted_quantile):
    grid=initial_adjustment_plan(weighted_quantile(scenarios,weights,.8),original,prices,initial,value)
    reserves=dp_reserves(grid,scenarios,weights,prices,value)
    mean=weights@scenarios
    scale=np.maximum(20.,np.sqrt(weights@((scenarios-mean)**2)))
    blocks=(np.arange(start_slot,144)//24-start_slot//24).astype(np.int64)
    assert len(blocks)==len(prices)
    n=int(blocks.max()+1);zero=np.zeros(n)
    objective=lambda z:closed_loop_loss(z,grid,scale,reserves,scenarios,weights,initial,prices,value,original,blocks)
    before=float(objective(zero)[0])
    fit=minimize(objective,zero,jac=True,method='L-BFGS-B',bounds=[(-2.,2.)]*n,
        options=dict(maxiter=60,ftol=1e-9,gtol=1e-4,maxls=25))
    z=fit.x if np.isfinite(fit.fun) and fit.fun<before else zero
    corrected=np.maximum(0,grid+z[blocks]*scale)
    changed=dp_reserves(corrected,scenarios,weights,prices,value)
    args=(scenarios,weights,initial,prices,value,original,blocks)
    old_loss=closed_loop_loss(zero,corrected,scale,reserves,*args)[0]
    new_loss=closed_loop_loss(zero,corrected,scale,changed,*args)[0]
    if new_loss<old_loss:reserves=changed
    after=float(closed_loop_loss(zero,corrected,scale,reserves,*args)[0])
    assert after<=before+1e-6
    return corrected,reserves,dict(objective_before=before,objective_after=after,
        optimizer_success=bool(fit.success),iterations=int(fit.nit),offsets=z.tolist())

# 三、已发布光伏预报、日内负载校正和条件情景。

def interpolate_release(day,release,raw,actual_pv):
    begin=release*36
    anchor=actual_pv[day,begin-1] if begin else (actual_pv[day-1,-1] if day else 0.)
    hours=release*6+np.arange(25)
    values=np.r_[anchor,raw[day,release]]
    endpoints=(np.arange(begin,144)+1)/6
    return np.maximum(0,np.interp(endpoints,hours,values))


def calibration_weight(day,release,actual_pv,base_pv,external):
    begin=release*36
    ids=np.arange(max(START,day-56),day)
    if len(ids)<7:return .5,len(ids),int(ids[-1]) if len(ids) else -1
    delta=external[ids,release,begin:]-base_pv[ids,begin:]
    error=actual_pv[ids,begin:]-base_pv[ids,begin:]
    weights=2.**(-(day-ids)/14)
    denominator=float(np.sum(weights[:,None]*delta*delta))
    alpha=float(np.clip(np.sum(weights[:,None]*delta*error)/denominator,0.,1.)) if denominator>1e-8 else 0.
    return alpha,len(ids),int(ids[-1])

def observed_features(day,release,load,base_load,pv):
    begin=36*release
    if not begin:return np.zeros(3)
    error=load[day,:begin]-base_load[day,:begin]
    revision=np.mean(pv[day,release,begin:begin+36]-pv[day,0,begin:begin+36])
    return np.array([error[-6:].mean(),error.mean(),revision])


def predict_load(day,release,actual,base):
    begin=36*release;result=base[day].copy();records=[]
    if not begin:return result,records
    ids=np.arange(max(15,day-56),day)
    if len(ids)<7:return result,records
    errors=actual[:day]-base[:day]
    x=np.column_stack([errors[ids,begin-6:begin].mean(axis=1),errors[ids,:begin].mean(axis=1)])
    now=np.array([(actual[day,begin-6:begin]-base[day,begin-6:begin]).mean(),
        (actual[day,:begin]-base[day,:begin]).mean()])
    weights=2.**(-(day-ids)/14)
    mean=np.average(x,axis=0,weights=weights)
    scale=np.maximum(10.,np.sqrt(np.average((x-mean)**2,axis=0,weights=weights)))
    design=np.column_stack([np.ones(len(ids)),(x-mean)/scale])
    vector=np.r_[1.,(now-mean)/scale]
    lhs=design.T@(weights[:,None]*design)+10*np.eye(3)
    for block in range(release,4):
        target=errors[ids,36*block:36*(block+1)].mean(axis=1)
        coef=np.linalg.solve(lhs,design.T@(weights*target))
        correction=float(vector@coef)
        result[36*block:36*(block+1)]=np.maximum(0,result[36*block:36*(block+1)]+correction)
        records.append(dict(day=day,release=release,block=block,history_last=int(ids[-1]),
            history_count=len(ids),correction_kw=correction,coefficients=coef.tolist(),mean=mean.tolist(),scale=scale.tolist()))
    return result,records


def historical_ids(day):
    ids=np.arange(max(15,day-56),day)
    if len(ids)>28:ids=ids[np.rint(np.linspace(0,len(ids)-1,28)).astype(int)]
    return ids


def scenarios(day,release,pred,actual,features,conditional=False):
    begin=36*release;ids=historical_ids(day)
    sc=pred[day,release,begin:][None]+actual[ids,begin:]-pred[ids,release,begin:]
    weights=2.**(-(day-ids)/14);weights/=weights.sum()
    if conditional and release:
        x=features[ids,release];now=features[day,release]
        scale=np.maximum(20.,np.sqrt(np.average((x-np.average(x,axis=0,weights=weights))**2,axis=0,weights=weights)))
        distance=np.mean(((x-now)/scale)**2,axis=1)
        near=weights*np.exp(-.5*np.minimum(distance,50));near/=near.sum()
        weights=.25*weights+.75*near
    return sc,weights,ids

def refine_plan(grid,reserve,sc,w,p,state,original,begin,value,step):
    mean=w@sc;scale=np.maximum(20.,np.sqrt(w@((sc-mean)**2)))
    blocks=(np.arange(begin,144)//step-begin//step).astype(np.int64)
    zero=np.zeros(int(blocks.max())+1)
    args=(sc,w,state,p,value,original,blocks)
    objective=lambda z:closed_loop_loss(z,grid,scale,reserve,*args)
    before=float(objective(zero)[0])
    fit=minimize(objective,zero,jac=True,method='L-BFGS-B',bounds=[(-2.,2.)]*len(zero),
        options=dict(maxiter=100,ftol=1e-9,gtol=1e-4,maxls=30))
    z=fit.x if np.isfinite(fit.fun) and fit.fun<before else zero
    changed=np.maximum(0,grid+z[blocks]*scale)
    new_reserve=dp_reserves(changed,sc,w,p,value)
    old=float(closed_loop_loss(zero,changed,scale,reserve,*args)[0])
    if closed_loop_loss(zero,changed,scale,new_reserve,*args)[0]<old:reserve=new_reserve
    after=float(closed_loop_loss(zero,changed,scale,reserve,*args)[0])
    assert after<=before+1e-6
    return changed,reserve,dict(refine_before=before,refine_after=after,refine_dimensions=len(zero),refine_iterations=int(fit.nit))

def simulate_day(day,state,pred,actual,prices,features,config):
    original,rule,log=plan_stage(day,0,state,None,pred,actual,prices,features,config)
    booked=original.copy();f=np.zeros((5,144));s=np.zeros(145);s[0]=state
    trades=np.zeros((4,2,144));snapshots=np.zeros((4,144));snapshots[0]=booked
    plans=np.full((4,144),np.nan);plans[0]=original;logs=[log]
    for release in range(4):
        begin=release*36
        if release:
            plan,reserve,log=plan_stage(day,release,float(s[begin]),original,pred,actual,prices,features,config)
            change=plan[:36]-booked[begin:begin+36]
            trades[release,0,begin:begin+36]=np.maximum(change,0)
            trades[release,1,begin:begin+36]=np.maximum(-change,0)
            booked[begin:begin+36]=plan[:36];rule[begin:]=reserve
            snapshots[release]=booked;plans[release,begin:]=plan;logs.append(log)
        for t in range(begin,begin+36):
            f[0,t]=booked[t]
            f[1,t],f[2,t],f[3,t],f[4,t],s[t+1]=step_action(float(actual[day,t]),float(booked[t]),float(s[t]),float(rule[t]))
    return original,f,s,rule,trades,snapshots,plans,logs

def clock_label(slot):
    return f'{slot//6}:{slot%6*10:02d}'


def emergency_events(values):
    """将连续紧急区间合并。下游使用原高精度电量，显示舍入不参与求和。"""
    events = []; t = 0
    while t < 144:
        if values[t] <= 1e-6:
            t += 1
            continue
        start = t
        while t < 144 and values[t] > 1e-6:
            t += 1
        events.append((f'{clock_label(start)}-{clock_label(t)}',float(values[start:t].sum())))
    return events


def required_input(value,label):
    if not str(value).strip():
        raise ValueError(label+'路径尚未填写。请编辑脚本顶部相关路径，或使用命令行参数。')
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(label+'文件不存在，请检查配置或命令行参数。')
    return path


def array_digest(values):
    """只用于核对附件是否与留档相符，不作为模型特征。"""
    return hashlib.sha256(np.ascontiguousarray(values, dtype=np.float64).view(np.uint8)).hexdigest()


def load_archived_components(path, dates, load, pv, prices):
    """读取显式提供的样本外四分量；不读取最终购电计划、SOC或费用。"""
    saved = read_npz(path)
    needed = {'dates', 'component_names', 'component_predictions_kw', 'provenance_json',
              'load_sha256', 'pv_sha256', 'price_sha256'}
    if set(saved) != needed:
        raise ValueError('预测留档字段不符；请使用本次附带的原交付四分量预测.npz。')
    if list(saved['component_names']) != COMPONENT_NAMES:
        raise ValueError('预测留档的四分量名称或顺序不符。')
    if not np.array_equal(saved['dates'], dates.strftime('%Y-%m-%d').to_numpy(dtype='U10')):
        raise ValueError('预测留档日期与附件不符。')
    for name, values in [('load',load),('pv',pv),('price',prices)]:
        if str(saved[name+'_sha256'].item()) != array_digest(values):
            raise ValueError('预测留档与当前附件的'+name+'数据不同，不能混用；请选择retrain模式。')
    components = saved['component_predictions_kw']
    if components.shape != (2,4,365,144) or not np.isfinite(components[:,:,START:]).all():
        raise ValueError('预测留档维度或有效预测不完整。')
    if (components[:,:,START:] < 0).any():
        raise ValueError('预测留档包含负功率。')
    # 留档内的来源信息可能包含原机器路径或其他个人信息，因此不复制到新输出。
    provenance = dict(archive_file_sha256=digest(path),
                      archived_provenance_sha256=hashlib.sha256(
                          str(saved['provenance_json'].item()).encode('utf-8')).hexdigest(),
                      base_models_retrained_in_this_run=False,
                      validation_scope='本次校验输入匹配；逐日融合与调度重新计算。基础预测来源及因果性依赖原研究留档审计。')
    return components, provenance

# 四、主方案与基准共用的运行入口。以下金额只用于核对，绝不参与训练和优化。
REFERENCE_COSTS = {'main':13016898.126054544,'baseline':13094652.121583482}
POLICIES = {'main':dict(nowcast=True,conditional=True,refine_step=1),'baseline':{}}
OUTPUT_NAMES = {'main':'主方案_1301','baseline':'对照方案_1309'}


def read_attachment3(path):
    workbook=load_workbook(path,read_only=True,data_only=True)
    raw=np.full((365,4,24),np.nan);seen=set();current=None
    try:
        for row in list(workbook.worksheets[0].values)[1:]:
            if row[0] is not None and str(row[0]).strip():current=pd.Timestamp(row[0])
            if current is None:raise ValueError('附件3首行缺少日期。')
            day=(current-pd.Timestamp('2025-01-01')).days
            hour=row[1].hour if hasattr(row[1],'hour') else int(str(row[1]).split(':')[0])
            if not 0<=day<365 or hour not in (0,6,12,18) or (day,hour) in seen:
                raise ValueError('附件3日期/发布时间无效或重复。')
            seen.add((day,hour));raw[day,hour//6]=np.asarray(row[2:26],float)
    finally:workbook.close()
    if len(seen)!=1460 or not np.isfinite(raw).all() or raw.min()<0:
        raise ValueError('附件3须包含365天×4个发布时间×24个有效预报。')
    return raw


def plan_stage(day,release,state,original,pred,actual,prices,features,config):
    """日内只提交紧接六小时，余下日内曲线用于储能价值的辅助规划。"""
    begin=36*release;p=prices[begin:];value=float(prices[:36].min()/.9)
    sc,w,ids=scenarios(day,release,pred,actual,features,config.get('conditional',False))
    if not release:
        plan,reserve,details=make_main_plan(sc,w,prices,state)
    else:
        plan,reserve,details=make_adjustment(sc,w,p,state,original[begin:],begin,value,weighted_quantile)
    details.update(day=day,release=release,last=int(ids[-1]),scenario_count=len(ids),
        effective_scenarios=float(1/(w@w)),state=float(state),tree=False)
    if release and config.get('refine_step'):
        plan,reserve,extra=refine_plan(plan,reserve,sc,w,p,state,original[begin:],begin,value,config['refine_step'])
        details.update(extra)
    assert int(ids[-1])<day
    return plan,reserve,details


def prepare_information(dates,load,pv,raw,components,output):
    """重算四分量融合、发布版本融合及低维负载校正，不读取最终答案。"""
    midnight=[];blend_logs=[]
    for component,actual,cube in zip(['load','pv'],[load,pv],components):
        forecast,records=online_blend(actual,cube.transpose(1,0,2),dates)
        midnight.append(forecast)
        for record in records:
            day=dates.get_loc(record['date'])
            assert record['weight_last_day']<day and record['monthly_choice_last_day']<day
            blend_logs.append(dict(component=component,**record))
    base_load,base_pv=midnight
    external=np.full((365,4,144),np.nan);blend=np.full_like(external,np.nan)
    alphas=np.full((365,4),np.nan);pvlogs=[]
    for day in range(365):
        for release in range(4):
            begin=release*36
            external[day,release,begin:]=interpolate_release(day,release,raw,pv)
    for day in range(START,365):
        for release in range(4):
            begin=release*36
            alpha,n,last=calibration_weight(day,release,pv,base_pv,external)
            assert last<day
            alphas[day,release]=alpha
            blend[day,release,begin:]=(1-alpha)*base_pv[day,begin:]+alpha*external[day,release,begin:]
            pvlogs.append(dict(date=str(dates[day].date()),release_hour=release*6,external_weight=alpha,history_last=last,history_count=n))
    base=(base_load[:,None,:]-blend)/6;corrected=base.copy()
    features=np.zeros((365,4,3));loadbank=np.broadcast_to(base_load[:,None,:],base.shape).copy();loadlogs=[]
    for day in range(START,365):
        for release in range(4):
            features[day,release]=observed_features(day,release,load,base_load,blend)
            prediction,records=predict_load(day,release,load,base_load)
            assert all(x['history_last']<day for x in records)
            loadlogs.extend(records);loadbank[day,release]=prediction
            corrected[day,release]=(prediction-blend[day,release])/6
    bank=dict(base=base,load_corrected=corrected,features=features,load_nowcast_kw=loadbank,
        predicted_load_kw=base_load,predicted_pv_kw=base_pv,blended_pv_kw=blend,external_pv_kw=external,external_weights=alphas)
    atomic_npz(output/'重建预测.npz',**bank)
    write_json(output/'每日基础融合记录.json',blend_logs)
    write_json(output/'日内负载拟合记录.json',loadlogs)
    pd.DataFrame(pvlogs).to_csv(output/'光伏发布融合权重.csv',index=False,encoding='utf-8-sig')
    return bank


def january_warmup(dates,load,pv,prices,bank,output):
    """6000仅在1月1日赋值；前15天闲置电池，16日起按第二问连续预运行。"""
    state=6000.;net=(load-pv)*DT;pred=(bank['predicted_load_kw']-bank['predicted_pv_kw'])*DT
    soc=np.full((31,145),6000.);flows=np.zeros((31,5,144));records=[]
    for day in range(31):
        if day<START:
            grid=np.maximum(net[day-1],0) if day else np.zeros(144)
            emergency=np.maximum(net[day]-grid,0)
            flows[day]=np.stack([grid,np.zeros(144),np.zeros(144),emergency,grid+emergency-net[day]])
        else:
            sc,w,last=scenario_set(net[:day],pred[:day],pred[day]);assert last<day
            grid,reserve,_=make_main_plan(sc,w,prices,state)
            flows[day],soc[day]=execute(net[day],grid,reserve,state)
        state=float(soc[day,-1]);records.append(dict(date=str(dates[day].date()),soc_start=float(soc[day,0]),soc_end=state))
    atomic_npz(output/'一月预运行.npz',soc=soc,flows=flows)
    pd.DataFrame(records).to_csv(output/'一月储电量.csv',index=False,encoding='utf-8-sig')
    return state


def replay_q3(name,dates,load,pv,prices,bank,initial,output):
    config=POLICIES[name];pred=bank['load_corrected' if config.get('nowcast') else 'base']
    actual=(load-pv)/6;value=float(prices[:36].min()/.9);state=initial
    output.mkdir(parents=True,exist_ok=True)
    a=dict(original_plan=np.zeros((334,144)),flows=np.zeros((334,5,144)),soc=np.zeros((334,145)),
        reserves=np.zeros((334,144)),trades=np.zeros((334,4,2,144)),bookings=np.zeros((334,4,144)),auxiliary=np.full((334,4,144),np.nan))
    records=[];logs=[];started=time.time()
    for day in range(FORMAL,365):
        result=simulate_day(day,state,pred,actual,prices,bank['features'],config)
        for key,x in zip(a,result[:-1]):a[key][day-FORMAL]=x
        original,f,s,_,_,_,_,detail=result;logs.extend(detail);state=float(s[-1])
        pc=float(prices@original);up=float(1.5*prices@np.maximum(f[0]-original,0))
        refund=float(.5*prices@np.maximum(original-f[0],0));ec=float(5*prices@f[3]);cost=pc+up-refund+ec
        records.append(dict(date=str(dates[day].date()),planned_cost=pc,up_cost=up,refund=refund,emergency_cost=ec,total_cost=cost,
            inventory_adjusted_cost=cost+value*(s[0]-s[-1]),soc_start=float(s[0]),soc_end=state,
            emergency_kwh=float(f[3].sum()),spill_kwh=float(f[4].sum()),emergency_intervals=int((f[3]>1e-6).sum()),
            up_kwh=float(np.maximum(f[0]-original,0).sum()),down_kwh=float(np.maximum(original-f[0],0).sum())))
        if (day-FORMAL)%30==0 or day==364:print(f'[{name}] {dates[day].date()} {time.time()-started:.1f}s',flush=True)
    daily=pd.DataFrame(records)
    summary={k:float(daily[k].sum()) for k in ['planned_cost','up_cost','refund','emergency_cost','total_cost','inventory_adjusted_cost','emergency_kwh','spill_kwh','up_kwh','down_kwh']}
    summary.update(policy=name,config=config,days=334,initial_soc=initial,final_soc=state,
        emergency_days=int((daily.emergency_intervals>0).sum()),emergency_intervals=int(daily.emergency_intervals.sum()),
        archived_cost_reference=REFERENCE_COSTS[name],difference_from_reference=summary['total_cost']-REFERENCE_COSTS[name],seconds=time.time()-started)
    atomic_npz(output/'arrays.npz',dates=dates[FORMAL:].strftime('%Y-%m-%d').to_numpy(dtype='U10'),**a)
    daily.to_csv(output/'daily.csv',index=False,encoding='utf-8-sig');write_json(output/'decisions.json',logs)
    write_json(output/'summary.json',summary)
    audit=verify_dispatch(a,actual[FORMAL:],prices,initial,summary)
    write_json(output/'物理与结算核验.json',audit)
    save_detail_tables(dates,a,prices,load,pv,daily,output)
    return a,summary


def verify_dispatch(a,actual,prices,initial,summary):
    f=a['flows'];s=a['soc'];g=a['original_plan'];tr=a['trades']
    errors=dict(balance=float(np.max(abs(f[:,0]+f[:,2]+f[:,3]-actual-f[:,1]-f[:,4]))),
        soc_equation=float(np.max(abs(np.diff(s,axis=1)-.9*f[:,1]+f[:,2]/.9))),
        midnight=float(np.max(abs(s[1:,0]-s[:-1,-1]))),initial=float(abs(s[0,0]-initial)))
    assert max(errors.values())<1e-7
    assert s.min()>=S_MIN-1e-7 and s.max()<=S_MAX+1e-7
    assert f.min()>=-1e-7 and f[:,1:3].max()<=P_MAX+1e-7
    assert np.minimum(f[:,1],f[:,2]).max()<1e-8 and not np.count_nonzero(tr[:,0])
    book=g.copy()
    for r in range(4):
        begin=r*36;stop=begin+36
        if r:
            assert not np.count_nonzero(tr[:,r,:,:begin]) and not np.count_nonzero(tr[:,r,:,stop:])
            book+=tr[:,r,0]-tr[:,r,1]
        assert np.max(abs(book-a['bookings'][:,r]))<1e-8
        assert np.max(abs(f[:,0,begin:stop]-book[:,begin:stop]))<1e-8
    cost=float(np.sum(g@prices+np.maximum(f[:,0]-g,0)@(1.5*prices)-np.maximum(g-f[:,0],0)@(.5*prices)+f[:,3]@(5*prices)))
    assert abs(cost-summary['total_cost'])<1e-6
    return dict(errors=errors,total_recomputed=cost,contract_windows_passed=True,
        current_interval_feedback_assumption=True,strict_interval_start_nonanticipativity_claimed=False,
        same_year_architecture_selection_disclosed=True)


def save_detail_tables(dates,a,prices,load,pv,daily,output):
    n=334*144;f=a['flows'];g=a['original_plan'];s=a['soc']
    frame=pd.DataFrame(dict(date=np.repeat(dates[FORMAL:].strftime('%Y-%m-%d'),144),slot=np.tile(np.arange(144),334),
        interval=np.tile([f'{clock_label(t)}-{clock_label(t+1)}' for t in range(144)],334),
        price_yuan_per_kwh=np.tile(prices,334),load_kw=load[FORMAL:].ravel(),pv_kw=pv[FORMAL:].ravel(),
        original_plan_kwh=g.ravel(),adjusted_purchase_kwh=f[:,0].ravel(),adjustment_delta_kwh=(f[:,0]-g).ravel(),
        charge_kwh=f[:,1].ravel(),discharge_kwh=f[:,2].ravel(),emergency_kwh=f[:,3].ravel(),surplus_kwh=f[:,4].ravel(),
        soc_start_kwh=s[:,:-1].ravel(),soc_end_kwh=s[:,1:].ravel()))
    assert len(frame)==n
    frame.to_csv(output/'interval_strategy.csv',index=False,encoding='utf-8-sig')
    blocks=[];events=[]
    for i,dt in enumerate(dates[FORMAL:]):
        for b in range(6):blocks.append(dict(date=str(dt.date()),interval=f'{b*4}:00-{(b+1)*4}:00',
            charge_kwh=float(f[i,1,b*24:(b+1)*24].sum()),discharge_kwh=float(f[i,2,b*24:(b+1)*24].sum()),
            soc_0000_kwh=float(s[i,0]),soc_2400_kwh=float(s[i,-1])))
        for interval,quantity in emergency_events(f[i,3]) or [('无',0.)]:events.append(dict(date=str(dt.date()),interval=interval,emergency_kwh=quantity))
    pd.DataFrame(blocks).to_csv(output/'storage_blocks.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame(events).to_csv(output/'emergency_events.csv',index=False,encoding='utf-8-sig')
    selected=['2025-03-20','2025-06-21','2025-09-23','2025-12-21']
    frame[frame.date.isin(selected)&frame.slot.isin([60,72,84,96,108,120])].to_csv(output/'paper_table1_intervals.csv',index=False,encoding='utf-8-sig')
    daily[daily.date.isin(selected)].to_csv(output/'paper_table1_daily.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame(blocks).query('date in @selected').to_csv(output/'paper_table2.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame(events).query('date in @selected').to_csv(output/'paper_table3.csv',index=False,encoding='utf-8-sig')


def validate_template(path):
    w=load_workbook(path,read_only=True)
    try:
        if w.sheetnames!=['计划购电量','调整购电量','充放电量','紧急购电量']:
            raise ValueError('请提供附件五原始result3模板，包含四张规定工作表。')
        for sheet in w.worksheets[:2]:
            if (sheet.max_row,sheet.max_column)!=(335,147) or sheet['B1'].value!='0:10-0:20' or sheet['EO1'].value!='0:00-0:10+1':
                raise ValueError('result3购电表日期/时间布局不符合原模板。')
            if not pd.DatetimeIndex([row[0] for row in sheet.iter_rows(min_row=2,max_row=335,max_col=1,values_only=True)]).equals(pd.date_range('2025-02-01','2025-12-31')):
                raise ValueError('result3购电表日期必须连续覆盖2—12月。')
        if [x.value for x in w.worksheets[2][1]]!=['日期','时间段','充电量','放电量','时刻','储电量']:
            raise ValueError('储能表头不符。')
        if [x.value for x in w.worksheets[3][1]]!=['日期','购电时间段','购电量']:
            raise ValueError('紧急购电表头不符。')
    finally:w.close()


def export_result3(template,destination,dates,a,prices):
    """Python直接填写原模板。调整表填写最终常规购电量A，费用为最终合同结算。"""
    if Path(template).resolve()==Path(destination).resolve():raise ValueError('禁止覆盖模板。')
    validate_template(template);w=load_workbook(template)
    plan,adjust,storage,emergency=w.worksheets
    def prototypes(sheet,rows,cols):
        return [([dict(value=sheet.cell(r,c).value,style=copy.copy(sheet.cell(r,c)._style)) for c in range(1,cols+1)],copy.copy(sheet.row_dimensions[r])) for r in rows]
    srows=prototypes(storage,range(2,8),6);erows=prototypes(emergency,range(2,5),3)
    storage.delete_rows(2,storage.max_row-1);emergency.delete_rows(2,emergency.max_row-1)
    def copy_row(sheet,row,prototype):
        cells,dimensions=prototype;dimensions=copy.copy(dimensions);dimensions.index=row;sheet.row_dimensions[row]=dimensions
        for column,saved in enumerate(cells,1):
            cell=sheet.cell(row,column);cell._style=copy.copy(saved['style']);cell.value=saved['value']
    erow=2
    for i,dt in enumerate(dates[FORMAL:]):
        row=i+2;g=a['original_plan'][i];f=a['flows'][i]
        for sheet,quantities,cost in [(plan,g,float(prices@g)),(adjust,f[0],settlement(f[0],g,prices))]:
            for t in range(144):sheet.cell(row,t+2,float(quantities[t])).number_format='0.00'
            sheet.cell(row,146,float(quantities.sum())).number_format='0.00'
            sheet.cell(row,147,cost).number_format='0.00'
        for block in range(6):
            r=2+i*6+block;copy_row(storage,r,srows[block]);storage.cell(r,1).value=dt.to_pydatetime() if block==0 else None
            for column,channel in [(3,1),(4,2)]:storage.cell(r,column,float(f[channel,24*block:24*(block+1)].sum())).number_format='0.00'
            storage.cell(r,6).value=float(a['soc'][i,0]) if block==0 else (float(a['soc'][i,-1]) if block==1 else None)
            if block<2:storage.cell(r,6).number_format='0.00'
        events=emergency_events(f[3]) or [('无',0.)];count=max(3,len(events))
        for j in range(count):
            copy_row(emergency,erow,erows[0 if j==0 else (2 if j==count-1 else 1)])
            emergency.cell(erow,1).value=dt.to_pydatetime() if j==0 else None
            emergency.cell(erow,2).value=events[j][0] if j<len(events) else None
            emergency.cell(erow,3).value=events[j][1] if j<len(events) else None
            emergency.cell(erow,3).number_format='0.00';erow+=1
    mapping=[dict(column_index=t+2,template_header=plan.cell(1,t+2).value,
        model_interval=f'{clock_label(t)}-{clock_label(t+1)}',attachment2_endpoint=clock_label(t+1)) for t in range(144)]
    destination=Path(destination);destination.parent.mkdir(parents=True,exist_ok=True)
    w.save(destination);w.close()
    pd.DataFrame(mapping).to_csv(destination.parent/'模板与模型时间映射.csv',index=False,encoding='utf-8-sig')
    verify_result3(template,destination,dates,a,prices)


def verify_result3(template,destination,dates,a,prices):
    source=load_workbook(template);w=load_workbook(destination,data_only=True)
    try:
        assert w.sheetnames==source.sheetnames
        for before,after in zip(source,w):
            assert [c.value for c in before[1]]==[c.value for c in after[1]]
            assert list(before.merged_cells.ranges)==list(after.merged_cells.ranges)
            assert before.freeze_panes==after.freeze_panes
            for key,dim in before.column_dimensions.items():assert after.column_dimensions[key].width==dim.width
        for sheet,quantities in zip(w.worksheets[:2],[a['original_plan'],a['flows'][:,0]]):
            values=np.asarray([[sheet.cell(r,c).value for c in range(2,146)] for r in range(2,336)],float)
            assert np.max(abs(values-quantities))<1e-8
            assert (sheet.max_row,sheet.max_column)==(335,147)
            for i in range(334):assert abs(sheet.cell(i+2,146).value-quantities[i].sum())<1e-7
        for i in range(334):
            assert abs(w.worksheets[0].cell(i+2,147).value-prices@a['original_plan'][i])<1e-7
            assert abs(w.worksheets[1].cell(i+2,147).value-settlement(a['flows'][i,0],a['original_plan'][i],prices))<1e-7
            r=2+i*6;storage=w.worksheets[2]
            assert pd.Timestamp(storage.cell(r,1).value)==dates[FORMAL+i]
            assert abs(storage.cell(r,6).value-a['soc'][i,0])<1e-7
            assert abs(storage.cell(r+1,6).value-a['soc'][i,-1])<1e-7
            for b in range(6):
                for col,ch in [(3,1),(4,2)]:assert abs(storage.cell(r+b,col).value-a['flows'][i,ch,b*24:(b+1)*24].sum())<1e-7
        urgent=sum(float(row[2] or 0) for row in w.worksheets[3].iter_rows(min_row=2,values_only=True))
        assert abs(urgent-a['flows'][:,3].sum())<1e-6 and w.worksheets[2].max_row==2005
        assert not any(c.data_type=='e' for sheet in w for row in sheet for c in row)
        write_json(Path(destination).parent/'Excel核验.json',dict(headers_and_layout_preserved=True,
            all_96192_purchase_values_checked=True,all_storage_blocks_checked=True,emergency_kwh=urgent,
            adjustment_sheet_means_final_regular_purchase=True,adjustment_cost_means_total_contract_settlement=True))
    finally:source.close();w.close()


def main():
    parser=argparse.ArgumentParser(description='第三问独立求解：完整重训/显式基础预测留档→融合→连续调度→result3.xlsx')
    for name,default in [('attachment1',ATTACHMENT1_PATH),('attachment2',ATTACHMENT2_PATH),('attachment3',ATTACHMENT3_PATH),('template',RESULT3_TEMPLATE_PATH)]:parser.add_argument('--'+name,default=default)
    parser.add_argument('--mode',choices=['retrain','archive'],default=RUN_MODE)
    parser.add_argument('--archive',default=ARCHIVED_PREDICTIONS_PATH)
    parser.add_argument('--output-dir',default=OUTPUT_DIRECTORY)
    parser.add_argument('--policy',choices=['main','baseline','both'],default='main')
    parser.add_argument('--workers',type=int,choices=[1,2],default=TRAINING_WORKERS)
    parser.add_argument('--threads',type=int,default=LIGHTGBM_THREADS)
    parser.add_argument('--end-index',type=int,default=365,help='仅用于重训断点；少于365不生成全年结果')
    args=parser.parse_args();started=time.time()
    paths=[required_input(getattr(args,k),k) for k in ['attachment1','attachment2','attachment3','template']]
    archive=required_input(args.archive,'基础预测留档') if args.mode=='archive' else None
    if args.mode=='retrain' and str(args.archive).strip():raise ValueError('retrain模式禁止指定基础预测留档。')
    if args.threads<1 or not 16<=args.end_index<=365:raise ValueError('线程数至少1，end-index须在16—365。')
    if args.mode=='archive' and args.end_index!=365:raise ValueError('archive模式应完整复算全年。')
    output=Path(args.output_dir).expanduser().resolve() if str(args.output_dir).strip() else Path(__file__).resolve().parent/('完整重训输出' if args.mode=='retrain' else '留档复算输出')
    if any((output/OUTPUT_NAMES[n]/'result3.xlsx').resolve() in paths for n in OUTPUT_NAMES):raise ValueError('输出不得覆盖输入。')
    validate_template(paths[3]);dates,load,pv,prices=read_inputs(paths[0],paths[1]);raw=read_attachment3(paths[2])
    versions={name:metadata.version(name) for name in ['numpy','pandas','scipy','scikit-learn','lightgbm','numba','llvmlite','openpyxl','threadpoolctl']};versions['python']=sys.version
    input_files=list(zip(['attachment1','attachment2','attachment3','template'],paths))
    if archive is not None:input_files.append(('archive',archive))
    inputs_hash={label:digest(path) for label,path in input_files}
    signature=hashlib.sha256((digest(__file__)+json.dumps(versions,sort_keys=True)+str(args.threads)+array_digest(load)+array_digest(pv)).encode()).hexdigest()
    provenance=dict(run_mode=args.mode,versions=versions,source_sha256=digest(__file__),input_sha256=inputs_hash,
        signature=signature,workers=args.workers,threads=args.threads,base_models_retrained_in_this_run=args.mode=='retrain',
        no_external_project_modules=True,no_embedded_forecasts_or_answers=True,status='running')
    existing=output/'运行来源与环境.json'
    if existing.exists():
        previous=json.loads(existing.read_text(encoding='utf-8'))
        if any(previous.get(k)!=provenance[k] for k in ['run_mode','source_sha256','input_sha256']):raise ValueError('输出目录属于不同运行，请另设目录。')
    output.mkdir(parents=True,exist_ok=True);write_json(existing,provenance)
    if args.mode=='archive':
        print('显式读取基础四分量预测；所有融合、1月预运行和第三问调度重新计算。',flush=True)
        components,record=load_archived_components(archive,dates,load,pv,prices);write_json(output/'预测留档信息.json',record)
    else:
        cache=output/'训练缓存';cache.mkdir(exist_ok=True)
        jobs=[(name,x,dates.strftime('%Y-%m-%d').to_numpy(dtype='U10'),str(cache/f'{name}.npz'),signature,args.threads,args.end_index) for name,x in [('load',load),('pv',pv)]]
        if args.workers==1:
            for job in jobs:train_component(*job)
        else:
            with ProcessPoolExecutor(max_workers=2) as pool:
                futures=[pool.submit(train_component,*job) for job in jobs]
                for future in futures:future.result()
        if args.end_index<365:
            provenance.update(status='partial_training',end_index=args.end_index,elapsed_seconds=time.time()-started);write_json(existing,provenance)
            print('阶段训练已保存；默认end-index=365可继续。尚未生成全年结果。',flush=True);return
        states=[read_npz(cache/f'{name}.npz') for name in ['load','pv']]
        assert all(s['done'][START:].all() and np.isfinite(s['predictions'][:,START:]).all() for s in states)
        components=np.stack([s['predictions'] for s in states]);records=[]
        for component,state in zip(['load','pv'],states):
            for day in range(START,365):
                choice_day=START if day<FORMAL else day-dates[day].day+1
                history=validation_origins(choice_day,START,35,3)
                records.append(dict(component=component,date=str(dates[day].date()),max_training_index=day-1,
                    parameter_choice_index=choice_day,choice_history_last=int(history[-1]) if len(history) else -1,
                    lgb_config=state['selected'][day,0],ridge_config=state['selected'][day,1],calendar_alpha=[.01,.1,1.][int(state['calendar_active'][day])]))
        pd.DataFrame(records).to_csv(output/'训练与选参记录.csv',index=False,encoding='utf-8-sig')
    atomic_npz(output/'本次基础预测.npz',component_predictions_kw=components)
    with threadpool_limits(1):
        bank=prepare_information(dates,load,pv,raw,components,output)
        initial=january_warmup(dates,load,pv,prices,bank,output)
        summaries=[]
        for name in (['main','baseline'] if args.policy=='both' else [args.policy]):
            dest=output/OUTPUT_NAMES[name]
            a,summary=replay_q3(name,dates,load,pv,prices,bank,initial,dest)
            summary.update(run_mode=args.mode,base_models_retrained=args.mode=='retrain');write_json(dest/'summary.json',summary)
            export_result3(paths[3],dest/'result3.xlsx',dates,a,prices);summaries.append(summary)
            print(name+' TOTAL='+format(summary['total_cost'],'.8f'),flush=True)
    write_json(output/'方案费用汇总.json',summaries)
    provenance.update(status='complete',elapsed_seconds=time.time()-started);write_json(existing,provenance)
    print('第三问结果生成完成。',flush=True)


if __name__=='__main__':
    from multiprocessing import freeze_support
    freeze_support()
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    try:main()
    except (ValueError,FileNotFoundError,PermissionError) as exc:
        print('运行失败：'+str(exc),file=sys.stderr);raise SystemExit(2)
