from __future__ import annotations

import argparse
import csv
import io
import os
import sys
import tempfile
from contextlib import redirect_stdout
from dataclasses import dataclass
from datetime import datetime, time
from pathlib import Path

try:
    import numpy as np
    from openpyxl import load_workbook
    from scipy.optimize import linprog
except ModuleNotFoundError as exc:
    print(f"缺少依赖 {exc.name}。请先执行：python -m pip install -r requirements.txt", file=sys.stderr)
    raise SystemExit(1) from exc


BASE_DIR = Path(__file__).resolve().parent
PERIODS = 144
TOLERANCE = 1.0e-6


def clock_label(minutes: int, *, next_day: bool = False) -> str:
    if next_day and minutes >= 1440:
        return f"{minutes // 60 % 24}:{minutes % 60:02d}+1"
    return f"{minutes // 60}:{minutes % 60:02d}"


def model_interval(index: int) -> str:
    return f"{clock_label(index * 10)}-{clock_label((index + 1) * 10)}"


def template_interval(index: int) -> str:
    return f"{clock_label((index + 1) * 10, next_day=True)}-{clock_label((index + 2) * 10, next_day=True)}"


def source_minutes(value: object) -> int:
    """兼容附件中混用的 Excel 时间值、23:50、0:00+1。"""
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f"附件时间包含非整分钟值：{value}")
        return value.hour * 60 + value.minute
    text = str(value).strip()
    next_day = text.endswith("+1")
    if next_day:
        text = text[:-2]
    try:
        hour_text, minute_text = text.split(":")
        hour, minute = int(hour_text), int(minute_text)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"无法识别附件时间：{value}") from exc
    if not (0 <= hour <= 24 and 0 <= minute < 60) or (hour == 24 and minute != 0):
        raise ValueError(f"附件时间超出范围：{value}")
    return hour * 60 + minute + (1440 if next_day else 0)


@dataclass(frozen=True)
class BatteryParameters:
    interval_hours: float = 1.0 / 6.0
    capacity_kwh: float = 12000.0
    soc_min_kwh: float = 1200.0
    soc_max_kwh: float = 10800.0
    max_power_kw: float = 5000.0
    efficiency: float = 0.90
    initial_soc_kwh: float = 6000.0

    @property
    def max_interval_energy_kwh(self) -> float:
        return self.max_power_kw * self.interval_hours


def read_attachment1(path: Path) -> tuple[list[object], np.ndarray, np.ndarray, np.ndarray]:
    """读取附件 1 的时间、电价、负载功率和光伏预测功率。"""
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    rows = [row for row in rows if any(value is not None for value in row)]
    if len(rows) != 144:
        raise ValueError(f"附件 1 应包含 144 个时段，实际为 {len(rows)} 个。")

    times = [row[0] for row in rows]
    for index, value in enumerate(times):
        if source_minutes(value) != (index + 1) * 10:
            raise ValueError(f"附件1第 {index + 2} 行时间错序、重复或缺失，应为 {clock_label((index + 1) * 10, next_day=True)}，实际为 {value}。")
    if any(len(row) < 4 or any(value is None or isinstance(value, bool) for value in row[1:4]) for row in rows):
        raise ValueError("附件1的电价、负载或光伏功率含有空值或非数值。")
    price = np.asarray([row[1] for row in rows], dtype=float)
    load_kw = np.asarray([row[2] for row in rows], dtype=float)
    pv_kw = np.asarray([row[3] for row in rows], dtype=float)

    if not all(np.all(np.isfinite(x)) for x in (price, load_kw, pv_kw)):
        raise ValueError("附件 1 含有空值或非数值。")
    if np.any(price < 0) or np.any(load_kw < 0) or np.any(pv_kw < 0):
        raise ValueError("电价、负载和光伏功率不应为负数。")
    return times, price, load_kw, pv_kw


def solve_dispatch(
    price: np.ndarray,
    load_kw: np.ndarray,
    pv_kw: np.ndarray,
    battery: BatteryParameters,
) -> dict[str, np.ndarray | float]:
    """求解确定性日内购电/储能线性规划。"""
    periods = len(price)
    if not (len(load_kw) == periods and len(pv_kw) == periods):
        raise ValueError("电价、负载和光伏序列长度必须一致。")

    # 决策变量依次为：购电 g[0:T]、充电 c[0:T]、放电 d[0:T]、
    # 储电量 s[0:T+1]。c、d 均按储能设备交流侧电量计量。
    g_slice = slice(0, periods)
    c_slice = slice(periods, 2 * periods)
    d_slice = slice(2 * periods, 3 * periods)
    s_slice = slice(3 * periods, 4 * periods + 1)
    variable_count = 4 * periods + 1

    objective = np.zeros(variable_count)
    objective[g_slice] = price
    # 极小的通量惩罚只用于消除等价最优解中的无意义充放电循环。
    objective[c_slice] = 1.0e-9
    objective[d_slice] = 1.0e-9

    # 功率平衡：g_t + PV_t*dt + d_t >= L_t*dt + c_t。
    a_ub = np.zeros((periods, variable_count))
    b_ub = (pv_kw - load_kw) * battery.interval_hours
    for t in range(periods):
        a_ub[t, g_slice.start + t] = -1.0
        a_ub[t, c_slice.start + t] = 1.0
        a_ub[t, d_slice.start + t] = -1.0

    # 储能状态：s_{t+1} = s_t + eta*c_t - d_t/eta。
    # 另加 s_0=s_T=6000 kWh。
    a_eq = np.zeros((periods + 2, variable_count))
    b_eq = np.zeros(periods + 2)
    eta = battery.efficiency
    for t in range(periods):
        a_eq[t, c_slice.start + t] = -eta
        a_eq[t, d_slice.start + t] = 1.0 / eta
        a_eq[t, s_slice.start + t] = -1.0
        a_eq[t, s_slice.start + t + 1] = 1.0
    a_eq[periods, s_slice.start] = 1.0
    b_eq[periods] = battery.initial_soc_kwh
    a_eq[periods + 1, s_slice.stop - 1] = 1.0
    b_eq[periods + 1] = battery.initial_soc_kwh

    max_energy = battery.max_interval_energy_kwh
    bounds = (
        [(0.0, None)] * periods
        + [(0.0, max_energy)] * periods
        + [(0.0, max_energy)] * periods
        + [(battery.soc_min_kwh, battery.soc_max_kwh)] * (periods + 1)
    )

    solution = linprog(
        c=objective,
        A_ub=a_ub,
        b_ub=b_ub,
        A_eq=a_eq,
        b_eq=b_eq,
        bounds=bounds,
        method="highs",
    )
    if not solution.success:
        raise RuntimeError(f"线性规划求解失败：{solution.message}")

    g = np.clip(solution.x[g_slice], 0.0, None)
    charge = np.clip(solution.x[c_slice], 0.0, None)
    discharge = np.clip(solution.x[d_slice], 0.0, None)
    soc = solution.x[s_slice]

    # 独立复核约束。允许的数值误差远小于结果文件保留的精度。
    supply_margin = g + pv_kw * battery.interval_hours + discharge - (
        load_kw * battery.interval_hours + charge
    )
    soc_residual = soc[1:] - soc[:-1] - eta * charge + discharge / eta
    if np.min(supply_margin) < -1.0e-6:
        raise RuntimeError("求解结果未通过电量平衡检查。")
    if np.max(np.abs(soc_residual)) > 1.0e-6:
        raise RuntimeError("求解结果未通过储能状态方程检查。")
    if np.any(soc < battery.soc_min_kwh - 1.0e-6) or np.any(
        soc > battery.soc_max_kwh + 1.0e-6
    ):
        raise RuntimeError("求解结果超出储电量上下限。")
    if max(abs(soc[0] - battery.initial_soc_kwh), abs(soc[-1] - battery.initial_soc_kwh)) > TOLERANCE:
        raise RuntimeError("0:00 与 24:00 的储电量不一致。")
    if max(float(charge.max()), float(discharge.max())) > max_energy + TOLERANCE:
        raise RuntimeError("求解结果超出储能充放电功率上限。")
    if np.max(np.minimum(charge, discharge)) > TOLERANCE:
        raise RuntimeError("求解结果出现同时充放电，需进一步处理互斥约束。")

    return {
        "grid_kwh": g,
        "charge_kwh": charge,
        "discharge_kwh": discharge,
        "soc_kwh": soc,
        "total_grid_kwh": float(np.sum(g)),
        "total_cost_yuan": float(np.dot(price, g)),
        "minimum_supply_margin_kwh": float(np.min(supply_margin)),
        "maximum_soc_residual_kwh": float(np.max(np.abs(soc_residual))),
    }


def validate_template(workbook: object) -> None:
    """确认使用官方原始布局，拒绝错位或被修改过的时间标签。"""
    if workbook.sheetnames != ["计划购电量", "充放电量"]:
        raise ValueError("模板应仅包含且依次为：计划购电量、充放电量。")
    purchase = workbook["计划购电量"]
    storage = workbook["充放电量"]
    if (purchase.max_row, purchase.max_column) != (145, 2):
        raise ValueError("模板计划购电量应为145行、2列。")
    if (storage.max_row, storage.max_column) != (7, 5):
        raise ValueError("模板充放电量应为7行、5列。")
    if [cell.value for cell in purchase[1]] != ["时间段", "购电量"]:
        raise ValueError("模板计划购电量表头与附件5不一致。")
    if [cell.value for cell in storage[1]] != ["时间段", "充电量", "放电量", "时刻", "储电量"]:
        raise ValueError("模板充放电量表头与附件5不一致。")
    for index in range(PERIODS):
        if purchase.cell(index + 2, 1).value != template_interval(index):
            raise ValueError(f"模板第 {index + 2} 行的时间标签与官方附件不一致。")
    for block in range(6):
        if storage.cell(block + 2, 1).value != f"{4 * block}:00-{4 * (block + 1)}:00":
            raise ValueError("模板充放电区间与附件5不一致。")
    if storage["D2"].value != "0:00" or storage["D3"].value != "24:00":
        raise ValueError("模板期初期末时刻与附件5不一致。")


def write_result_workbook(
    template_path: Path,
    output_path: Path,
    result: dict[str, np.ndarray | float],
) -> None:
    """只填158个结果单元格；临时文件保存成功后才替换目标。"""
    if template_path.resolve() == output_path.resolve():
        raise ValueError("输出文件不能覆盖原始模板。")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(template_path)
    try:
        validate_template(workbook)
    except Exception:
        workbook.close()
        raise

    grid = np.asarray(result["grid_kwh"])
    charge = np.asarray(result["charge_kwh"])
    discharge = np.asarray(result["discharge_kwh"])
    soc = np.asarray(result["soc_kwh"])

    purchase_sheet = workbook["计划购电量"]
    if purchase_sheet.max_row - 1 != len(grid):
        raise ValueError("result1.xlsx 模板的计划购电时段数与附件 1 不一致。")
    for row, value in enumerate(grid, start=2):
        purchase_sheet.cell(row=row, column=2, value=float(value))

    storage_sheet = workbook["充放电量"]
    # 144 个时段按时间顺序均分为 6 个连续的 4 小时区间，每段 24 个时段。
    for block in range(6):
        start = 24 * block
        end = start + 24
        storage_sheet.cell(
            row=2 + block, column=2, value=float(np.sum(charge[start:end]))
        )
        storage_sheet.cell(
            row=2 + block, column=3, value=float(np.sum(discharge[start:end]))
        )

    storage_sheet["E2"] = float(soc[0])
    storage_sheet["E3"] = float(soc[-1])
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(dir=output_path.parent, suffix=".xlsx", delete=False) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        workbook.close()
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def print_paper_tables(result: dict[str, np.ndarray | float]) -> None:
    """输出论文表 1 和表 2 需要填写的数值。"""
    # 论文指定时刻按自然日索引提取，不能按偏移的模板标签寻找。
    label_to_index = {model_interval(index): index for index in range(PERIODS)}
    requested_intervals = [
        "10:00-10:10",
        "12:00-12:10",
        "14:00-14:10",
        "16:00-16:10",
        "18:00-18:10",
        "20:00-20:10",
    ]

    grid = np.asarray(result["grid_kwh"])
    charge = np.asarray(result["charge_kwh"])
    discharge = np.asarray(result["discharge_kwh"])
    soc = np.asarray(result["soc_kwh"])

    print("\n表 1：指定时段购电量（kWh）")
    for label in requested_intervals:
        print(f"  {label}: {grid[label_to_index[label]]:.4f}")
    print(f"  全天购电量: {float(result['total_grid_kwh']):.4f} kWh")
    print(f"  全天购电费: {float(result['total_cost_yuan']):.4f} 元")

    print("\n表 2：4 小时分段充放电量（kWh）")
    for block in range(6):
        start = 24 * block
        end = start + 24
        print(
            f"  {4 * block}:00-{4 * (block + 1)}:00: "
            f"充电 {np.sum(charge[start:end]):.4f}, "
            f"放电 {np.sum(discharge[start:end]):.4f}"
        )
    print(f"  0:00 储电量: {soc[0]:.4f} kWh")
    print(f"  24:00 储电量: {soc[-1]:.4f} kWh")

    print("\n约束检查")
    print(
        "  最小供电裕量: "
        f"{float(result['minimum_supply_margin_kwh']):.3e} kWh"
    )
    print(
        "  最大储能状态残差: "
        f"{float(result['maximum_soc_residual_kwh']):.3e} kWh"
    )
    print(f"  最大同时充放电量: {np.max(np.minimum(charge, discharge)):.3e} kWh")


def write_supporting_files(output_dir: Path, times: list[object], price: np.ndarray,
                           load_kw: np.ndarray, pv_kw: np.ndarray,
                           result: dict[str, np.ndarray | float]) -> None:
    """明细和时间映射放在结果旁边，不给官方工作簿增加工作表。"""
    grid, charge, discharge, soc = [np.asarray(result[key]) for key in
                                   ("grid_kwh", "charge_kwh", "discharge_kwh", "soc_kwh")]
    with (output_dir / "dispatch_detail.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["period_index", "source_time", "model_interval", "template_row", "template_interval",
                         "price_yuan_per_kwh", "load_kw", "pv_kw", "grid_kwh", "charge_kwh",
                         "discharge_kwh", "soc_start_kwh", "soc_end_kwh", "curtailment_kwh", "cost_yuan"])
        for index in range(PERIODS):
            curtailment = grid[index] + pv_kw[index] / 6 + discharge[index] - load_kw[index] / 6 - charge[index]
            writer.writerow([index + 1, clock_label(source_minutes(times[index]), next_day=True),
                             model_interval(index), index + 2, template_interval(index),
                             price[index], load_kw[index], pv_kw[index], grid[index], charge[index],
                             discharge[index], soc[index], soc[index + 1], max(0.0, curtailment),
                             price[index] * grid[index]])
    with (output_dir / "模板与模型时间映射.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["模板行号", "模板原时间标签", "模型实际区间", "附件1行号", "附件1时间", "购电量(kWh)"])
        for index in range(PERIODS):
            writer.writerow([index + 2, template_interval(index), model_interval(index), index + 2,
                             clock_label(source_minutes(times[index]), next_day=True), grid[index]])
    buffer = io.StringIO()
    with redirect_stdout(buffer):
        print("问题1计算结果（自然日00:00—24:00，电量单位kWh，费用单位元）")
        print("附件1时间按右端点解释：0:10对应0:00—0:10；0:00+1对应23:50—24:00。")
        print("官方购电模板标签比模型实际区间晚10分钟；原表头与行顺序保留，对照详见时间映射CSV。")
        print_paper_tables(result)
    summary = buffer.getvalue()
    (output_dir / "运行摘要.txt").write_text(summary, encoding="utf-8-sig")
    print(summary)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="求解 2026 数模 C 题问题 1。")
    parser.add_argument("--input", type=Path, default=BASE_DIR / "附件/附件1.xlsx", help="默认读取本程序目录内的附件1。")
    parser.add_argument(
        "--template", type=Path, default=BASE_DIR / "附件/附件5/result1.xlsx", help="默认使用本程序目录内的官方模板。"
    )
    parser.add_argument("--output", type=Path, default=BASE_DIR / "result1.xlsx", help="结果路径；默认放在本程序旁，明细与摘要跟随输出目录。")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    args.input, args.template, args.output = [path.expanduser().resolve() for path in (args.input, args.template, args.output)]
    if args.output.suffix.lower() != ".xlsx":
        raise ValueError("输出文件必须使用.xlsx扩展名。")
    generated_paths = {args.output, *(args.output.parent / name for name in
                       ("dispatch_detail.csv", "模板与模型时间映射.csv", "运行摘要.txt"))}
    if generated_paths & {args.input, args.template}:
        raise ValueError("输出不能覆盖输入附件或原始模板。")
    for label, path in (("输入附件", args.input), ("结果模板", args.template)):
        if not path.is_file():
            raise FileNotFoundError(f"未找到{label}。请检查命令行参数或独立运行文件夹。")
    template = load_workbook(args.template, read_only=True)
    try:
        validate_template(template)
    finally:
        template.close()
    times, price, load_kw, pv_kw = read_attachment1(args.input)
    battery = BatteryParameters()
    result = solve_dispatch(price, load_kw, pv_kw, battery)
    write_result_workbook(args.template, args.output, result)
    write_supporting_files(args.output.parent, times, price, load_kw, pv_kw, result)
    print("\n完整结果已生成。")


if __name__ == "__main__":
    try:
        main()
    except (ValueError, RuntimeError, OSError) as exc:
        print(f"运行失败：{exc}", file=sys.stderr)
        if isinstance(exc, PermissionError):
            print("请关闭正在打开的结果Excel文件，并确认输出目录可写后重试。", file=sys.stderr)
        raise SystemExit(1) from exc
